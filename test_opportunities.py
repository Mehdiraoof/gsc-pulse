"""
Tests for the Resolved-adjacent v2.4 data work: gone pages, the opportunity engine and its tab, and
the new site settings. All data is invented.
"""

import json
import unittest
from unittest import mock

import deck_fixtures
import decks
import deck_runner
import opportunities as op
import reports
import resolved
import sites
import client_fixtures as fx
from test_reports import RID, ReportsTestCase, site_with

QUERIES_HEADER = ["Query", "Clicks", "Impressions", "CTR (%)", "Avg. Position"]
PAGES_HEADER = ["Page", "Clicks", "Impressions", "CTR (%)", "Avg. Position"]
PAIRS_HEADER = ["Query", "Page", "Clicks", "Impressions", "CTR (%)", "Avg. Position"]


def inputs(**kw):
    return op.Inputs(**kw)


def q(query, clicks, impr, pos):
    return op.QueryRow(query, clicks, impr, pos)


def pg(url, clicks, impr, pos):
    return op.PageRow(url, clicks, impr, pos)


def pair(query, url, clicks, impr, pos=5.0):
    return op.Pair(query, url, clicks, impr, pos)


def of_type(rows, kind):
    return [dict(zip(op.HEADER, r)) for r in rows if r[0] == kind]


def month_row(kind, url, **kw):
    base = {"Period": "2026-09", "Type": kind, "Page": url, "Clicks": "50", "Clicks prev": "100",
            "Clicks diff": "-50", "Impressions": "1000", "Impr prev": "1200", "Impr diff": "-200",
            "Position": "6.0", "Position prev": "6.0", "Position diff": "0"}
    base.update({k: str(v) for k, v in kw.items()})
    return base


class MathTests(unittest.TestCase):

    def test_expected_click_rate_by_position(self):
        self.assertEqual([op.expected_ctr(p) for p in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10)],
                         [.28, .15, .10, .07, .05, .04, .03, .025, .02, .018])
        self.assertEqual(op.expected_ctr(3.9), .10)             # the rank is the whole number part
        self.assertEqual(op.expected_ctr(0.4), .28)
        self.assertEqual(op.expected_ctr(11), .01)
        self.assertEqual(op.expected_ctr(20.9), .01)
        self.assertEqual(op.expected_ctr(25), .005)

    def test_gain_is_impressions_a_month_times_the_missing_click_rate_rounded_to_5(self):
        self.assertEqual(op.gain(3000, 0.02, 0.08), 60)
        self.assertEqual(op.gain(1000, 0.0153, 0.08), 20)       # 21.6 -> 20
        self.assertEqual(op.gain(1000, 0.0, 0.08), 25)          # 26.7 -> 25
        self.assertEqual(op.gain(100, 0.09, 0.08), 0)           # already above the target
        self.assertEqual(op.gain(5475, 120 / 5475, 0.08), 105)
        self.assertEqual(op.gain(3000, 0.02, 0.12), 100)        # the target comes from the setting

    def test_split_text(self):
        self.assertEqual(op.split_text("Bottom of page 1. Refresh content."), ("Bottom of page 1.", "Refresh content."))
        self.assertEqual(op.split_text("No action"), ("No action", ""))


class ParseTests(unittest.TestCase):

    def test_the_data_sheet_tabs_are_parsed_and_fragments_merged(self):
        queries = op.parse_queries([QUERIES_HEADER, ["a b", "10", "1,000", "1.0", "5.5"], ["", "1", "1", "1", "1"]])
        self.assertEqual((queries[0].query, queries[0].clicks, queries[0].impressions, queries[0].position),
                         ("a b", 10.0, 1000.0, 5.5))
        pages = op.parse_pages([PAGES_HEADER, ["https://x.test/p", "10", "100", "10", "4.0"],
                                ["https://x.test/p#frag", "5", "300", "1.7", "8.0"]])
        self.assertEqual(len(pages), 1)
        self.assertEqual((pages[0].clicks, pages[0].impressions), (15.0, 400.0))
        self.assertAlmostEqual(pages[0].position, (4.0 * 100 + 8.0 * 300) / 400)
        pairs = op.parse_pairs([PAIRS_HEADER, ["k", "https://x.test/p", "1", "10", "1", "3"],
                                ["k", "https://x.test/p#a", "2", "30", "1", "5"]])
        self.assertEqual((pairs[0].clicks, pairs[0].impressions), (3.0, 40.0))


class QueryRuleTests(unittest.TestCase):
    MIN = dict(min_impressions=300, target_ctr=0.08)

    def build(self, queries, **kw):
        return op.build(inputs(queries=queries, **self.MIN, **kw))

    def test_quick_win_is_position_4_to_20_with_enough_impressions_and_a_gain(self):
        rows = self.build([q("blue shoes", 120, 5475, 9.4), q("too high", 100, 3000, 3.9),
                           q("too low", 5, 3000, 20.5), q("small", 3, 299, 9.0),
                           q("edge four", 30, 3000, 4.0), q("edge twenty", 3, 3000, 20.0)],
                          pairs=[pair("blue shoes", "https://x.test/a/", 100, 4000),
                                 pair("blue shoes", "https://x.test/b/", 20, 1475)],
                          segment_of=lambda u: "Guides" if "/a/" in u else "")
        wins = of_type(rows, "quick_win")
        by = {w["keyword"]: w for w in wins}
        self.assertEqual(set(by), {"blue shoes", "edge four", "edge twenty"})
        first = by["blue shoes"]
        self.assertEqual((first["page"], first["potential_gain_month"], first["position"]),
                         ("https://x.test/a/", 105, 9.4))
        self.assertEqual((first["impressions_90d"], first["clicks_90d"], first["segment"]), (5475, 120, "Guides"))
        self.assertAlmostEqual(first["ctr"], 0.0219, places=4)
        self.assertEqual(first["expected_ctr"], 0.02)
        self.assertEqual(first["diagnosis"], "Ranks on page 1 but below the top 3.")
        self.assertEqual(by["edge twenty"]["diagnosis"], "Ranks on page 2 but below the top 3.")

    def test_a_query_that_already_beats_the_target_is_not_a_quick_win(self):
        self.assertEqual(self.build([q("popular", 600, 3000, 6.0)]), [])           # 20% CTR

    def test_brand_searches_are_skipped_by_every_query_rule(self):
        rows = self.build([q("Acme shoes", 120, 5475, 9.4), q("acme how to", 5, 3000, 8.0),
                           q("other shoes", 120, 5475, 9.4)], brand_terms=["acme"])
        self.assertEqual({r[1] for r in rows}, {"other shoes"})

    def test_ctr_fix_needs_a_top_5_position_and_under_half_the_expected_click_rate(self):
        rows = self.build([q("weak title", 20, 4000, 3.0), q("fine title", 300, 4000, 3.0),
                           q("borderline", 200, 4000, 3.0), q("not top five", 10, 4000, 6.5)])
        fixes = of_type(rows, "ctr_fix")
        self.assertEqual([f["keyword"] for f in fixes], ["weak title"])
        self.assertEqual(fixes[0]["potential_gain_month"], 100)
        self.assertEqual(fixes[0]["action"], "Rewrite title and meta description; check SERP features (images, AI Overview).")
        self.assertEqual(fixes[0]["diagnosis"], "Click rate 0.5% where about 10.0% is usual at this position.")

    def test_questions_in_persian_and_english(self):
        rows = self.build([q("قدیس یعنی چه", 11, 1756, 8.0), q("how to clean a bike", 0, 931, 9.0),
                           q("what is a trail", 40, 800, 2.0), q("what is a pin", 100, 600, 2.0),
                           q("is it good", 2, 90, 9.0), q("blue shoes", 2, 3000, 9.0)])
        questions = of_type(rows, "question")
        self.assertEqual({r["keyword"] for r in questions},
                         {"قدیس یعنی چه", "how to clean a bike", "what is a trail"})
        # a question at position 2 only counts when its click rate is below the expected 15%
        self.assertNotIn("what is a pin", {r["keyword"] for r in questions})
        # fewer than a third of the minimum impressions (100) is ignored
        self.assertNotIn("is it good", {r["keyword"] for r in questions})
        self.assertEqual(questions[0]["action"], "Write a short explainer page or add an FAQ answer.")

    def test_competing_pages_for_one_query(self):
        pairs = [pair("shared", "https://x.test/a/", 5, 90), pair("shared", "https://x.test/b/", 3, 60),
                 pair("shared", "https://x.test/c/", 1, 30), pair("shared", "https://x.test/d/", 0, 29),
                 pair("shared", "https://x.test/", 9, 500),                        # the homepage does not count
                 pair("lonely", "https://x.test/a/", 5, 90), pair("lonely", "https://x.test/b/", 1, 10)]
        rows = self.build([q("shared", 9, 5000, 7.0), q("lonely", 5, 100, 9.0)], pairs=pairs)
        cannibal = of_type(rows, "cannibal")
        self.assertEqual([c["keyword"] for c in cannibal], ["shared"])
        self.assertEqual(cannibal[0]["diagnosis"], "3 pages compete for this search.")
        self.assertEqual(cannibal[0]["page"], "https://x.test/a/")
        self.assertEqual(cannibal[0]["action"], "Merge into one page or make the pages target different intents.")


class PageRuleTests(unittest.TestCase):

    def update_rows(self, page, pairs=(), seasonal=None):
        rows = op.build(inputs(pages=[page], pairs=list(pairs), seasonal=seasonal or {},
                               next_month="October", min_impressions=300, target_ctr=0.08))
        return of_type(rows, "update_page")

    def test_the_five_diagnoses_in_order(self):
        url = "https://x.test/p/"
        split = self.update_rows(pg(url, 100, 13000, 7.5), [
            pair("topic", url, 100, 900), pair("topic", "https://x.test/q/", 5, 40), pair("topic", "https://x.test/r/", 2, 15)])
        self.assertEqual((split[0]["diagnosis"], split[0]["action"]),
                         ("Split with 2 other pages for the same keyword.", "Merge into one strong page."))
        self.assertEqual(split[0]["keyword"], "topic")
        title = self.update_rows(pg(url, 140, 13000, 7.5))                         # CTR 1.1% vs 3% expected
        self.assertEqual((title[0]["diagnosis"], title[0]["action"]),
                         ("Ranks well, title doesn't win the click.", "Rewrite title + meta, add a quick summary."))
        bottom = self.update_rows(pg(url, 200, 13000, 9.0))                        # 1.5%: under 2%, over 1%
        self.assertEqual(bottom[0]["diagnosis"], "Bottom of page 1.")
        seasonal = self.update_rows(pg(url, 500, 15000, 6.0), seasonal={url: (100, 150)})   # 3.3% vs 4%
        self.assertEqual((seasonal[0]["diagnosis"], seasonal[0]["action"]),
                         ("Seasonal.", "Update before October."))
        plain = self.update_rows(pg(url, 500, 15000, 6.0))
        self.assertEqual(plain[0]["diagnosis"], "Close to the top.")
        self.assertEqual(plain[0]["action"], "Answer the main question in the first lines, add FAQ.")

    def test_one_other_page_is_not_enough_to_call_it_a_split(self):
        url = "https://x.test/p/"
        rows = self.update_rows(pg(url, 140, 13000, 7.5), [pair("topic", url, 100, 900),
                                                          pair("topic", "https://x.test/q/", 5, 40)])
        self.assertEqual(rows[0]["diagnosis"], "Ranks well, title doesn't win the click.")

    def test_pages_outside_the_rule_are_ignored(self):
        url = "https://x.test/p/"
        for page in (pg(url, 100, 999, 7.0), pg(url, 100, 5000, 3.9), pg(url, 100, 5000, 10.5),
                     pg(url, 400, 5000, 7.0)):                                      # CTR 8% beats 3%
            self.assertEqual(self.update_rows(page), [], page)

    def decliners(self, rows, **kw):
        out = op.build(inputs(month_pages=rows, **kw))
        return of_type(out, "decliner")

    def both(self, url, **kw):
        return [month_row("Declining", url, **kw), month_row("YoY Declining", url)]

    def test_a_decliner_is_down_on_last_month_and_last_year_and_has_a_diagnosis(self):
        url = "https://x.test/old/"
        only_month = [month_row("Declining", url)]
        self.assertEqual(self.decliners(only_month), [])
        ranking = self.decliners(self.both(url, **{"Position diff": 1.4}))
        self.assertEqual((ranking[0]["diagnosis"], ranking[0]["action"]), ("Ranking fell.", "Update content."))
        self.assertEqual(ranking[0]["potential_gain_month"], 50)                    # the clicks lost
        demand = self.decliners(self.both(url, Impressions=900, **{"Impr prev": 1200}))
        self.assertTrue(demand[0]["diagnosis"].startswith("Fewer searches"))
        ctr = self.decliners(self.both(url, Clicks=60, **{"Clicks prev": 100, "Impressions": 1000, "Impr prev": 1000}))
        self.assertTrue(ctr[0]["diagnosis"].startswith("Search results changed"))
        generic = self.decliners(self.both(url, Clicks=90, **{"Clicks prev": 100, "Clicks diff": -10}))
        self.assertTrue(generic[0]["diagnosis"].startswith("Down on last month and last year"))

    def test_gone_resolved_and_invisible_pages_are_not_decliners(self):
        url = "https://x.test/old/"
        self.assertEqual(self.decliners(self.both(url), gone={url}), [])
        self.assertEqual(self.decliners(self.both(url), resolved=resolved.parse([["page", "/old/", "", ""]])), [])
        self.assertEqual(self.decliners(self.both(url, Impressions=0)), [])

    def test_check_page_is_an_80_percent_impression_drop_that_is_not_gone(self):
        url = "https://x.test/x/"
        drop = op.build(inputs(month_pages=[month_row("Declining", url, Impressions=150, **{"Impr prev": 1000})]))
        rows = of_type(drop, "check_page")
        self.assertEqual(rows[0]["diagnosis"], "Impressions down 85% on last month.")
        self.assertEqual(rows[0]["action"], "Check this page (index status, recent changes).")
        smaller = op.build(inputs(month_pages=[month_row("Declining", url, Impressions=300, **{"Impr prev": 1000})]))
        self.assertEqual(of_type(smaller, "check_page"), [])
        gone = op.build(inputs(month_pages=[month_row("Declining", url, Impressions=150, **{"Impr prev": 1000})],
                               gone={url}))
        self.assertEqual(of_type(gone, "check_page"), [])

    def test_aio_up_and_clicks_down(self):
        url = "https://x.test/a/"
        aio = [{"Page": url, "Type": "Growing", "AIO Impr": "300", "AIO Impr prev": "100", "AIO Impr diff": "200"},
               {"Page": "https://x.test/b/", "Type": "Growing", "AIO Impr": "30", "AIO Impr prev": "10", "AIO Impr diff": "20"}]
        rows = op.build(inputs(aio_pages=aio, month_pages=[month_row("Declining", url)]))
        found = of_type(rows, "aio")
        self.assertEqual([f["page"] for f in found], [url])                           # b's clicks did not fall
        self.assertEqual(found[0]["diagnosis"], "AI Overview impressions +200, clicks -50.")

    def test_seasonal_pages(self):
        seasonal = {"https://x.test/a/": (100, 160), "https://x.test/b/": (100, 120),
                    "https://x.test/c/": (10, 90), "https://x.test/d/": (100, 200)}
        rows = op.build(inputs(seasonal=seasonal, next_month="October", gone={"https://x.test/d/"}))
        found = of_type(rows, "seasonal")
        self.assertEqual([f["page"] for f in found], ["https://x.test/a/"])
        self.assertEqual((found[0]["diagnosis"], found[0]["action"]),
                         ("Grew +60% into October last year.", "Update before October."))
        self.assertEqual(found[0]["potential_gain_month"], 60)


class OrderAndTabTests(unittest.TestCase):

    def test_rows_are_sorted_by_type_then_gain_and_capped_per_type(self):
        queries = [q(f"query {i}", 10, 3000 + i * 300, 9.0) for i in range(6)]
        with mock.patch.object(op, "MAX_PER_TYPE", 4):
            rows = op.build(inputs(queries=queries, min_impressions=300, target_ctr=0.08))
        self.assertEqual(len(rows), 4)
        gains = [r[8] for r in rows]
        self.assertEqual(gains, sorted(gains, reverse=True))
        self.assertEqual(op.counts(rows), {"quick_win": 4})

    def test_types_come_in_the_documented_order(self):
        rows = op.build(inputs(queries=[q("weak", 20, 4000, 3.0), q("blue", 120, 5475, 9.4)],
                               seasonal={"https://x.test/a/": (100, 200)}, next_month="October"))
        self.assertEqual([r[0] for r in rows], ["quick_win", "ctr_fix", "seasonal"])

    def test_the_header_is_the_agreed_columns(self):
        self.assertEqual(op.HEADER, ["type", "keyword", "page", "impressions_90d", "clicks_90d", "position",
                                     "ctr", "expected_ctr", "potential_gain_month", "diagnosis", "action",
                                     "segment"])


class GoneRowsTests(unittest.TestCase):

    def test_a_page_with_impressions_last_period_and_none_now_is_gone(self):
        prev = {"https://x.test/a": (10, 100, 3.0), "https://x.test/b": (5, 50, 4.0), "https://x.test/c": (0, 0, 0.0)}
        cur = {"https://x.test/b": (4, 40, 4.0)}
        rows = reports.gone_rows("2026-09", cur, prev, {"https://x.test/y": (3, 30, 5.0)})
        self.assertEqual([(r[1], r[2]) for r in rows],
                         [("Gone", "https://x.test/a"), ("Gone", "https://x.test/y")])
        self.assertEqual(rows[0], ["2026-09", "Gone", "https://x.test/a", 10, 100, 0, 0, "/a"])
        self.assertEqual(rows[1][3:7], [0, 0, 3, 30])

    def test_a_page_with_zero_impressions_now_counts_as_gone_too(self):
        rows = reports.gone_rows("2026-09", {"https://x.test/a": (0, 0, 0.0)}, {"https://x.test/a": (3, 20, 2.0)})
        self.assertEqual(rows[0][2], "https://x.test/a")

    def test_biggest_clicks_first_and_capped(self):
        prev = {f"https://x.test/p{i}": (i, 100, 3.0) for i in range(1, 8)}
        rows = reports.gone_rows("2026-09", {}, prev, None, cap=3)
        self.assertEqual([r[3] for r in rows], [7, 6, 5])

    def test_weekly_rows_have_no_year_on_year(self):
        rows = reports.gone_rows("2026-W41", {}, {"https://x.test/a": (4, 10, 2.0)}, None)
        self.assertEqual(rows[0][5:7], ["", ""])

    def test_a_period_without_gone_pages_gets_a_marker_row(self):
        self.assertEqual(reports.gone_rows("2026-09", {"https://x.test/a": (1, 10, 2.0)},
                                           {"https://x.test/a": (1, 10, 2.0)}),
                         [["2026-09", "None", "", "", "", "", "", ""]])


class EngineIntegrationTests(ReportsTestCase):
    """run_reports with the Data sheet filled with invented rows."""

    def setUp(self):
        super().setUp()
        self.sheets.add_book("main-id")
        self.fill("Queries", QUERIES_HEADER, [["alpha query", "30", "2000", "1.5", "8.0"],
                                              ["beta query", "300", "1000", "30", "1.5"]])
        self.fill("Pages", PAGES_HEADER, [["https://www.alpha.test/en/article/a", "30", "2000", "1.5", "8.0"]])
        self.fill("Query×Page", PAIRS_HEADER, [["alpha query", "https://www.alpha.test/en/article/a",
                                                    "30", "2000", "1.5", "8.0"]])

    def fill(self, tab, header, rows):
        from test_history import Tab
        t = Tab(self.sheets.next_id, len(rows) + 1, len(header))
        self.sheets.next_id += 1
        t.data = [list(header)] + [list(r) for r in rows]
        self.sheets.books["main-id"][tab] = t

    def test_the_opportunities_tab_is_written_from_the_data_sheet(self):
        result = self.run_reports()
        self.assertEqual(self.tab(reports.OPPORTUNITIES)[0], op.HEADER)
        found = of_type(self.rows(reports.OPPORTUNITIES), "quick_win")
        self.assertEqual([f["keyword"] for f in found], ["alpha query"])
        self.assertEqual(found[0]["page"], "https://www.alpha.test/en/article/a")
        self.assertEqual(result.opportunities.get("quick_win"), 1)

    def test_the_tab_is_replaced_each_run_not_appended(self):
        self.run_reports()
        first = [list(r) for r in self.rows(reports.OPPORTUNITIES)]
        self.run_reports()
        self.assertEqual(self.rows(reports.OPPORTUNITIES), first)

    def test_brand_terms_and_thresholds_come_from_the_site(self):
        site = site_with(opportunities={"brand_terms": ["alpha"], "min_impressions_90d": 5000})
        self.run_reports(site)
        self.assertEqual([r for r in self.rows(reports.OPPORTUNITIES) if r[0] in ("quick_win", "ctr_fix", "question")], [])

    def test_a_data_sheet_that_cannot_be_read_skips_the_tab_without_failing(self):
        del self.sheets.books["main-id"]["Queries"]
        result = self.run_reports()
        self.assertEqual(result.opportunities, {})
        self.assertTrue(any("Opportunities skipped" in l for l in self.logs))

    def test_gone_pages_are_found_once_per_month(self):
        gone_url = "https://www.alpha.test/en/guide/b"
        from datetime import date
        self.records[:] = [r for r in self.records
                           if not (r["page"] == gone_url and r["date"] >= date(2026, 9, 1))]
        self.run_reports()
        rows = self.rows(reports.GONE)
        gone = [r for r in rows if r[1] == "Gone" and r[0] == "2026-09"]
        self.assertIn(gone_url, [r[2] for r in gone])
        row = [r for r in gone if r[2] == gone_url][0]
        self.assertEqual(row[7], "/en/guide/b")
        self.assertGreater(float(row[3]), 0)                                    # it had clicks in August
        count = len(rows)
        self.run_reports()
        self.assertEqual(len(self.rows(reports.GONE)), count)                   # not fetched again
        periods_done = {r[0] for r in self.rows(reports.GONE)}
        self.assertIn("2026-09", periods_done)
        self.assertIn("2026-08", periods_done)                                  # months without gone pages: a marker row

    def test_the_gone_pages_header(self):
        self.run_reports()
        self.assertEqual(self.tab(reports.GONE)[0], reports.GONE_HEADER)


class SiteSettingTests(unittest.TestCase):

    def parse(self, **extra):
        entry = {"name": "site-a", "site_url": "sc-domain:a.test", "sheet_id": "m", **extra}
        return sites.parse_sites_json(json.dumps([entry]))[0]

    def test_defaults(self):
        s = self.parse()
        self.assertEqual(s.opportunity_setting("min_impressions_90d"), 300)
        self.assertEqual(s.opportunity_setting("target_ctr_top3"), 0.08)
        self.assertEqual(s.opportunity_setting("brand_terms"), [])

    def test_settings(self):
        s = self.parse(opportunities={"brand_terms": ["Acme", " acme shop "], "min_impressions_90d": 500,
                                      "target_ctr_top3": 0.1})
        self.assertEqual(s.opportunity_setting("brand_terms"), ["Acme", "acme shop"])
        self.assertEqual(s.opportunity_setting("min_impressions_90d"), 500)
        self.assertEqual(s.opportunity_setting("target_ctr_top3"), 0.1)
        self.assertIn("Acme", s.sensitive_values())                       # brand words stay out of logs

    def test_bad_values_are_refused_without_echoing_them(self):
        for bad in ({"brand_terms": "secret"}, {"brand_terms": [1]}, {"min_impressions_90d": 0},
                    {"target_ctr_top3": 8}, {"target_ctr_top3": True}, {"secret_setting": 1}, "secret"):
            with self.assertRaises(sites.SitesConfigError) as cm:
                self.parse(opportunities=bad)
            self.assertNotIn("secret", str(cm.exception))


class WeeklyGoneSlideTests(unittest.TestCase):

    def deck(self, gone):
        rep = deck_fixtures.weekly_report()
        rep.gone = gone
        from pptx import Presentation
        import io
        data = decks.build_deck(rep, fx.Cfg(), deck_runner.TEMPLATE_DIR)
        return Presentation(io.BytesIO(data))

    def rows(self, n):
        return [dict(zip(reports.GONE_HEADER, ["2026-W41", "Gone", f"https://example.test/old/p{i}", 40 - i,
                                                300, "", "", f"/old/p{i}"])) for i in range(n)]

    def test_the_weekly_deck_gets_a_table_of_pages_that_disappeared(self):
        prs = self.deck(self.rows(12))
        self.assertEqual(len(prs.slides), 7)
        slide = prs.slides[-1]
        text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        self.assertIn("Pages that disappeared", text)
        table = [s for s in slide.shapes if s.has_table][0].table
        self.assertEqual(len(table.rows), 11)                                  # header + the top 10
        self.assertEqual(table.cell(1, 0).text_frame.text, "/old/p0")
        self.assertEqual(table.cell(1, 1).text_frame.text, "40")

    def test_no_gone_pages_means_no_extra_slide(self):
        self.assertEqual(len(self.deck([]).slides), 6)


if __name__ == "__main__":
    unittest.main()
