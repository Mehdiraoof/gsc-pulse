"""
PulseO: PowerPoint decks (v2.1)
===============================
Builds the monthly and weekly SEO report as a .pptx with python-pptx, entirely in
memory, from a designed template (templates/gsc_pulse_template.pptx, or the _fa one for
Persian). The look (fonts, colours, positions of titles, footer, rule, logo, slide number)
comes from the template's layouts; this code only fills their placeholders and draws
cards, tables and native charts inside the layouts' content areas.

Slides are described in one place (MONTHLY_SLIDES / WEEKLY_SLIDES) so a new slide is one
more function in a list.
"""

import copy
import io
import os
import re
import urllib.request
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import charts
import client_i18n
import deck_data as dd
import i18n

IN = 914400
FONT_FA = "Vazirmatn"

INK, SLATE, TINT = "1B1E3C", "4A4F6A", "F3F4F8"
AMBER, PERIWINKLE = "F2A541", "6C7AE0"
GREEN, RED, LINE, MUTED, CALLOUT = "1A9E6A", "D64545", "D9DBE5", "8A8FA6", "FFF6E8"

# layout name -> {role: placeholder idx}. The templates identify placeholders by idx.
PLACEHOLDERS = {
    "Title": {"site": 101, "title": 102, "subtitle": 103},
    "Summary": {"title": 103, "footer": 104, "body": 106, "next_steps": 107},
    "Alarms": {"title": 103, "footer": 104},
    "Scorecard": {"title": 103, "footer": 104},
    "Chart": {"title": 103, "footer": 104, "caption": 105},
    "Chart + Table": {"title": 103, "footer": 104},
    "Two Tables": {"title": 103, "footer": 104, "left_heading": 105, "right_heading": 106},
    "Notes": {"title": 103, "footer": 104, "body": 105},
}


class TemplateError(Exception):
    """The template does not have what the decks need. The message names the missing layout."""


def inch(v: float) -> int:
    return int(round(v * IN))


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


# --------------------------------------------------------------------------- #
# The deck being built
# --------------------------------------------------------------------------- #

@dataclass
class Ctx:
    prs: object
    layouts: dict
    lang: str
    T: dict
    label: str
    report: dd.Report
    report_title: str
    period_title: str
    accent: str = AMBER                    # clicks colour (config: accent_color)
    brand: str = PERIWINKLE                # impressions colour (config: brand_color)
    logo: bytes = None
    summary: tuple = None                  # (bullets, steps) from the AI, or None
    aggregated: str = ""
    slide_count: int = 0
    mark: bytes = None                     # the PulseO mark, taken from the template

    @property
    def rtl(self) -> bool:
        return i18n.is_rtl(self.lang)

    @property
    def font(self) -> str:
        return FONT_FA if self.rtl else ""

    @property
    def footer(self) -> str:
        return self.T["footer"].format(label=self.label, report=self.report_title,
                                       period=self.period_title)


def open_template(path: str):
    prs = Presentation(path)
    layouts = {l.name: l for l in prs.slide_layouts}
    for name in PLACEHOLDERS:
        if name not in layouts:
            raise TemplateError(f"the template has no slide layout named '{name}'")
    sld_ids = prs.slides._sldIdLst                      # the template ships demo slides: drop them
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)
    return prs, layouts


def template_mark(layouts) -> bytes:
    """The PulseO mark (a picture on the layouts), reused for the 'What it means' callout."""
    for shape in layouts["Scorecard"].shapes:
        if shape.shape_type == 13:                       # PICTURE
            return shape.image.blob
    return None


# --------------------------------------------------------------------------- #
# Text
# --------------------------------------------------------------------------- #

_ARABIC = i18n._ARABIC


def split_script(text: str, neutral_to_persian: bool = False) -> list:
    """
    [(piece, is_persian)]: Persian words and everything else, in order. Spaces, digits and
    punctuation (neutrals) go with the letters around them; where a neutral stretch sits
    between the two scripts it goes to the script of the paragraph (Latin in an English deck,
    Persian in a Persian one), so a space never ends up on the wrong side of a Persian word.
    """
    kinds = []
    for ch in text:
        if _ARABIC.match(ch) or ch in "\u200c\u200f":
            kinds.append(True)
        elif ch.isalpha():
            kinds.append(False)
        else:
            kinds.append(None)
    n = len(kinds)
    i = 0
    while i < n:
        if kinds[i] is not None:
            i += 1
            continue
        j = i
        while j < n and kinds[j] is None:
            j += 1
        before = kinds[i - 1] if i > 0 else None
        after = kinds[j] if j < n else None
        if before is not None and before == after:
            fill = before
        elif before is None and after is None:
            fill = neutral_to_persian
        elif before is None or after is None:
            fill = before if before is not None else after
            if before is None:
                fill = after
        else:
            fill = neutral_to_persian
        for k in range(i, j):
            kinds[k] = fill
        i = j
    pieces = []
    for ch, k in zip(text, kinds):
        if pieces and pieces[-1][1] == k:
            pieces[-1][0] += ch
        else:
            pieces.append([ch, k])
    return [(p, k) for p, k in pieces]


def _style_run(run, ctx: Ctx, size, bold, color, persian: bool, number: bool = False):
    """
    Persian runs (and every run of a Persian deck) get Vazirmatn and are marked right-to-left.
    A number inside Persian text is its own run in the same font but *not* marked right-to-left
    (language en-US, wrapped in LRM): marked right-to-left, a minus sign lands on the wrong
    side of the digits.
    """
    font = run.font
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color:
        font.color.rgb = rgb(color)
    if number:
        font.name = FONT_FA
        run._r.get_or_add_rPr().set("lang", "en-US")
    elif persian or ctx.rtl:
        font.name = FONT_FA
        r_pr = run._r.get_or_add_rPr()
        r_pr.set("lang", "fa-IR")
        etree.SubElement(r_pr, qn("a:cs")).set("typeface", FONT_FA)
        etree.SubElement(r_pr, qn("a:rtl")).set("val", "1")


def run_pieces(chunk: str, ctx: Ctx) -> list:
    """[(text, is_persian, is_number)] for one chunk of text."""
    if not ctx.rtl:
        return [(p, k, False) for p, k in split_script(chunk)]
    out, pos = [], 0
    for m in i18n._NUMBER.finditer(chunk):
        if m.start() > pos:
            out += [(p, k, False) for p, k in split_script(chunk[pos:m.start()], True)]
        out.append((i18n.LRM + m.group(0) + i18n.LRM, False, True))
        pos = m.end()
    if pos < len(chunk):
        out += [(p, k, False) for p, k in split_script(chunk[pos:], True)]
    return out

def write_paragraph(paragraph, text: str, ctx: Ctx, size=None, bold=None, color=None,
                    align=None, rtl=None, runs=None):
    """
    Fills a paragraph. `runs` is a list of (text, {size, bold, color}) for mixed formatting;
    otherwise `text` is one run. Every Persian run gets Vazirmatn (all runs in a Persian deck),
    and in a Persian deck every run is right-to-left, with numbers wrapped in LRM.
    """
    for r in list(paragraph.runs):
        r._r.getparent().remove(r._r)
    items = runs if runs else [(text, {})]
    for chunk, fmt in items:
        for piece, persian, number in run_pieces(chunk, ctx):
            run = paragraph.add_run()
            run.text = piece
            _style_run(run, ctx, fmt.get("size", size), fmt.get("bold", bold),
                       fmt.get("color", color), persian, number)
    p_pr = paragraph._p.get_or_add_pPr()
    is_rtl = ctx.rtl if rtl is None else rtl
    if is_rtl:
        p_pr.set("rtl", "1")
    if align:
        p_pr.set("algn", align)            # physical alignment: l, r or ctr
    elif is_rtl:
        p_pr.set("algn", "r")


def set_bullet(paragraph, kind: str, space_after: int = 10):
    """kind: 'bullet' or 'number'. The layouts have no bullets of their own."""
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        p_pr.remove(child)
    p_pr.set("marL", "342900")
    p_pr.set("indent", "-342900")
    aft = etree.SubElement(p_pr, qn("a:spcAft"))
    etree.SubElement(aft, qn("a:spcPts")).set("val", str(space_after * 100))
    if kind == "bullet":
        etree.SubElement(p_pr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(p_pr, qn("a:buChar")).set("char", "•")
    else:
        etree.SubElement(p_pr, qn("a:buFont")).set("typeface", "+mj-lt")
        etree.SubElement(p_pr, qn("a:buAutoNum")).set("type", "arabicPeriod")


def fill_placeholder(ph, paragraphs: list, ctx: Ctx, anchor=None, **fmt):
    """paragraphs: list of str or (str, {format}). First one reuses the placeholder's paragraph.
    The text starts exactly at the placeholder's edge (no inner margin), like the sample decks."""
    tf = ph.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Emu(0)
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, item in enumerate(paragraphs):
        text, own = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        write_paragraph(p, text, ctx, **{**fmt, **own})

def text_box(slide, x, y, w, h, paragraphs: list, ctx: Ctx, size=14, color=SLATE, bold=None,
             align=None, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Emu(inch(x)), Emu(inch(y)), Emu(inch(w)), Emu(inch(h)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(paragraphs):
        text, own = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        write_paragraph(p, text, ctx, **{"size": size, "color": color, "bold": bold,
                                          "align": align, **own})
    return box


def shape(slide, kind, x, y, w, h, fill: str, radius: float = None):
    """A flat shape with no outline, shadow or theme styling."""
    s = slide.shapes.add_shape(kind, Emu(inch(x)), Emu(inch(y)), Emu(inch(w)), Emu(inch(h)))
    style = s._element.find(qn("p:style"))
    if style is not None:
        s._element.remove(style)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    if radius is not None:
        s.adjustments[0] = radius
    return s


def card(slide, x, y, w, h, fill=TINT, radius=0.06):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius)


def mirror_x(ctx: Ctx, x: float, w: float, left: float = 0.6, width: float = 12.13) -> float:
    """In a right-to-left deck, a shape at x moves to the opposite side of the content area."""
    return left + width - (x - left) - w if ctx.rtl else x


# --------------------------------------------------------------------------- #
# Slides from layouts
# --------------------------------------------------------------------------- #

def new_slide(ctx: Ctx, layout_name: str, title: str = None, notes: str = None):
    layout = ctx.layouts[layout_name]
    slide = ctx.prs.slides.add_slide(layout)
    by_idx = {p.placeholder_format.idx: p for p in slide.placeholders}
    roles = {}
    for role, idx in PLACEHOLDERS[layout_name].items():
        if idx not in by_idx:
            raise TemplateError(f"the layout '{layout_name}' has no placeholder {idx} ({role})")
        roles[role] = by_idx[idx]
    used = {p.placeholder_format.idx for p in roles.values()}
    for p in list(slide.placeholders):                   # nothing is left empty
        if p.placeholder_format.idx not in used:
            drop_placeholder(p)
    for shp in layout.placeholders:                      # the slide number is not cloned: copy it
        if shp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(shp._element))
    ctx.slide_count += 1
    if title is not None:
        text = title if len(title) <= 78 else title[:77] + "…"
        fill_placeholder(roles["title"], [text], ctx, anchor=MSO_ANCHOR.MIDDLE,
                         **({"size": 26} if len(text) > 56 else {}))
    if "footer" in roles:
        fill_placeholder(roles["footer"], [ctx.footer], ctx)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide, roles


def drop_placeholder(ph):
    ph._element.getparent().remove(ph._element)


# --------------------------------------------------------------------------- #
# Tables
# --------------------------------------------------------------------------- #

def _borders(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        tc_pr.remove(child)
    for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tc_pr, qn(side))
        ln.set("w", "9525")
        etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", LINE)
    etree.SubElement(etree.SubElement(tc_pr, qn("a:solidFill")), qn("a:srgbClr")).set("val", fill)
    tc_pr.set("anchor", "ctr")
    for margin in ("marL", "marR"):
        tc_pr.set(margin, str(inch(0.1)))
    tc_pr.set("marT", str(inch(0.04)))
    tc_pr.set("marB", str(inch(0.04)))


@dataclass
class Cell:
    text: str
    kind: str = "text"            # text | query | num | change
    color: str = None
    bold: bool = None


def add_table(slide, ctx: Ctx, x, y, col_widths: list, header: list, rows: list,
              row_height: float = 0.42, size: int = 12, text_cols: tuple = ()):
    """header: list of str. rows: list of list of Cell. Returns the table graphic frame."""
    w = sum(col_widths)
    frame = slide.shapes.add_table(len(rows) + 1, len(col_widths), Emu(inch(x)), Emu(inch(y)),
                                   Emu(inch(w)), Emu(inch(row_height * (len(rows) + 1))))
    table = frame.table
    tbl_pr = table._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tbl_pr.set(attr, "0")
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"       # No Style, No Grid
    if ctx.rtl:
        tbl_pr.set("rtl", "1")                                      # first column on the right
    for j, cw in enumerate(col_widths):
        table.columns[j].width = Emu(inch(cw))
    for i in range(len(rows) + 1):
        table.rows[i].height = Emu(inch(row_height))

    def put(cell, spec: Cell, header_row: bool, force_right: bool = False):
        _borders(cell, TINT if header_row else "FFFFFF")
        persian = dd_has_arabic(spec.text)
        if spec.kind in ("query", "page"):
            align = "r" if (persian or ctx.rtl) else "l"
        elif spec.kind == "text":
            align = "r" if ctx.rtl else "l"
        else:                                           # numbers sit at the end of the row
            align = "l" if ctx.rtl else "r"
        if force_right:
            align = "r"
        p = cell.text_frame.paragraphs[0]
        write_paragraph(p, spec.text, ctx, size=size - 1 if header_row else size,
                        bold=True if header_row else spec.bold,
                        color=SLATE if header_row else (spec.color or INK), align=align,
                        rtl=True if (persian and spec.kind in ("query", "page")) else None)

    persian_cols = {j for row in rows for j, c in enumerate(row)
                    if c.kind in ("query", "page") and dd_has_arabic(c.text)}
    for j, h in enumerate(header):
        put(table.cell(0, j), Cell(h, "text" if j in text_cols else ("num" if j > 0 else "query")), True,
            force_right=j in persian_cols)
    for i, row in enumerate(rows, start=1):
        for j, spec in enumerate(row):
            put(table.cell(i, j), spec, False)
    return frame


def dd_has_arabic(text: str) -> bool:
    return i18n.has_arabic_script(text)


def sentence(slide, ctx: Ctx, text: str, x=0.6, y=2.4, w=5.9):
    """Empty data is a sentence, never an empty table."""
    return text_box(slide, x, y, w, 0.8, [text], ctx, size=18, color=MUTED)


# --------------------------------------------------------------------------- #
# Monthly / weekly slides
# --------------------------------------------------------------------------- #

def slide_title(ctx: Ctx):
    slide, ph = new_slide(ctx, "Title")
    fill_placeholder(ph["site"], [ctx.label if ctx.rtl else ctx.label.upper()], ctx)
    fill_placeholder(ph["title"], [ctx.report_title], ctx)
    fill_placeholder(ph["subtitle"], [ctx.period_title], ctx)
    if ctx.logo:
        try:
            slide.shapes.add_picture(io.BytesIO(ctx.logo), Emu(inch(10.8 if not ctx.rtl else 0.6)),
                                     Emu(inch(0.7)), height=Emu(inch(0.9)))
        except Exception:                                 # a bad logo never breaks the deck
            pass
    return slide


def slide_summary(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    slide, ph = new_slide(ctx, "Summary", title=dd.title_for_clicks(rep, T))
    if ctx.summary:
        bullets, steps = ctx.summary
        fill_placeholder(ph["body"], bullets, ctx, size=18)
        for p in ph["body"].text_frame.paragraphs:
            set_bullet(p, "bullet", 12)
        notes = None
    else:
        bullets, steps = [T["summary_placeholder"]], [T["summary_hint"]]
        fill_placeholder(ph["body"], bullets, ctx, size=22, color=MUTED)
        notes = f"{T['notes_heading']}\n\n{ctx.aggregated}"
    fill_placeholder(ph["next_steps"],
                     [(T["next_steps"], {"bold": True, "size": 18, "color": INK})] + steps, ctx)
    for p in ph["next_steps"].text_frame.paragraphs[1:]:
        set_bullet(p, "number", 10)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _alarm_cards(slide, ctx: Ctx, alarms: list):
    top, height, left, width = 1.5, 5.1, 0.6, 12.13
    n = len(alarms)
    gap = 0.12
    h = min(1.0, (height - gap * (n - 1)) / n)
    for i, a in enumerate(alarms):
        y = top + i * (h + gap)
        card(slide, left, y, width, h)
        high = a["severity"] == "high"
        pill_w, pill_h = 1.25, 0.38
        pill_x = mirror_x(ctx, left + 0.3, pill_w)
        pill = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, pill_x, y + (h - pill_h) / 2, pill_w, pill_h,
                     RED if high else AMBER, 0.5)
        tf = pill.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        write_paragraph(tf.paragraphs[0], ctx.T["sev_" + a["severity"]], ctx, size=12, bold=True,
                        color="FFFFFF" if high else INK, align="ctr")
        text_x = mirror_x(ctx, left + 1.85, width - 2.15)
        area = ctx.T.get("area_" + a["area"], a["area"])
        message = dd.alarm_text(a, ctx.lang)
        small = h < 0.72
        text_box(slide, text_x, y, width - 2.15, h,
                 [(area, {"bold": True, "color": INK, "size": 13 if small else 15}),
                  (message, {"size": 11.5 if small else 13})], ctx, anchor=MSO_ANCHOR.MIDDLE)


def slide_alarms(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    weekly = rep.kind == "weekly"
    alarms = rep.alarms[: dd.MAX_ALARMS]
    slide, ph = new_slide(ctx, "Alarms", title=dd.alarms_title(alarms, T, weekly))
    if not alarms:
        sentence(slide, ctx, T["alarms_none_week" if weekly else "alarms_none"], y=2.2, w=12.13)
    else:
        _alarm_cards(slide, ctx, alarms)
    return slide


def slide_scorecard(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    cards = dd.scorecard(rep, T)
    slide, ph = new_slide(ctx, "Scorecard", title=dd.title_for_clicks(rep, T))
    gap, left, total = 0.25, 0.6, 12.13
    w = (total - gap * (len(cards) - 1)) / max(len(cards), 1)
    for i, c in enumerate(cards):
        x = mirror_x(ctx, left + i * (w + gap), w)
        card(slide, x, 1.7, w, 3.0, TINT, 0.07)
        text_box(slide, x + 0.25, 1.95, w - 0.5, 0.4, [T[c.label_key]], ctx, size=14, color=SLATE)
        text_box(slide, x + 0.25, 2.55, w - 0.5, 0.9, [c.value], ctx, size=40, color=INK, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE)
        box = text_box(slide, x + 0.25, 3.7, w - 0.5, 0.9, [""], ctx, size=14)
        paragraphs = box.text_frame.paragraphs
        items = ([("note", T[c.note_key])] if c.note_key else []) + [("line", l) for l in c.lines]
        for j, (kind, item) in enumerate(items):
            para = paragraphs[0] if j == 0 else box.text_frame.add_paragraph()
            para.space_after = Pt(6)
            if kind == "note":
                write_paragraph(para, item, ctx, size=14, color=MUTED, bold=True)
            else:
                color = GREEN if item.good is True else (RED if item.good is False else SLATE)
                write_paragraph(para, "", ctx, runs=[(T[item.label_key] + "  ", {"color": MUTED, "size": 13}),
                                                     (item.text, {"color": color, "bold": True, "size": 14})])
    meaning = dd.what_it_means(rep, T)
    if meaning:
        card(slide, 0.6, 5.0, 12.13, 1.45, CALLOUT, 0.06)
        if ctx.mark:
            slide.shapes.add_picture(io.BytesIO(ctx.mark), Emu(inch(mirror_x(ctx, 0.9, 0.7))),
                                     Emu(inch(5.4)), height=Emu(inch(0.65)))
        box = text_box(slide, mirror_x(ctx, 1.9, 10.5), 5.1, 10.5, 1.25, [""], ctx, size=15,
                       anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(box.text_frame.paragraphs[0], "", ctx, size=15,
                        runs=[(T["what_it_means"] + "  ", {"bold": True, "color": INK}),
                              (meaning, {"color": SLATE})])
    return slide


def _trend_slide(ctx: Ctx, count: int, label_of):
    T, rep = ctx.T, ctx.report
    labels, clicks, impr = dd.trend_series(rep, count)
    unit = T["units_months" if rep.kind == "monthly" else "units_weeks"]
    peak = max(clicks) if clicks else 0
    now = clicks[-1] if clicks else 0
    if len(labels) < 2 or not peak:
        title = T["trend_title_plain"].format(n=len(labels), unit=unit)
    elif now >= peak:
        title = T["trend_title_high"].format(n=len(labels), unit=unit)
    elif now < peak * 0.9:
        title = T["trend_title_peak"].format(n=len(labels), unit=unit,
                                             peak=label_of(labels[clicks.index(peak)]),
                                             pct=f"{(1 - now / peak) * 100:.0f}%")
    else:
        title = T["trend_title_plain"].format(n=len(labels), unit=unit)
    slide, ph = new_slide(ctx, "Chart", title=title)
    charts.combo_chart(slide, (inch(0.6), inch(1.5), inch(12.13), inch(4.6)),
                       [label_of(l) for l in labels], (T["series_clicks"], clicks),
                       (T["series_impressions"], impr), (ctx.accent, ctx.brand),
                       (T["axis_clicks"], T["axis_impressions"]), ctx.font)
    if labels:
        hi_i, lo_i = clicks.index(max(clicks)), clicks.index(min(clicks))
        fill_placeholder(ph["caption"], [T["trend_caption"].format(
            hi_label=label_of(labels[hi_i]), hi=i18n.fmt_int(clicks[hi_i]),
            lo_label=label_of(labels[lo_i]), lo=i18n.fmt_int(clicks[lo_i]))], ctx)
    else:
        drop_placeholder(ph["caption"])
    return slide


def slide_trend_monthly(ctx: Ctx):
    return _trend_slide(ctx, 13, lambda l: i18n.short_month(l, ctx.lang))


def slide_trend_weekly(ctx: Ctx):
    return _trend_slide(ctx, 12, lambda l: "W" + l[6:])


def slide_segments(ctx: Ctx):
    """Only for sites that have page types: clicks by segment, bars plus a table."""
    T, rep = ctx.T, ctx.report
    segs = [s for s in rep.segments if rep.row(segment=s)]
    if len(segs) < 2:
        return None
    slide, ph = new_slide(ctx, "Chart + Table", title=T["segments_title"])
    label = lambda s: T["segment_all"] if s == "All" else s      # the whole site is never called "All"
    ordered = sorted(segs, key=lambda s: -(dd.num(rep.row(segment=s)["Clicks"]) or 0))
    chart_x = 0.6 if not ctx.rtl else 5.6
    charts.bar_chart(slide, (inch(chart_x), inch(1.5), inch(7.0), inch(5.1)), [label(s) for s in ordered],
                     T["series_clicks"], [dd.num(rep.row(segment=s)["Clicks"]) or 0 for s in ordered],
                     ctx.accent, ctx.font)
    rows = []
    for s in ordered:
        r = rep.row(segment=s)
        mom, yoy = dd.change_cell(r["Clicks MoM %"]), dd.change_cell(r["Clicks YoY %"])
        rows.append([Cell(label(s), "query"), Cell(i18n.fmt_int(dd.num(r["Clicks"]) or 0), "num"),
                     Cell(mom[0] + ("%" if mom[0] != "—" else ""), "change",
                          GREEN if mom[1] else (RED if mom[1] is False else None), True),
                     Cell(yoy[0] + ("%" if yoy[0] != "—" else ""), "change",
                          GREEN if yoy[1] else (RED if yoy[1] is False else None), True),
                     Cell(i18n.fmt_int(dd.num(r["Impressions"]) or 0), "num")])
    add_table(slide, ctx, 7.85 if not ctx.rtl else 0.6, 1.7,
              [1.45, 0.85, 0.85, 0.85, 0.88],
              [T["col_segment"], T["col_clicks"], T["col_mom"], T["col_yoy"], T["col_impr"]], rows,
              row_height=0.5)
    return slide


def _two_tables(ctx: Ctx, title: str, notes: str = None):
    slide, ph = new_slide(ctx, "Two Tables", title=title, notes=notes)
    return slide, ph


def _put_two(slide, ph, ctx: Ctx, first, second):
    """first / second: (heading, header cols, col widths, rows, empty sentence). In a Persian deck
    the first table is on the right, because that is where reading starts."""
    sides = [("left_heading", 0.6), ("right_heading", 6.83)]
    if ctx.rtl:
        sides.reverse()
    for (role, x), spec in zip(sides, (first, second)):
        heading, header, widths, rows, empty = spec
        fill_placeholder(ph[role], [heading], ctx)
        if rows:
            add_table(slide, ctx, x, 1.95, widths, header, rows)
        else:
            sentence(slide, ctx, empty, x=x, y=2.3)


def _kw_rows(rows: list, key: str, cols: list) -> list:
    out = []
    for r in rows[: dd.TOP_ROWS]:
        label = dd.short_path(r[key]) if key == "Page" else r[key]
        cells = [Cell(label, "query")]
        for col in cols:
            if col.endswith("diff"):
                text, good = dd.change_cell(r[col])
                cells.append(Cell(text, "change", GREEN if good else (RED if good is False else None), True))
            else:
                cells.append(Cell(i18n.fmt_int(dd.num(r[col]) or 0), "num"))
        out.append(cells)
    return out


def slide_lost_new(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    lost, new = dd.of_type(rep.keywords, "Lost"), dd.of_type(rep.keywords, "New")
    key = "kw_title_one" if len(lost) == 1 else "kw_title"
    slide, ph = _two_tables(ctx, T[key].format(lost=len(lost), new=len(new)))
    prev_m = i18n.month_name(rep.prev_period, ctx.lang, short=True)
    cur_m = i18n.month_name(rep.period, ctx.lang, short=True)
    _put_two(slide, ph, ctx,
             (T["h_lost"], [T["col_query"], T["col_impr_month"].format(m=prev_m)], [4.3, 1.6],
              _kw_rows(lost, "Query", ["Impr prev"]), T["no_lost"]),
             (T["h_new"], [T["col_query"], T["col_impr_month"].format(m=cur_m)], [4.3, 1.6],
              _kw_rows(new, "Query", ["Impressions"]), T["no_new"]))
    return slide


def _sum_abs(rows: list, col: str) -> int:
    return int(sum(abs(dd.num(r[col]) or 0) for r in rows))


def slide_keyword_movers(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    up, down = dd.of_type(rep.keywords, "Growing"), dd.of_type(rep.keywords, "Declining")
    title = (T["movers_title"].format(up=i18n.fmt_int(_sum_abs(up, "Clicks diff")),
                                      down=i18n.fmt_int(_sum_abs(down, "Clicks diff")))
             if (up or down) else T["movers_title_plain"])
    slide, ph = _two_tables(ctx, title)
    cols = [T["col_query"], T["col_clicks"], T["col_change"]]
    _put_two(slide, ph, ctx,
             (T["h_kw_up"], cols, [3.5, 1.15, 1.25], _kw_rows(up, "Query", ["Clicks", "Clicks diff"]), T["no_kw_up"]),
             (T["h_kw_down"], cols, [3.5, 1.15, 1.25], _kw_rows(down, "Query", ["Clicks", "Clicks diff"]), T["no_kw_down"]))
    return slide


def _page_notes(rows: list) -> str:
    return "\n".join(r["Page"] for r in rows[: dd.TOP_ROWS])


def slide_page_movers(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    up, down = dd.of_type(rep.pages, "Growing"), dd.of_type(rep.pages, "Declining")
    if down and (dd.num(down[0]["Clicks diff"]) or 0) < 0:
        path = dd.short_path(down[0]["Page"], 30)
        n = i18n.fmt_int(abs(dd.num(down[0]["Clicks diff"])))
        title = (T["pages_title_loss_generic"].format(n=n) if (dd_has_arabic(path) and not ctx.rtl)
                 else T["pages_title_loss"].format(page=path, n=n))
    else:
        title = T["pages_title_plain"]
    notes = "Full URLs\n\nGrowing:\n" + _page_notes(up) + "\n\nDeclining:\n" + _page_notes(down)
    slide, ph = _two_tables(ctx, title, notes)
    cols = [T["col_page"], T["col_clicks"], T["col_change"]]
    _put_two(slide, ph, ctx,
             (T["h_pg_up"], cols, [3.5, 1.15, 1.25], _kw_rows(up, "Page", ["Clicks", "Clicks diff"]), T["no_pg_up"]),
             (T["h_pg_down"], cols, [3.5, 1.15, 1.25], _kw_rows(down, "Page", ["Clicks", "Clicks diff"]), T["no_pg_down"]))
    return slide


def slide_yoy_pages(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    up, down = dd.of_type(rep.pages, "YoY Growing"), dd.of_type(rep.pages, "YoY Declining")
    if not (up or down):
        slide, ph = new_slide(ctx, "Notes", title=T["yoy_title_none"])
        fill_placeholder(ph["body"], [T["no_yoy"]], ctx, size=20, color=MUTED)
        return slide
    title = T["yoy_title"].format(up=i18n.fmt_int(_sum_abs(up, "Clicks diff")),
                                  down=i18n.fmt_int(_sum_abs(down, "Clicks diff")))
    notes = "Full URLs\n\nGrowing vs last year:\n" + _page_notes(up) + "\n\nDeclining vs last year:\n" + _page_notes(down)
    slide, ph = _two_tables(ctx, title, notes)
    period = i18n.month_title(rep.year_period, ctx.lang)
    cols = [T["col_page"], T["col_clicks"], T["col_change"]]
    _put_two(slide, ph, ctx,
             (T["h_yoy_up"].format(period=period), cols, [3.5, 1.15, 1.25],
              _kw_rows(up, "Page", ["Clicks", "Clicks diff"]), T["no_yoy_up"]),
             (T["h_yoy_down"].format(period=period), cols, [3.5, 1.15, 1.25],
              _kw_rows(down, "Page", ["Clicks", "Clicks diff"]), T["no_yoy_down"]))
    return slide


def slide_aio(ctx: Ctx):
    """One or two slides: the AI search trend and pages, or how to import the data."""
    T, rep = ctx.T, ctx.report
    month = i18n.month_name(rep.period, ctx.lang)
    cur = rep.row()
    aio = dd.num(cur["AIO Impressions"]) if cur else None
    if aio is None:
        slide, ph = new_slide(ctx, "Chart", title=T["aio_title_none"].format(month=month))
        card(slide, 0.6, 1.6, 12.13, 3.6)
        ym = rep.period
        steps = [T["aio_step1"], T["aio_step2"], T["aio_step3"].format(ym=ym)]
        box = text_box(slide, mirror_x(ctx, 1.1, 11.1), 1.9, 11.1, 3.0,
                       [(T["aio_howto_heading"], {"bold": True, "color": INK, "size": 18})] + steps,
                       ctx, size=16)
        for p in box.text_frame.paragraphs[1:]:
            set_bullet(p, "number", 10)
        fill_placeholder(ph["caption"], [T["aio_manual"]], ctx)
        return slide
    labels, values = dd.aio_series(rep)
    slide, ph = new_slide(ctx, "Chart", title=T["aio_title"].format(n=i18n.fmt_int(aio), month=month))
    charts.column_chart(slide, (inch(0.6), inch(1.5), inch(12.13), inch(4.6)),
                        [i18n.short_month(l, ctx.lang) for l in labels], T["sc_aio"], values,
                        ctx.brand, T["axis_impressions"], ctx.font)
    fill_placeholder(ph["caption"], [T["aio_chart_caption"]], ctx)

    up, down = dd.of_type(rep.aio_pages, "Growing")[:5], dd.of_type(rep.aio_pages, "Declining")[:5]
    slide2, ph2 = _two_tables(ctx, T["aio_pages_title"],
                              "Full URLs\n\n" + _page_notes(up) + "\n\n" + _page_notes(down))
    cols = [T["col_page"], T["col_aio_impr"], T["col_change"]]
    _put_two(slide2, ph2, ctx,
             (T["h_aio_up"], cols, [3.5, 1.15, 1.25], _kw_rows(up, "Page", ["AIO Impr", "AIO Impr diff"]), T["no_aio_up"]),
             (T["h_aio_down"], cols, [3.5, 1.15, 1.25], _kw_rows(down, "Page", ["AIO Impr", "AIO Impr diff"]), T["no_aio_down"]))
    return slide2


def slide_about(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    slide, ph = new_slide(ctx, "Notes", title=T["about_title"])
    unit_free = {"AI Overviews", "Alarms", "Page types", "All and segments"}
    if ctx.lang == "en" and rep.notes:
        bullets = [T["about_source"]] + [f"{topic}: {dd.first_sentence(text)}" for topic, text in rep.notes
                                         if topic not in unit_free and text]
        bullets.append(T["about_types"].format(names=", ".join(rep.page_types)) if rep.page_types
                       else T["about_types_none"])
    else:
        bullets = [T["about_source"], T["about_complete"], T["about_changes"],
                   T["about_keywords"].format(n=rep.lost_new_floor),
                   T["about_types"].format(names=", ".join(rep.page_types)) if rep.page_types
                   else T["about_types_none"]]
    bullets.append(T["about_history"])
    fill_placeholder(ph["body"], bullets[:8], ctx, size=16)
    for p in ph["body"].text_frame.paragraphs:
        set_bullet(p, "bullet", 12)
    return slide


def slide_weekly_movers(ctx: Ctx):
    T, rep = ctx.T, ctx.report
    up, down = dd.of_type(rep.movers, "Growing"), dd.of_type(rep.movers, "Declining")
    slide, ph = _two_tables(ctx, T["wk_movers_title"])
    cols = [T["col_query"], T["col_clicks"], T["col_change"]]
    _put_two(slide, ph, ctx,
             (T["wk_h_up"], cols, [3.5, 1.15, 1.25], _kw_rows(up, "Query", ["Clicks", "Clicks diff"]), T["no_wk_movers"]),
             (T["wk_h_down"], cols, [3.5, 1.15, 1.25], _kw_rows(down, "Query", ["Clicks", "Clicks diff"]), T["no_wk_movers"]))
    return slide


def slide_weekly_gone(ctx: Ctx):
    """Weekly only: pages that had impressions last week and none this week (top 10 by clicks
    last week), so redirects can be checked. Skipped when no page disappeared."""
    rep = ctx.report
    gone = rep.gone[:10]
    if not gone:
        return None
    CT = client_i18n.table(ctx.lang)
    slide, ph = new_slide(ctx, "Notes", title=CT["c_wk_gone"])
    drop_placeholder(ph["body"])
    sentence(slide, ctx, CT["c_wk_gone_intro"], x=0.6, y=1.5, w=12.13)
    rows = [[Cell(dd.short_path(r["Page"], 60), "page"),
             Cell(i18n.fmt_int(dd.num(r["Clicks prev"]) or 0), "num"),
             Cell(i18n.fmt_int(dd.num(r["Impr prev"]) or 0), "num")] for r in gone]
    add_table(slide, ctx, mirror_x(ctx, 0.6, 12.13), 2.3, [8.1, 2.0, 2.03],
              [ctx.T["col_page"], CT["c_col_clicks_before"], CT["c_col_impr_before"]], rows,
              row_height=0.4, size=12)
    return slide




MONTHLY_SLIDES = [slide_title, slide_summary, slide_alarms, slide_scorecard, slide_trend_monthly,
                  slide_segments, slide_lost_new, slide_keyword_movers, slide_page_movers,
                  slide_yoy_pages, slide_aio, slide_about]
WEEKLY_SLIDES = [slide_title, slide_summary, slide_alarms, slide_scorecard, slide_trend_weekly,
                 slide_weekly_movers, slide_weekly_gone]


# --------------------------------------------------------------------------- #
# Building a deck
# --------------------------------------------------------------------------- #

def period_title(rep: dd.Report, lang: str) -> str:
    return dd.period_title(rep, lang)


def build_deck(report: dd.Report, cfg, template_dir: str, summary=None, logo: bytes = None) -> bytes:
    """The .pptx as bytes. cfg is the site's DeckConfig."""
    lang = cfg.language
    path = os.path.join(template_dir, "gsc_pulse_template_fa.pptx" if lang == "fa" else "gsc_pulse_template.pptx")
    prs, layouts = open_template(path)
    T = i18n.table(lang)
    ctx = Ctx(prs=prs, layouts=layouts, lang=lang, T=T, label=cfg.site_label, report=report,
              report_title=T["monthly_report" if report.kind == "monthly" else "weekly_report"],
              period_title=period_title(report, lang),
              accent=(cfg.accent_color or AMBER), brand=(cfg.brand_color or PERIWINKLE),
              logo=logo, summary=summary, mark=template_mark(layouts),
              aggregated=dd.aggregated_text(report, cfg.site_label))
    for build in (MONTHLY_SLIDES if report.kind == "monthly" else WEEKLY_SLIDES):
        build(ctx)
    core = prs.core_properties
    core.title = f"{ctx.report_title} · {cfg.site_label} · {ctx.period_title}"
    core.author = "PulseO"
    core.subject = ctx.report_title
    core.comments = ""
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def fetch_logo(url: str, limit: int = 2_000_000):
    """The logo bytes, or None. Failure to download must never fail the deck."""
    if not url or not url.startswith("https://"):
        return None
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "PulseO"}),
                                    timeout=10) as response:
            data = response.read(limit + 1)
        return data if 0 < len(data) <= limit else None
    except Exception:                                    # noqa: BLE001
        return None
