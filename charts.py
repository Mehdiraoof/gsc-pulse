"""
PulseO: native PowerPoint charts
================================
Charts are real, editable PowerPoint charts (python-pptx chart API), not images, so
they stay sharp and PowerPoint draws any Persian text itself. python-pptx cannot make
a combo chart with a second value axis, so the trend chart is a column chart whose
second series is moved into a line chart on a secondary axis by editing the chart XML.
"""


from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

INK, SLATE, LINE = "1B1E3C", "4A4F6A", "D9DBE5"
AMBER, PERIWINKLE = "F2A541", "6C7AE0"

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _style_axis_text(axis, size: int, color: str, font: str):
    axis.tick_labels.font.size = Pt(size)
    axis.tick_labels.font.color.rgb = _rgb(color)
    if font:
        axis.tick_labels.font.name = font
    axis.format.line.color.rgb = _rgb(LINE)


def _axis_title(axis, text: str, font: str):
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(12)
    run.font.bold = False
    run.font.color.rgb = _rgb(SLATE)
    if font:
        run.font.name = font


def _base_chart(slide, kind, x, y, w, h, categories, series, font: str):
    data = CategoryChartData()
    data.categories = categories
    for name, values, number_format in series:
        data.add_series(name, values, number_format)
    frame = slide.shapes.add_chart(kind, Emu(x), Emu(y), Emu(w), Emu(h), data)
    chart = frame.chart
    chart.has_title = False
    chart.font.size = Pt(12)
    chart.font.color.rgb = _rgb(SLATE)
    if font:
        chart.font.name = font
    return frame, chart


def _legend(chart, font: str):
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(13)
    chart.legend.font.color.rgb = _rgb(SLATE)
    if font:
        chart.legend.font.name = font


def combo_chart(slide, box, categories, bars, line, colors, titles, font=""):
    """
    Columns on the left axis plus a line with markers on the right axis.
    box: (x, y, w, h) in EMU.  bars / line: (series name, values).
    colors: (bar hex, line hex).  titles: (left axis title, right axis title).
    """
    frame, chart = _base_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, *box, categories,
                               [(bars[0], bars[1], "#,##0"), (line[0], line[1], "#,##0")], font)
    _legend(chart, font)
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.vary_by_categories = False
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = _rgb(colors[0])
    plot.series[0].invert_if_negative = False

    value_axis, category_axis = chart.value_axis, chart.category_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = _rgb(LINE)
    value_axis.tick_labels.number_format = "#,##0"
    value_axis.tick_labels.number_format_is_linked = False
    _style_axis_text(value_axis, 12, SLATE, font)
    _style_axis_text(category_axis, 12, SLATE, font)
    _axis_title(value_axis, titles[0], font)

    _move_second_series_to_secondary_line(chart, colors[1], titles[1], font)
    return frame


def _el(tag, **attrs):
    prefix, name = tag.split(":")
    ns = C_NS if prefix == "c" else A_NS
    el = etree.Element(f"{{{ns}}}{name}")
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _move_second_series_to_secondary_line(chart, line_color: str, title: str, font: str):
    plot_area = chart._chartSpace.chart.plotArea
    bar = plot_area.find(qn("c:barChart"))
    sers = bar.findall(qn("c:ser"))
    ser = sers[1]
    bar.remove(ser)
    for tag in ("c:invertIfNegative", "c:spPr"):          # column-only settings
        for child in ser.findall(qn(tag)):
            ser.remove(child)

    # series formatting for a line with round markers
    sp_pr = etree.fromstring(
        f'<c:spPr xmlns:c="{C_NS}" xmlns:a="{A_NS}"><a:ln w="38100" cap="rnd">'
        f'<a:solidFill><a:srgbClr val="{line_color.lstrip("#").upper()}"/></a:solidFill>'
        f'<a:round/></a:ln></c:spPr>')
    marker = etree.fromstring(
        f'<c:marker xmlns:c="{C_NS}" xmlns:a="{A_NS}"><c:symbol val="circle"/><c:size val="8"/>'
        f'<c:spPr><a:solidFill><a:srgbClr val="{line_color.lstrip("#").upper()}"/></a:solidFill>'
        f'<a:ln w="9525"><a:solidFill><a:srgbClr val="{line_color.lstrip("#").upper()}"/>'
        f'</a:solidFill></a:ln></c:spPr></c:marker>')
    # a line series is: idx, order, tx, spPr, marker, dPt, dLbls, trendline, errBars, cat, val, smooth
    tx = ser.find(qn("c:tx"))
    anchor = tx if tx is not None else ser.find(qn("c:order"))
    anchor.addnext(sp_pr)
    sp_pr.addnext(marker)
    ser.append(_el("c:smooth", val=0))

    sec_cat, sec_val = "50010", "50020"
    line_chart = etree.fromstring(
        f'<c:lineChart xmlns:c="{C_NS}"><c:grouping val="standard"/><c:varyColors val="0"/>'
        f'</c:lineChart>')
    line_chart.append(ser)
    line_chart.append(_el("c:marker", val=1))
    line_chart.append(_el("c:axId", val=sec_cat))
    line_chart.append(_el("c:axId", val=sec_val))
    bar.addnext(line_chart)

    last_axis = plot_area.findall(qn("c:valAx"))[-1]
    cat_ax = etree.fromstring(
        f'<c:catAx xmlns:c="{C_NS}"><c:axId val="{sec_cat}"/><c:scaling><c:orientation val="minMax"/>'
        f'</c:scaling><c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/>'
        f'<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/><c:crossAx val="{sec_val}"/>'
        f'<c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/>'
        f'<c:noMultiLvlLbl val="0"/></c:catAx>')
    font_xml = (f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/>' if font else "")
    val_ax = etree.fromstring(
        f'<c:valAx xmlns:c="{C_NS}" xmlns:a="{A_NS}"><c:axId val="{sec_val}"/><c:scaling>'
        f'<c:orientation val="minMax"/></c:scaling><c:delete val="0"/><c:axPos val="r"/>'
        f'<c:title><c:tx><c:rich><a:bodyPr rot="-5400000" vert="horz"/><a:lstStyle/><a:p><a:pPr>'
        f'<a:defRPr sz="1200" b="0"/></a:pPr><a:r><a:rPr lang="en-US" sz="1200" b="0">'
        f'<a:solidFill><a:srgbClr val="{SLATE}"/></a:solidFill>{font_xml}</a:rPr>'
        f'<a:t>{_escape(title)}</a:t></a:r></a:p></c:rich></c:tx><c:overlay val="0"/></c:title>'
        f'<c:numFmt formatCode="#,##0" sourceLinked="0"/><c:majorTickMark val="out"/>'
        f'<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
        f'<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="{LINE}"/></a:solidFill></a:ln></c:spPr>'
        f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="1200">'
        f'<a:solidFill><a:srgbClr val="{SLATE}"/></a:solidFill>{font_xml}</a:defRPr></a:pPr>'
        f'<a:endParaRPr lang="en-US"/></a:p></c:txPr>'
        f'<c:crossAx val="{sec_cat}"/><c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>')
    last_axis.addnext(cat_ax)
    cat_ax.addnext(val_ax)
    etree.cleanup_namespaces(chart._chartSpace)          # no repeated xmlns declarations


def _escape(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def column_chart(slide, box, categories, series_name, values, color, title="", font=""):
    """A single-series column chart (AIO impressions per month, weekly trends)."""
    frame, chart = _base_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, *box, categories,
                               [(series_name, values, "#,##0")], font)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = _rgb(color)
    plot.series[0].invert_if_negative = False
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb(LINE)
    va.tick_labels.number_format = "#,##0"
    va.tick_labels.number_format_is_linked = False
    _style_axis_text(va, 12, SLATE, font)
    _style_axis_text(chart.category_axis, 12, SLATE, font)
    if title:
        _axis_title(va, title, font)
    return frame


def bar_chart(slide, box, categories, series_name, values, color, font="", title=""):
    """Horizontal bars (clicks by segment), the biggest on top."""
    frame, chart = _base_chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, *box, categories,
                               [(series_name, values, "#,##0")], font)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.number_format = "#,##0"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(12)
    plot.data_labels.font.color.rgb = _rgb(INK)
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = _rgb(color)
    plot.series[0].invert_if_negative = False
    ca, va = chart.category_axis, chart.value_axis
    ca.reverse_order = True                      # first category on top
    va.has_major_gridlines = False
    va.visible = False
    _style_axis_text(ca, 12, SLATE, font)
    if title:
        _chart_title(chart, title, font)
    return frame


def _chart_title(chart, text: str, font: str):
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(14)
    run.font.bold = False
    run.font.color.rgb = _rgb(INK)
    if font:
        run.font.name = font
    chart.chart_title.include_in_layout = False


def line_chart(slide, box, categories, series_name, values, color, title="", font=""):
    """One line with round markers (clicks per month)."""
    frame, chart = _base_chart(slide, XL_CHART_TYPE.LINE_MARKERS, *box, categories,
                               [(series_name, values, "#,##0")], font)
    chart.has_legend = False
    series = chart.plots[0].series[0]
    series.smooth = False
    series.format.line.color.rgb = _rgb(color)
    series.format.line.width = Pt(3)
    series.marker.style = XL_MARKER_STYLE.CIRCLE
    series.marker.size = 8
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb = _rgb(color)
    series.marker.format.line.color.rgb = _rgb(color)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb(LINE)
    va.tick_labels.number_format = "#,##0"
    va.tick_labels.number_format_is_linked = False
    va.minimum_scale = 0
    _style_axis_text(va, 12, SLATE, font)
    _style_axis_text(chart.category_axis, 12, SLATE, font)
    if title:
        _chart_title(chart, title, font)
    return frame
