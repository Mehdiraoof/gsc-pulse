"""
Tests for the deck builder: structure, layouts, native charts, languages, empty data.
No network, no Google: invented data from deck_fixtures.
"""

import io
import os
import re
import shutil
import tempfile
import unittest
from datetime import date

from pptx import Presentation
from pptx.oxml.ns import qn

import deck_data as dd
import deck_fixtures as fx
import decks
import i18n
import sites

TEMPLATES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")
def _png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


PNG = _png()

MONTHLY_LAYOUTS = ["Title", "Summary", "Alarms", "Scorecard", "Chart", "Two Tables", "Two Tables",
                   "Two Tables", "Two Tables", "Chart", "Notes"]
WEEKLY_LAYOUTS = ["Title", "Summary", "Alarms", "Scorecard", "Chart", "Two Tables"]


def cfg(lang="en", **kw):
    base = dict(enabled=True, language=lang, site_label="Example", recipients=["x@example.test"])
    base.update(kw)
    return sites.DeckConfig(**base)


def build(report, lang="en", **kw):
    summary, logo = kw.pop("summary", None), kw.pop("logo", None)
    data = decks.build_deck(report, cfg(lang, **kw), TEMPLATES, summary, logo)
    return Presentation(io.BytesIO(data))


def texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out.extend(c.text_frame.text for c in row.cells)
    return out


def all_text(slide) -> str:
    return "\n".join(texts(slide))


def tables(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]


def charts_of(slide):
    return [sh.chart for sh in slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart]


class MonthlyStructureTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.prs = build(fx.monthly_report())

    def test_slides_follow_the_plan_and_use_template_layouts(self):
        self.assertEqual([s.slide_layout.name for s in self.prs.slides], MONTHLY_LAYOUTS)
        known = {l.name for l in self.prs.slide_layouts}
        self.assertTrue({s.slide_layout.name for s in self.prs.slides} <= known)

    def test_the_templates_demo_slides_are_gone(self):
        self.assertEqual(len(self.prs.slides), 11)
        for s in self.prs.slides:
            self.assertNotIn("Layout:", all_text(s))

    def test_title_slide(self):
        t = texts(self.prs.slides[0])
        self.assertIn("EXAMPLE", t)
        self.assertIn("Monthly SEO Report", t)
        self.assertIn("September 2026", t)

    def test_footer_on_every_content_slide_and_a_slide_number_field(self):
        for slide in list(self.prs.slides)[1:]:
            self.assertIn("Example · Monthly SEO Report · September 2026", texts(slide))
            self.assertTrue(slide._element.findall(".//" + qn("a:fld")), "slide number missing")

    def test_action_titles_are_generated_from_the_data(self):
        titles = [texts(s)[0] for s in list(self.prs.slides)[1:]]
        self.assertEqual(titles[0], "Clicks down 32% this month, 89% vs last year")        # summary
        self.assertEqual(titles[2], "Clicks down 32% this month, 89% vs last year")        # scorecard
        self.assertTrue(any(t == "12 keywords lost, 8 new" for t in titles))
        self.assertTrue(any(t.startswith("Last 13 months: clicks peaked in Jan 26") for t in titles))
        self.assertTrue(any(t.startswith("4 alarms this month") for t in titles))

    def test_summary_slide_is_the_placeholder_with_the_data_in_its_notes(self):
        slide = self.prs.slides[1]
        self.assertIn("Summary not generated automatically", all_text(slide))
        notes = slide.notes_slide.notes_text_frame.text
        self.assertIn("Clicks: 325", notes)
        self.assertIn("Lost keywords", notes)
        self.assertNotIn("example.test", notes.replace("/guide/", ""))     # paths, not URLs

    def test_summary_slide_with_an_ai_summary(self):
        prs = build(fx.monthly_report(), summary=(["Clicks fell 32%.", "Impressions fell too."],
                                                  ["Check indexing.", "Review the top pages."]))
        slide = prs.slides[1]
        body = all_text(slide)
        self.assertIn("Clicks fell 32%.", body)
        self.assertIn("Check indexing.", body)
        self.assertNotIn("Summary not generated automatically", body)
        self.assertEqual(slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else "", "")

    def test_alarm_cards_show_severity_and_area(self):
        text = all_text(self.prs.slides[2])
        for word in ("HIGH", "MEDIUM", "Traffic", "Page loss", "Lost keyword"):
            self.assertIn(word, text)
        self.assertIn("and 4 more pages lost 50%+ of their clicks", text)
        self.assertNotIn("https://", text)                                    # URLs are shortened

    def test_scorecard_numbers_and_colours(self):
        slide = self.prs.slides[3]
        text = all_text(slide)
        for expected in ("325", "69,600", "0.47%", "3.7", "−32.4%", "+0.3 worse", "−3.0 better",
                         "Not imported", "What it means"):
            self.assertIn(expected, text)
        colours = {}
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text in ("+0.3 worse", "−3.0 better", "−32.4%"):
                            colours[r.text] = str(r.font.color.rgb)
        self.assertEqual(colours, {"+0.3 worse": "D64545", "−3.0 better": "1A9E6A",
                                   "−32.4%": "D64545"})

    def test_trend_chart_is_native_with_two_series_on_two_axes(self):
        chart = charts_of(self.prs.slides[4])[0]
        plots = list(chart.plots)
        self.assertEqual(len(plots), 2)
        bars, line = plots
        self.assertEqual(bars.series[0].name, "Clicks")
        self.assertEqual(list(bars.series[0].values), [float(c) for c in fx.CLICKS])
        self.assertEqual(line.series[0].name, "Impressions")
        self.assertEqual(list(line.series[0].values), [float(c) for c in fx.IMPR])
        self.assertEqual(list(bars.categories)[0], "Sep 25")
        self.assertEqual(len(list(bars.categories)), 13)
        xml = chart._chartSpace.xml
        self.assertIn("<c:lineChart", xml)
        self.assertEqual(len(re.findall(r"<c:valAx[ >]", xml)), 2)             # a second value axis
        self.assertIn('crosses val="max"', xml)
        self.assertEqual(str(bars.series[0].format.fill.fore_color.rgb), "F2A541")
        self.assertIn("6C7AE0", xml)

    def test_config_colours_only_change_the_chart_series(self):
        prs = build(fx.monthly_report(), brand_color="#1F3A5F", accent_color="#E07A2E")
        chart = charts_of(prs.slides[4])[0]
        self.assertEqual(str(list(chart.plots)[0].series[0].format.fill.fore_color.rgb), "E07A2E")
        self.assertIn("1F3A5F", chart._chartSpace.xml)
        self.assertNotIn("6C7AE0", chart._chartSpace.xml)

    def test_keyword_and_page_tables(self):
        lost_new = self.prs.slides[5]
        self.assertEqual(len(tables(lost_new)), 2)
        lost = tables(lost_new)[0].table
        self.assertEqual(len(lost.rows), 11)                                  # header + top 10
        self.assertEqual([c.text for c in lost.rows[0].cells], ["Query", "Impr. Aug"])
        new = tables(lost_new)[1].table
        self.assertEqual(len(new.rows), 9)                                    # only 8 new keywords exist
        movers = tables(self.prs.slides[6])
        self.assertEqual([c.text for c in movers[0].table.rows[0].cells], ["Query", "Clicks", "Change"])
        change = movers[0].table.cell(1, 2)
        self.assertEqual(change.text, "+11")
        self.assertEqual(str(change.text_frame.paragraphs[0].runs[0].font.color.rgb), "1A9E6A")
        self.assertEqual(movers[1].table.cell(1, 2).text, "−14")

    def test_table_cells_use_at_least_11pt(self):
        for slide in self.prs.slides:
            for t in tables(slide):
                for row in t.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                self.assertGreaterEqual(r.font.size.pt, 11)

    def test_pages_show_paths_and_notes_hold_the_full_urls(self):
        slide = self.prs.slides[7]
        pages = tables(slide)[0].table
        self.assertTrue(pages.cell(1, 0).text.startswith("/guide/"))
        self.assertNotIn("example.test", pages.cell(1, 0).text)
        notes = slide.notes_slide.notes_text_frame.text
        self.assertIn("https://example.test/guide/", notes)

    def test_year_over_year_slide(self):
        slide = self.prs.slides[8]
        text = all_text(slide)
        self.assertIn("Growing vs September 2025", text)
        self.assertEqual(len(tables(slide)), 2)

    def test_aio_slide_without_import_explains_the_steps(self):
        text = all_text(self.prs.slides[10 - 1])
        self.assertIn("AI search: no data imported for September", text)
        self.assertIn("AIO 2026-09 Pages", text)
        self.assertIn("AIO 2026-09 Dates", text)

    def test_about_slide(self):
        text = all_text(self.prs.slides[10])
        self.assertIn("About this report", text)
        self.assertIn("PulseO", text)
        self.assertIn("Anonymized queries", text)

    def test_no_slide_shows_internal_names_or_regexes(self):
        for s in self.prs.slides:
            text = all_text(s)
            self.assertNotRegex(text, r"[A-Za-z0-9_-]{25,}")                  # no sheet ids
            self.assertNotIn("SITES_JSON", text)

    def test_document_properties(self):
        core = self.prs.core_properties
        self.assertEqual(core.author, "PulseO")
        self.assertIn("Example", core.title)


class SegmentsAndAioTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.prs = build(fx.monthly_report(segments=True, aio=True), accent_color="#E07A2E")

    def test_more_slides_when_there_are_segments_and_aio(self):
        self.assertEqual([s.slide_layout.name for s in self.prs.slides],
                         ["Title", "Summary", "Alarms", "Scorecard", "Chart", "Chart + Table",
                          "Two Tables", "Two Tables", "Two Tables", "Two Tables", "Chart",
                          "Two Tables", "Notes"])

    def test_segments_slide_has_a_bar_chart_and_a_table(self):
        slide = self.prs.slides[5]
        chart = charts_of(slide)[0]
        plot = list(chart.plots)[0]
        self.assertEqual(sorted(plot.categories), sorted(["Whole site", "English", "Article", "Other"]))
        self.assertEqual(plot.categories[0], "Whole site")                     # the biggest first
        rows = tables(slide)[0].table
        self.assertEqual([c.text for c in rows.rows[0].cells], ["Segment", "Clicks", "MoM", "YoY", "Impr."])
        self.assertEqual(len(rows.rows), 5)

    def test_aio_trend_chart_and_pages(self):
        trend = self.prs.slides[10]
        self.assertIn("AI search: 3,100 impressions in September", all_text(trend))
        chart = charts_of(trend)[0]
        series = list(chart.plots)[0].series[0]
        self.assertEqual(list(series.values), [5200.0, 4800.0, 4400.0, 3100.0])
        pages = self.prs.slides[11]
        self.assertEqual(len(tables(pages)), 2)
        self.assertEqual(len(tables(pages)[0].table.rows), 6)                  # top 5 growing + header

    def test_scorecard_shows_aio(self):
        text = all_text(self.prs.slides[3])
        self.assertIn("3,100", text)
        self.assertNotIn("Not imported", text)
        self.assertIn("−29.5%", text)


class EmptyDataTests(unittest.TestCase):

    def setUp(self):
        rep = fx.monthly_report()
        rep.keywords, rep.pages, rep.alarms, rep.aio_pages = [], [], [], []
        self.prs = build(rep)

    def test_no_alarms_is_a_sentence(self):
        slide = self.prs.slides[2]
        self.assertEqual(texts(slide)[0], "No major changes this month")
        self.assertIn("No major changes this month.", all_text(slide))

    def test_empty_tables_become_sentences_never_empty_tables(self):
        text = {i: all_text(s) for i, s in enumerate(self.prs.slides)}
        self.assertIn("No lost keywords this month", text[5])
        self.assertIn("No new keywords this month", text[5])
        self.assertIn("No growing keywords this month", text[6])
        self.assertIn("No declining pages this month", text[7])
        for slide in list(self.prs.slides)[5:9]:
            self.assertEqual(tables(slide), [])
        self.assertIn("year-over-year comparison is not possible yet", text[8])

    def test_titles_fall_back_to_plain_wording(self):
        titles = [texts(s)[0] for s in self.prs.slides]
        self.assertIn("Growing and declining pages", titles)
        self.assertIn("Growing and declining keywords", titles)


class WeeklyTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.prs = build(fx.weekly_report())

    def test_structure(self):
        self.assertEqual([s.slide_layout.name for s in self.prs.slides], WEEKLY_LAYOUTS)
        t = texts(self.prs.slides[0])
        self.assertIn("Weekly SEO Report", t)
        self.assertIn("Week 41 · 5 Oct – 11 Oct 2026", t)

    def test_alarms_are_week_on_week_drops(self):
        text = all_text(self.prs.slides[2])
        self.assertIn("Clicks fell 28.9% vs the previous week (90 to 64).", text)
        self.assertIn("Impressions fell 29.1% vs the previous week", text)

    def test_scorecard_uses_wow(self):
        text = all_text(self.prs.slides[3])
        self.assertIn("WoW", text)
        self.assertNotIn("MoM", text)

    def test_trend_is_twelve_weeks(self):
        chart = charts_of(self.prs.slides[4])[0]
        cats = list(list(chart.plots)[0].categories)
        self.assertEqual(len(cats), 12)
        self.assertEqual((cats[0], cats[-1]), ("W30", "W41"))

    def test_movers_slide(self):
        tabs_ = tables(self.prs.slides[5])
        self.assertEqual(len(tabs_), 2)
        self.assertEqual(tabs_[0].table.cell(1, 0).text, "blue running shoes")
        self.assertEqual(tabs_[1].table.cell(1, 2).text, "−8")

    def test_no_big_drop_means_no_major_changes_sentence(self):
        rep = fx.weekly_report()
        rep.alarms = []
        prs = build(rep)
        self.assertIn("No major changes this week.", all_text(prs.slides[2]))


class LogoTests(unittest.TestCase):

    def test_a_logo_is_embedded_on_the_title_slide(self):
        prs = build(fx.monthly_report(), logo=PNG)
        pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
        self.assertEqual(len(pics), 1)

    def test_a_broken_logo_never_fails_the_deck(self):
        prs = build(fx.monthly_report(), logo=b"not an image")
        self.assertEqual(len(prs.slides), 11)

    def test_logo_download_failures_return_none(self):
        self.assertIsNone(decks.fetch_logo("http://insecure.test/logo.png"))
        self.assertIsNone(decks.fetch_logo(""))
        self.assertIsNone(decks.fetch_logo("https://invalid.invalid/logo.png"))


class TemplateTests(unittest.TestCase):

    def test_a_missing_layout_is_a_clear_error_naming_it(self):
        tmp = tempfile.mkdtemp()
        try:
            prs = Presentation(os.path.join(TEMPLATES, "gsc_pulse_template.pptx"))
            for layout in prs.slide_layouts:
                if layout.name == "Scorecard":
                    layout._element.find(qn("p:cSld")).set("name", "Renamed")
            prs.save(os.path.join(tmp, "gsc_pulse_template.pptx"))
            with self.assertRaises(decks.TemplateError) as cm:
                decks.build_deck(fx.monthly_report(), cfg(), tmp)
            self.assertIn("Scorecard", str(cm.exception))
        finally:
            shutil.rmtree(tmp)

    def test_both_templates_have_every_layout(self):
        for name in ("gsc_pulse_template.pptx", "gsc_pulse_template_fa.pptx"):
            prs, layouts = decks.open_template(os.path.join(TEMPLATES, name))
            self.assertTrue(set(decks.PLACEHOLDERS) <= set(layouts), name)
            self.assertEqual(len(prs.slides), 0)

    def test_the_mark_comes_from_the_template(self):
        prs, layouts = decks.open_template(os.path.join(TEMPLATES, "gsc_pulse_template.pptx"))
        self.assertTrue(decks.template_mark(layouts))


class PersianTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.prs = build(fx.monthly_report(), lang="fa", site_label="نمونه")

    def runs(self, slide):
        for sh in slide.shapes:
            frames = []
            if sh.has_text_frame:
                frames.append(sh.text_frame)
            if getattr(sh, "has_table", False) and sh.has_table:
                frames += [c.text_frame for row in sh.table.rows for c in row.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs:
                        yield p, r

    def test_every_text_is_right_to_left_in_vazirmatn(self):
        for slide in self.prs.slides:
            for p, r in self.runs(slide):
                is_number = r.text.startswith("‎") and r.text.endswith("‎")
                p_pr = p._p.find(qn("a:pPr"))
                self.assertIsNotNone(p_pr, r.text)
                self.assertEqual(p_pr.get("rtl"), "1", r.text)
                self.assertEqual(r.font.name, "Vazirmatn", r.text)
                r_pr = r._r.find(qn("a:rPr"))
                if is_number:
                    self.assertIsNone(r_pr.find(qn("a:rtl")), r.text)        # numbers stay LTR
                elif r.text.strip():
                    self.assertIsNotNone(r_pr.find(qn("a:rtl")), r.text)
                    self.assertEqual(r_pr.find(qn("a:cs")).get("typeface"), "Vazirmatn")

    def test_numbers_signs_and_percent_are_wrapped_in_lrm(self):
        found = [r.text for slide in self.prs.slides for _, r in self.runs(slide)
                 if re.fullmatch(r"‎[+\-−]?\d[\d,.]*%?‎", r.text)]
        self.assertIn("‎−32.4%‎", found)
        self.assertIn("‎325‎", found)

    def test_persian_labels_and_title(self):
        T = i18n.table("fa")
        scorecard = all_text(self.prs.slides[3])
        for key in ("sc_clicks", "sc_impressions", "sc_ctr", "sc_position", "not_imported"):
            self.assertIn(T[key], scorecard)
        title_slide = all_text(self.prs.slides[0]).replace("\u200e", "")
        self.assertIn(T["monthly_report"], title_slide)
        self.assertIn(T["jalali_months"][5] + " 1405", title_slide)

    def test_tables_flow_right_to_left(self):
        for t in tables(self.prs.slides[5]):
            self.assertEqual(t.table._tbl.tblPr.get("rtl"), "1")

    def test_the_first_table_is_on_the_right(self):
        lost, new = tables(self.prs.slides[5])
        self.assertGreater(lost.left, new.left)
        self.assertIn("کلمات از دست رفته", texts(self.prs.slides[5]))

    def test_uses_the_persian_template(self):
        self.assertEqual(self.prs.slides[0].slide_layout.name, "Title")

    def test_alarm_sentences_are_translated(self):
        text = all_text(self.prs.slides[2])
        self.assertIn("ترافیک", text)
        self.assertIn("از دست داد", text)


class EnglishDeckWithPersianKeywordsTests(unittest.TestCase):

    def test_persian_runs_use_vazirmatn_and_english_runs_do_not(self):
        prs = build(fx.monthly_report())
        persian_fonts, english_fonts = set(), set()
        for slide in prs.slides:
            for t in tables(slide):
                for row in t.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                (persian_fonts if i18n.has_arabic_script(r.text) else english_fonts).add(r.font.name)
        self.assertEqual(persian_fonts, {"Vazirmatn"})
        self.assertNotIn("Vazirmatn", english_fonts)

    def test_persian_cells_are_right_aligned_and_right_to_left(self):
        prs = build(fx.monthly_report())
        cell = tables(prs.slides[5])[0].table.cell(1, 0)
        p_pr = cell.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
        self.assertEqual((p_pr.get("algn"), p_pr.get("rtl")), ("r", "1"))

    def test_mixed_sentences_keep_their_spaces_on_the_right_side(self):
        pieces = decks.split_script("The page /guide/آموزش/ lost 59.9%")
        flat = "".join(p for p, _ in pieces)
        self.assertEqual(flat, "The page /guide/آموزش/ lost 59.9%")
        persian = [p for p, k in pieces if k]
        self.assertEqual(len(persian), 1)
        self.assertFalse(persian[0].startswith(" "))
        self.assertFalse(persian[0].endswith(" "))

    def test_split_script_in_a_persian_paragraph_keeps_boundary_spaces_with_persian(self):
        pieces = decks.split_script("سلام AI دنیا", True)
        self.assertEqual("".join(p for p, _ in pieces), "سلام AI دنیا")
        self.assertTrue(pieces[0][1])


class TitleAndScorecardRuleTests(unittest.TestCase):
    T = i18n.table("en")

    def report(self, **row):
        rep = fx.monthly_report()
        rep.rows[(rep.period, "All")].update(row)
        return rep

    def test_clicks_title_rules(self):
        t = lambda **row: dd.title_for_clicks(self.report(**row), self.T)           # noqa: E731
        self.assertEqual(t(**{"Clicks MoM %": -32.4, "Clicks YoY %": -90.6}),
                         "Clicks down 32% this month, 91% vs last year")
        self.assertEqual(t(**{"Clicks MoM %": 12.2, "Clicks YoY %": 40.0}),
                         "Clicks up 12% this month, 40% vs last year")
        self.assertEqual(t(**{"Clicks MoM %": -12.2, "Clicks YoY %": 40.0}),
                         "Clicks down 12% this month, up 40% vs last year")
        self.assertEqual(t(**{"Clicks MoM %": 8.0, "Clicks YoY %": ""}), "Clicks up 8% this month")
        self.assertEqual(t(**{"Clicks MoM %": 0.2, "Clicks YoY %": ""}), "Clicks are steady this month")
        self.assertEqual(t(**{"Clicks MoM %": "", "Clicks YoY %": ""}), "Clicks: 325")

    def test_position_change_lower_is_better(self):
        cards = dd.scorecard(self.report(), self.T)
        pos = [c for c in cards if c.label_key == "sc_position"][0]
        self.assertEqual([(l.text, l.good) for l in pos.lines], [("+0.3 worse", False), ("−3.0 better", True)])

    def test_ctr_changes_are_in_points(self):
        cards = dd.scorecard(self.report(), self.T)
        ctr = [c for c in cards if c.label_key == "sc_ctr"][0]
        self.assertEqual(ctr.value, "0.47%")
        self.assertEqual([l.text for l in ctr.lines], ["−0.01 pt", "−1.33 pt"])

    def test_missing_comparisons_are_dashes(self):
        rep = fx.monthly_report()
        rep.year_period = "2020-09"
        cards = dd.scorecard(rep, self.T)
        self.assertEqual([l.text for l in cards[3].lines][1], "—")

    def test_what_it_means_rules(self):
        T = self.T
        self.assertIn("image block", dd.what_it_means(fx.monthly_report(), T))
        both = self.report(**{"Clicks MoM %": -30, "Clicks YoY %": -5, "Impr MoM %": -30})
        both.rows[(both.year_period, "All")]["Avg. Position"] = 2.0       # no better ranking last year
        self.assertIn("Google is showing the site less often", dd.what_it_means(both, T))
        flat = self.report(**{"Clicks MoM %": 1, "Clicks YoY %": 1, "Impr MoM %": 1})
        self.assertEqual(dd.what_it_means(flat, T), "")

    def test_alarm_titles(self):
        T = self.T
        alarm = lambda s: {"severity": s, "area": "Page loss", "message": "m", "impact": 1}    # noqa: E731
        self.assertEqual(dd.alarms_title([], T), "No major changes this month")
        self.assertEqual(dd.alarms_title([alarm("medium")], T), "1 alarm this month, all medium")
        self.assertEqual(dd.alarms_title([alarm("medium")] * 4, T), "4 alarms this month, all medium")
        self.assertEqual(dd.alarms_title([alarm("high"), alarm("medium"), alarm("medium")], T),
                         "3 alarms this month, 1 high, 2 medium")

    def test_short_paths(self):
        self.assertEqual(dd.short_path("https://x.test/a/b/?utm=1"), "/a/b/?utm=1")
        long = "https://x.test/" + "%D8%A7" * 60
        self.assertTrue(dd.short_path(long).endswith("…"))
        self.assertEqual(len(dd.short_path(long)), 34)
        self.assertEqual(dd.short_path("https://x.test/%D8%A7%D9%84"), "/ال")
        self.assertEqual(dd.shorten_urls("The page https://x.test/guide/ lost 5%."), "The page /guide/ lost 5%.")


class I18nTests(unittest.TestCase):

    def test_both_languages_have_the_same_keys(self):
        self.assertEqual(set(i18n.EN), set(i18n.FA))

    def test_every_format_placeholder_matches_between_languages(self):
        for key, en in i18n.EN.items():
            if isinstance(en, str):
                fa = i18n.FA[key]
                self.assertEqual(set(re.findall(r"\{(\w+)\}", en)) - {"when"},
                                 set(re.findall(r"\{(\w+)\}", fa)) - {"when"}, key)

    def test_months_and_jalali(self):
        self.assertEqual(i18n.month_title("2026-09", "en"), "September 2026")
        self.assertEqual(i18n.short_month("2026-09", "en"), "Sep 26")
        self.assertEqual(i18n.gregorian_to_jalali(2026, 9, 15)[:2], (1405, 6))
        self.assertIn("(شهریور 1405)", i18n.month_title("2026-09", "fa"))
        self.assertEqual(i18n.fmt_signed(-14), "−14")
        self.assertEqual(i18n.fmt_signed(11), "+11")
        self.assertEqual(i18n.fmt_signed(-0.01, 2), "−0.01")
        self.assertEqual(i18n.fmt_signed(0), "0")

    def test_bidi_only_changes_persian_text(self):
        self.assertEqual(i18n.bidi("Clicks -5%", "en"), "Clicks -5%")
        self.assertEqual(i18n.bidi("5%", "fa"), "‎5%‎")

    def test_alarm_message_translation(self):
        T = i18n.table("fa")
        out = dd.alarm_text({"severity": "high", "area": "Segment clicks", "impact": 1,
                             "message": "Clicks for All fell by 32.4% compared with 2026-08 (481 to 325)."}, "fa")
        self.assertIn("32.4%", out)
        self.assertIn("کمتر شد", out)
        unknown = dd.alarm_text({"severity": "low", "area": "New URL pattern", "impact": 0,
                                 "message": "Some other sentence."}, "fa")
        self.assertEqual(unknown, "Some other sentence.")
        self.assertIn("{metric}", T["msg_clicks_up"])




class PersianNumberRunTests(unittest.TestCase):

    def ctx(self):
        return decks.Ctx(prs=None, layouts={}, lang="fa", T=i18n.table("fa"), label="x",
                         report=None, report_title="", period_title="")

    def test_a_period_like_2026_08_is_not_read_as_a_negative_number(self):
        pieces = decks.run_pieces("\u0646\u0633\u0628\u062a \u0628\u0647 2026-08 \u0628\u0648\u062f", self.ctx())
        numbers = [p for p, _, n in pieces if n]
        self.assertEqual(numbers, ["\u200e2026\u200e", "\u200e08\u200e"])
        self.assertIn("-", "".join(p for p, _, n in pieces if not n))

    def test_signs_stay_with_their_number_after_a_space_or_bracket(self):
        pieces = decks.run_pieces("(\u22121.33) +0.3 -5", self.ctx())
        self.assertEqual([p for p, _, n in pieces if n],
                         ["\u200e\u22121.33\u200e", "\u200e+0.3\u200e", "\u200e-5\u200e"])


class AlarmWordingTests(unittest.TestCase):
    """What a person reads: month names, 'vs the previous week', segment labels, no 'All'."""

    def text(self, message, lang="en", area="Segment clicks"):
        return dd.alarm_text({"severity": "medium", "area": area, "message": message, "impact": 1}, lang)

    def test_the_whole_site_is_not_named(self):
        self.assertEqual(self.text("Clicks for All fell by 32.4% compared with 2026-08 (481 to 325)."),
                         "Clicks fell 32.4% vs August (481 to 325).")
        self.assertEqual(self.text("Clicks for All rose by 22.6% compared with 2025-07 (2,377 to 2,915)."),
                         "Clicks rose 22.6% vs July (2,377 to 2,915).")

    def test_a_segment_is_named_by_its_label(self):
        self.assertEqual(self.text("Clicks for English fell by 40.0% compared with 2026-12 (100 to 60)."),
                         "Clicks for English fell 40.0% vs December (100 to 60).")

    def test_other_sentences(self):
        self.assertEqual(
            self.text('The keyword "blue shoes" brought 26 clicks in 2026-08 and is not in Search '
                      "Console's data in 2026-09.", area="Lost keyword"),
            "The keyword “blue shoes” brought 26 clicks in August and is not in Search "
            "Console's data in September.")
        self.assertEqual(
            self.text("The page https://x.test/guide/ lost 59.9% of its clicks compared with 2026-08 "
                      "(137 to 55).", area="Page loss"),
            "The page /guide/ lost 59.9% of its clicks vs August (137 to 55).")
        self.assertEqual(
            self.text("AI Overview impressions fell by 51.6% compared with 2026-08 (6,200 to 3,000).",
                      area="AI Overviews"),
            "AI Overview impressions fell 51.6% vs August (6,200 to 3,000).")
        self.assertEqual(self.text("and 7 more pages lost 50%+ of their clicks", area="Page loss"),
                         "and 7 more pages lost 50%+ of their clicks")
        self.assertEqual(self.text("and 1 more page lost 50%+ of its clicks", area="Page loss"),
                         "and 1 more page lost 50%+ of its clicks")

    def test_an_unknown_sentence_still_loses_its_raw_periods_but_not_its_urls(self):
        out = self.text("Something changed between 2026-08 and 2026-09 on https://x.test/2026-08/ page.",
                        area="Other area")
        self.assertEqual(out, "Something changed between August and September on /2026-08/ page.")
        self.assertEqual(self.text("Week 2026-W41 was odd.", area="Other area"), "Week week 41 was odd.")

    def test_the_import_problem_keeps_the_tab_names_it_asks_for(self):
        msg = 'The tab "AIO Sept Pages" was not imported: its name must look like "AIO 2026-09 Pages".'
        self.assertEqual(self.text(msg, area="AIO import"), msg)

    def test_weekly_alarms(self):
        base = {"severity": "medium", "area": "Traffic", "weekly": True, "pct": 28.9, "a": 90, "b": 64,
                "impact": 26, "message": ""}
        whole = dict(base, metric="clicks", segment="All")
        seg = dict(base, metric="impressions", segment="English")
        self.assertEqual(dd.alarm_text(whole, "en"), "Clicks fell 28.9% vs the previous week (90 to 64).")
        self.assertEqual(dd.alarm_text(seg, "en"),
                         "Impressions for English fell 28.9% vs the previous week (90 to 64).")

    def test_persian_wording(self):
        fa = i18n.table("fa")
        out = self.text("Clicks for All fell by 32.4% compared with 2026-08 (481 to 325).", "fa")
        self.assertIn(fa["months"][7], out)                                  # August
        self.assertNotIn("2026-08", out)
        self.assertNotIn("All", out)
        self.assertIn(fa["clicks_all"], out)
        seg = self.text("Clicks for English rose by 5.0% compared with 2026-08 (10 to 11).", "fa")
        self.assertIn(fa["clicks_seg"].format(seg="English"), seg)

    def test_human_period_names(self):
        self.assertEqual(dd.human_period("2026-08", "en"), "August")
        self.assertEqual(dd.human_period("2026-W41", "en"), "week 41")
        self.assertEqual(dd.human_period("2026-13", "en"), "2026-13")
        self.assertEqual(dd.human_period("hello", "en"), "hello")


RAW_MONTH = re.compile(r"\b\d{4}-\d{2}\b")
RAW_WEEK = re.compile(r"\b\d{4}-W\d{2}\b")
ALL_AS_NAME = re.compile(r"\bAll\b")
TAB_INSTRUCTION = re.compile(r"AIO \d{4}-\d{2} (?:Pages|Dates)")        # the import steps name the tabs


class NoInternalNamesTests(unittest.TestCase):
    """Nothing a human reads may contain a raw YYYY-MM / YYYY-Www key or 'All' as a segment name."""

    def human_texts(self, prs):
        for slide in prs.slides:
            yield all_text(slide)
            for chart in charts_of(slide):
                yield " ".join(str(c) for c in list(chart.plots)[0].categories)
                for plot in chart.plots:
                    for s in plot.series:
                        yield s.name
            if slide.has_notes_slide:
                yield slide.notes_slide.notes_text_frame.text

    def check(self, label, text):
        text = TAB_INSTRUCTION.sub("", text.replace("‎", ""))
        self.assertIsNone(RAW_MONTH.search(text), f"{label}: raw month key in: {text[:200]}")
        self.assertIsNone(RAW_WEEK.search(text), f"{label}: raw week key in: {text[:200]}")
        self.assertIsNone(ALL_AS_NAME.search(text), f"{label}: 'All' as a name in: {text[:200]}")

    def test_every_deck_in_every_variant(self):
        cases = {
            "monthly": lambda: fx.monthly_report(),
            "monthly+segments+aio": lambda: fx.monthly_report(segments=True, aio=True),
            "weekly": lambda: fx.weekly_report(),
        }
        for lang in ("en", "fa"):
            for name, make in cases.items():
                prs = build(make(), lang=lang)
                for text in self.human_texts(prs):
                    self.check(f"{name}/{lang}", text)

    def test_weekly_alarms_for_a_segment(self):
        rep = fx.weekly_report()
        rep.rows[(rep.period, "English")] = dict(rep.rows[(rep.period, "All")], Segment="English")
        rep.segments = ["All", "English"]
        rep.alarms = dd.weekly_alarms(rep)
        text = " ".join(dd.alarm_text(a, "en") for a in rep.alarms)
        self.assertIn("Clicks for English fell", text)
        self.assertIn("Clicks fell", text)
        self.check("weekly segments", text)

    def test_email_subject_and_body(self):
        import deck_runner
        for lang in ("en", "fa"):
            for make in (fx.monthly_report, fx.weekly_report):
                rep = make()
                T = i18n.table(lang)
                title = T["monthly_report" if rep.kind == "monthly" else "weekly_report"]
                subject, body = deck_runner.email_text(rep, cfg(lang), lang, title, decks.period_title(rep, lang))
                self.check(f"email/{lang}", subject)
                self.check(f"email/{lang}", body)
                self.assertIn("PulseO", subject)

    def test_the_ai_summary_input_and_the_notes_use_names_too(self):
        text = dd.aggregated_text(fx.monthly_report(), "Example")
        self.check("aggregated", text)
        self.assertIn("September 2026", text)
        self.assertIn("vs August", text)
        weekly = dd.aggregated_text(fx.weekly_report(), "Example")
        self.check("aggregated weekly", weekly)
        self.assertIn("Week 41", weekly)

    def test_the_whole_site_row_of_the_segments_slide_is_called_whole_site(self):
        prs = build(fx.monthly_report(segments=True))
        slide = [s for s in prs.slides if s.slide_layout.name == "Chart + Table"][0]
        self.assertIn("Whole site", all_text(slide))
        self.assertNotRegex(all_text(slide), ALL_AS_NAME)


if __name__ == "__main__":
    unittest.main()
