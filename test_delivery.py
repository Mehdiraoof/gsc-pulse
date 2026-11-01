"""
Tests for deck delivery: the email, SMTP handling, the AI summary, when decks are sent
(State tab), rebuild, privacy, and the weekly movers tab the weekly deck reads.
"""

import contextlib
import glob
import io
import json
import os
import re
import smtplib
import sys
import unittest
from datetime import date, timedelta
from unittest import mock

from pptx import Presentation

import deck_runner
import gsc_pulse
import mailer
import reports
import sites
import summary
from test_history import FakeSheets, patch_time
from test_reports import (RID, FakeSA, ReportsTestCase, build_records, site_with)

ROOT = os.path.dirname(os.path.abspath(__file__))
USER = "owner@example.test"


def deck(**kw):
    base = dict(enabled=True, language="en", site_label="Example", recipients=["env:SMTP_USER"])
    base.update(kw)
    return sites.DeckConfig(**base)


class FakeSMTP:
    instances, fail_times = [], 0

    def __init__(self, host, port, timeout=None):
        FakeSMTP.instances.append(self)
        self.host, self.port, self.log, self.messages = host, port, [], []
        if FakeSMTP.fail_times > 0:
            FakeSMTP.fail_times -= 1
            raise smtplib.SMTPConnectError(421, "service not available")

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def ehlo(self):
        self.log.append("ehlo")

    def starttls(self, context=None):
        self.log.append("starttls")

    def login(self, user, password):
        self.log.append(("login", user, password))

    def send_message(self, msg):
        self.messages.append(msg)
        FakeSMTP.sent.append(msg)


FakeSMTP.sent = []


def reset_smtp():
    FakeSMTP.instances, FakeSMTP.fail_times, FakeSMTP.sent = [], 0, []


# --------------------------------------------------------------------------- #

class DeckConfigTests(unittest.TestCase):

    def parse(self, **deck_fields):
        entry = {"name": "site-a", "site_url": "sc-domain:a.test", "sheet_id": "m",
                 "reports_sheet_id": "r", "deck": deck_fields}
        return sites.parse_sites_json(json.dumps([entry]))[0]

    def fails(self, *secret, **fields):
        with self.assertRaises(sites.SitesConfigError) as cm:
            self.parse(**fields)
        for text in secret:
            self.assertNotIn(text, str(cm.exception))

    def test_per_deck_type_recipients(self):
        d = self.parse(site_label="Acme", recipients=["a@b.test"], weekly_recipients=[],
                       monthly_recipients=["c@d.test", "env:SMTP_USER"]).deck
        self.assertEqual(d.recipients_for("weekly"), ["env:SMTP_USER"])        # [] means the owner only
        self.assertEqual(d.recipients_for("monthly"), ["c@d.test", "env:SMTP_USER"])
        plain = self.parse(site_label="Acme", recipients=["a@b.test"]).deck
        self.assertEqual((plain.weekly_recipients, plain.monthly_recipients), (None, None))
        self.assertEqual(plain.recipients_for("weekly"), ["a@b.test"])
        self.assertEqual(plain.recipients_for("monthly"), ["a@b.test"])
        self.assertIn("c@d.test", d.sensitive_values())                       # scrubbed from logs
        self.assertNotIn("env:SMTP_USER", d.sensitive_values())

    def test_per_deck_type_recipients_are_validated_without_echoing(self):
        self.fails("secret-text", site_label="Acme", recipients=[], weekly_recipients="secret-text")
        self.fails("not-an-address", site_label="Acme", recipients=[], monthly_recipients=["not-an-address"])
        self.fails(site_label="Acme", recipients=[], weekly_recipients=["a@b.test"] * 21)

    def test_split_recipients(self):
        self.assertEqual(mailer.split_recipients([], USER), ([USER], []))
        self.assertEqual(mailer.split_recipients([USER], USER), ([USER], []))
        self.assertEqual(mailer.split_recipients(["a@b.test"], USER), (["a@b.test"], [USER]))
        self.assertEqual(mailer.split_recipients(["a@b.test", USER, "A@B.test"], USER),
                         (["a@b.test"], [USER]))

    def test_a_full_block(self):
        s = self.parse(enabled=True, language="fa", site_label="Acme", recipients=["a@b.test", "env:SMTP_USER"],
                       weekly=True, monthly=False, brand_color="#1F3A5F", accent_color="#E07A2E",
                       logo_url="https://example.test/logo.png", ai_summary=True)
        d = s.deck
        self.assertEqual((d.enabled, d.language, d.site_label, d.weekly, d.monthly, d.ai_summary),
                         (True, "fa", "Acme", True, False, True))
        self.assertEqual(d.recipients, ["a@b.test", "env:SMTP_USER"])
        self.assertEqual((d.brand_color, d.accent_color), ("#1F3A5F", "#E07A2E"))

    def test_no_block_or_disabled_means_no_decks(self):
        entry = {"name": "site-a", "site_url": "sc-domain:a.test", "sheet_id": "m"}
        self.assertFalse(sites.parse_sites_json(json.dumps([entry]))[0].deck.enabled)
        self.assertFalse(self.parse(enabled=False).deck.enabled)

    def test_defaults(self):
        d = self.parse(site_label="Acme", recipients=["a@b.test"]).deck
        self.assertEqual((d.language, d.weekly, d.monthly, d.ai_summary), ("en", True, True, False))

    def test_validation_without_echoing_values(self):
        self.fails(site_label="", recipients=["a@b.test"])
        self.fails(recipients=["a@b.test"])
        self.fails("a@b.test", site_label="Acme", recipients="a@b.test")       # a list, or [] for the owner only
        self.fails("not-an-address", site_label="Acme", recipients=["not-an-address"])
        self.fails(site_label="Acme", recipients=["a@b.test"], language="de")
        self.fails("ff0000", site_label="Acme", recipients=["a@b.test"], brand_color="ff0000")
        self.fails("http://x.test/l.png", site_label="Acme", recipients=["a@b.test"], logo_url="http://x.test/l.png")
        self.fails(site_label="Acme", recipients=["a@b.test"], weekly=False, monthly=False)
        self.fails(site_label="Acme", recipients=["a@b.test"], ai_summary="yes")
        self.fails("secretfield", site_label="Acme", recipients=["a@b.test"], secretfield=1)
        self.fails(site_label="Acme", recipients=["a@b.test"] * 21)

    def test_values_that_must_not_reach_a_log_are_sensitive(self):
        s = self.parse(site_label="Acme Label", recipients=["a@b.test", "env:SMTP_USER"],
                       logo_url="https://example.test/logo.png")
        values = s.sensitive_values()
        for v in ("Acme Label", "a@b.test", "https://example.test/logo.png"):
            self.assertIn(v, values)
        self.assertNotIn("env:SMTP_USER", values)


class MailerTests(unittest.TestCase):
    SETTINGS = mailer.SmtpSettings("smtp.test", 587, USER, "app-password")

    def setUp(self):
        reset_smtp()

    def test_settings_from_env(self):
        env = {"SMTP_HOST": "smtp.gmail.com", "SMTP_PORT": "587", "SMTP_USER": USER, "SMTP_PASSWORD": "pw"}
        s = mailer.settings_from_env(env)
        self.assertEqual((s.host, s.port, s.user, s.password), ("smtp.gmail.com", 587, USER, "pw"))
        self.assertEqual(mailer.settings_from_env({**env, "SMTP_PORT": ""}).port, 587)
        for missing in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD"):
            self.assertIsNone(mailer.settings_from_env({k: v for k, v in env.items() if k != missing}))
        self.assertIsNone(mailer.settings_from_env({**env, "SMTP_PORT": "abc"}))

    def test_recipients_can_come_from_a_secret(self):
        self.assertEqual(mailer.resolve_recipients(["a@b.test", "env:SMTP_USER"], {"SMTP_USER": USER}),
                         ["a@b.test", USER])
        self.assertEqual(mailer.resolve_recipients(["env:MISSING"], {}), [])

    def test_file_names_read_like_the_subject_and_hold_no_unsafe_characters(self):
        self.assertEqual(mailer.deck_filename("Kalameh SEO Report \u00b7 September 2026"),
                         "Kalameh SEO Report \u00b7 September 2026.pptx")
        self.assertEqual(mailer.deck_filename("My Site / Ltd. SEO Report: 2026"), "My Site Ltd. SEO Report 2026.pptx")
        self.assertEqual(mailer.deck_filename("گزارش سئو نمونه \u00b7 2026"), "گزارش سئو نمونه \u00b7 2026.pptx")
        self.assertEqual(mailer.deck_filename("a:b*c?"), "a b c.pptx")
        self.assertEqual(mailer.deck_filename("///"), "SEO Report.pptx")
        self.assertNotRegex(mailer.deck_filename("x" * 300), r"x{121}")
        self.assertNotIn("PulseO", mailer.deck_filename("Example SEO Report"))

    def test_message_has_subject_recipients_and_the_attachment(self):
        msg = mailer.build_message(USER, ["a@b.test", "c@d.test"], "Example SEO Report", "Body", "x.pptx", b"PPTX")
        self.assertEqual(msg["Subject"], "Example SEO Report")
        self.assertEqual(msg["To"], "a@b.test, c@d.test")
        attachment = list(msg.iter_attachments())[0]
        self.assertEqual(attachment.get_filename(), "x.pptx")
        self.assertEqual(attachment.get_content(), b"PPTX")
        self.assertIn("presentationml.presentation", attachment.get_content_type())

    def test_send_uses_starttls_and_login(self):
        msg = mailer.build_message(USER, [USER], "s", "b", "x.pptx", b"x")
        mailer.send(self.SETTINGS, msg, sleep=lambda s: None, smtp_class=FakeSMTP)
        smtp = FakeSMTP.instances[0]
        self.assertEqual((smtp.host, smtp.port), ("smtp.test", 587))
        self.assertEqual(smtp.log, ["ehlo", "starttls", "ehlo", ("login", USER, "app-password")])
        self.assertEqual(len(smtp.messages), 1)

    def test_one_retry_then_success(self):
        FakeSMTP.fail_times = 1
        slept = []
        mailer.send(self.SETTINGS, mailer.build_message(USER, [USER], "s", "b", "x.pptx", b"x"),
                    sleep=slept.append, smtp_class=FakeSMTP)
        self.assertEqual(len(FakeSMTP.sent), 1)
        self.assertEqual(len(slept), 1)

    def test_two_failures_raise(self):
        FakeSMTP.fail_times = 2
        with self.assertRaises(mailer.MailError):
            mailer.send(self.SETTINGS, mailer.build_message(USER, [USER], "s", "b", "x.pptx", b"x"),
                        sleep=lambda s: None, smtp_class=FakeSMTP)
        self.assertEqual(FakeSMTP.sent, [])

    def test_the_default_smtp_class_is_looked_up_when_sending(self):
        with mock.patch("smtplib.SMTP", FakeSMTP):
            mailer.send(self.SETTINGS, mailer.build_message(USER, [USER], "s", "b", "x.pptx", b"x"))
        self.assertEqual(len(FakeSMTP.sent), 1)


class SummaryTests(unittest.TestCase):

    def reply(self, text):
        return json.dumps({"content": [{"type": "text", "text": text}]}).encode()

    def test_parse(self):
        ok = summary.parse_summary('Here: {"bullets": ["a", "b"], "next_steps": ["c"]} done')
        self.assertEqual(ok, (["a", "b"], ["c"]))
        self.assertIsNone(summary.parse_summary("no json"))
        self.assertIsNone(summary.parse_summary('{"bullets": "x", "next_steps": []}'))
        self.assertIsNone(summary.parse_summary('{"bullets": [], "next_steps": ["c"]}'))
        many = summary.parse_summary(json.dumps({"bullets": list("abcdefgh"), "next_steps": list("wxyz")}))
        self.assertEqual((len(many[0]), len(many[1])), (6, 3))

    def test_request_success_and_what_is_sent(self):
        seen = {}

        def post(url, headers, body, timeout=60):
            seen.update(url=url, headers=headers, body=json.loads(body))
            return self.reply('{"bullets": ["Clicks fell."], "next_steps": ["Check indexing."]}')

        out = summary.request_summary("Clicks: 325", "en", {"ANTHROPIC_API_KEY": "key-123"}, post)
        self.assertEqual(out, (["Clicks fell."], ["Check indexing."]))
        self.assertEqual(seen["url"], "https://api.anthropic.com/v1/messages")
        self.assertEqual(seen["headers"]["x-api-key"], "key-123")
        self.assertEqual(seen["body"]["model"], summary.DEFAULT_MODEL)
        prompt = seen["body"]["messages"][0]["content"]
        self.assertIn("Clicks: 325", prompt)
        self.assertIn("English", prompt)
        self.assertNotIn("key-123", json.dumps(seen["body"]))

    def test_model_and_language_are_configurable(self):
        seen = {}

        def post(url, headers, body, timeout=60):
            seen.update(json.loads(body))
            return self.reply('{"bullets": ["a"], "next_steps": ["b"]}')

        summary.request_summary("x", "fa", {"ANTHROPIC_API_KEY": "k", "AI_SUMMARY_MODEL": "some-model"}, post)
        self.assertEqual(seen["model"], "some-model")
        self.assertIn("Persian", seen["messages"][0]["content"])

    def test_any_problem_means_none(self):
        key = {"ANTHROPIC_API_KEY": "k"}
        self.assertIsNone(summary.request_summary("x", "en", {}, lambda *a, **k: b""))

        def boom(*a, **k):
            raise OSError("network down")

        self.assertIsNone(summary.request_summary("x", "en", key, boom))
        self.assertIsNone(summary.request_summary("x", "en", key, lambda *a, **k: b"not json"))
        self.assertIsNone(summary.request_summary("x", "en", key, lambda *a, **k: self.reply("sorry, no")))


class DeckRunnerTests(ReportsTestCase):

    def setUp(self):
        super().setUp()
        reset_smtp()
        self.site = site_with(deck=deck())
        self.run_reports(self.site)
        self.env = {"SMTP_USER": USER}
        self.smtp = mailer.SmtpSettings("smtp.test", 587, USER, "pw")
        self.sent = []

    def run_decks(self, send=None, env=None, site=None, rebuild=False, post=None, test=False):
        self.logs.clear()
        return deck_runner.run_decks(
            self.sheets, RID, site or self.site, self.smtp, self.env if env is None else env,
            rebuild=rebuild, log=self.logs.append, send=send or (lambda s, m: self.sent.append(m)), test=test,
            fetch_logo=lambda url: None, **({"post": post} if post else {}))

    def state(self):
        return {r[0]: r[1] for r in self.rows("State") if r[0].endswith("_deck_sent")}

    def test_first_run_sends_both_decks(self):
        result = self.run_decks()
        self.assertEqual(result.sent, ["monthly", "weekly"])
        self.assertEqual(result.failed, [])
        self.assertEqual(len(self.sent), 2)
        monthly, weekly = self.sent
        self.assertEqual(monthly["To"], USER)                                  # from the secret
        self.assertEqual(monthly["Subject"], "Example SEO Report \u00b7 September 2026")
        self.assertTrue(weekly["Subject"].startswith("Example SEO Report \u00b7 Week 41"))
        names = [list(m.iter_attachments())[0].get_filename() for m in self.sent]
        self.assertEqual(names, ["Example SEO Report \u00b7 September 2026.pptx", weekly["Subject"] + ".pptx"])
        counts = [len(Presentation(io.BytesIO(list(m.iter_attachments())[0].get_content())).slides)
                  for m in self.sent]
        self.assertGreaterEqual(counts[0], 5)                  # the client deck: slides without data are left out
        self.assertEqual(counts[1], 6)
        self.assertEqual(self.state(), {"monthly_deck_sent": "2026-09", "weekly_deck_sent": "2026-W41"})

    def test_both_emails_are_short_personal_notes_signed_mehdi(self):
        self.run_decks()
        body = self.sent[0].get_body().get_content()
        self.assertIn("Here is the SEO report for Example, September 2026.", body)
        self.assertIn("clicks from Google", body)
        weekly_body = self.sent[1].get_body().get_content()
        self.assertIn("Here is the SEO report for Example, Week 41", weekly_body)
        for text in (body, weekly_body):
            self.assertTrue(text.startswith("Hi,"))
            self.assertTrue(text.strip().endswith("Mehdi"))
            for banned in ("PulseO", "Top alarm", "Clicks:", "Impressions:", "Avg. position:"):
                self.assertNotIn(banned, text)

    def test_nothing_is_sent_twice(self):
        self.run_decks()
        self.run_decks()
        self.assertEqual(len(self.sent), 2)
        self.assertTrue(any("already sent" in l for l in self.logs))

    def test_a_new_week_sends_only_the_weekly_deck(self):
        self.run_decks()
        self.run_reports(self.site, final=date(2026, 10, 21), today=date(2026, 10, 23))
        result = self.run_decks()
        self.assertEqual(result.sent, ["weekly"])
        self.assertEqual(self.state()["weekly_deck_sent"], "2026-W42")
        self.assertEqual(self.state()["monthly_deck_sent"], "2026-09")

    def test_a_new_month_sends_the_monthly_deck_again(self):
        self.run_decks()
        self.run_reports(self.site, final=date(2026, 10, 31), today=date(2026, 11, 5))
        result = self.run_decks()
        self.assertIn("monthly", result.sent)
        self.assertEqual(self.state()["monthly_deck_sent"], "2026-10")

    def test_mark_only_moves_the_state_and_emails_nothing(self):
        self.sent.clear()
        deck_result = deck_runner.run_decks(
            self.sheets, RID, self.site, self.smtp, self.env, log=self.logs.append,
            send=lambda s, m: self.sent.append(m), fetch_logo=lambda url: None, mark_only=True)
        self.assertEqual(self.sent, [])
        self.assertEqual(deck_result.sent, [])
        self.assertEqual(self.state(), {"monthly_deck_sent": "2026-09", "weekly_deck_sent": "2026-W41"})
        self.run_decks()                                           # a normal run afterwards sends nothing
        self.assertEqual(self.sent, [])

    def test_rebuild_sends_both_even_if_already_sent(self):
        self.run_decks()
        result = self.run_decks(rebuild=True)
        self.assertEqual(result.sent, ["monthly", "weekly"])
        self.assertEqual(len(self.sent), 4)

    def test_a_rebuild_can_name_one_deck_kind(self):
        self.run_decks()
        self.sent.clear()
        result = self.run_decks(rebuild={"monthly"})
        self.assertEqual(result.sent, ["monthly"])
        self.assertEqual(len(self.sent), 1)
        self.assertNotIn("Week", self.sent[0]["Subject"])
        self.assertTrue(any("weekly already sent" in l for l in self.logs))

    def test_a_test_copy_goes_to_the_owner_only_is_marked_and_changes_no_state(self):
        site = site_with(deck=deck(monthly_recipients=[self.CLIENT], weekly_recipients=[self.CLIENT]))
        self.run_decks()                                                          # both decks went out normally
        sent_before = self.state()
        self.sent.clear()
        self.run_decks(site=site, rebuild={"monthly"}, test=True)
        (m,) = self.sent
        self.assertEqual((m["To"], m["Bcc"]), (USER, None))
        self.assertTrue(m["Subject"].startswith("[Test] Example SEO Report"))
        self.assertEqual(list(m.iter_attachments())[0].get_filename(), m["Subject"][len("[Test] "):] + ".pptx")
        self.assertEqual(self.state(), sent_before)                              # nothing was marked

    def test_which_site_and_which_decks_a_rebuild_covers(self):
        site_2 = gsc_pulse.SiteRun(site_with(name="site-2", deck=deck()))
        site_3 = gsc_pulse.SiteRun(site_with(name="site-3", deck=deck()))
        f = gsc_pulse.deck_rebuild_for
        self.assertIs(f(site_2, True, "", ""), True)                              # the old input: everything
        self.assertEqual(f(site_2, False, "site-2", "monthly"), {"monthly"})
        self.assertEqual(f(site_2, False, "site-2", "weekly"), {"weekly"})
        self.assertEqual(f(site_2, False, "site-2", ""), {"monthly", "weekly"})
        self.assertEqual(f(site_2, False, "site-2", "both"), {"monthly", "weekly"})
        self.assertEqual(f(site_2, False, "site-2", "nonsense"), set())
        self.assertFalse(f(site_3, False, "site-2", "monthly"))                   # other sites: untouched
        self.assertFalse(f(site_3, False, "", ""))

    def test_weekly_and_monthly_switches(self):
        only_weekly = site_with(deck=deck(monthly=False))
        self.assertEqual(self.run_decks(site=only_weekly).sent, ["weekly"])
        self.sent.clear()
        self.sheets.books[RID]["State"].data = self.sheets.books[RID]["State"].data[:1]
        only_monthly = site_with(deck=deck(weekly=False))
        self.assertEqual(self.run_decks(site=only_monthly).sent, ["monthly"])

    def parsed_site(self, **fields):
        entry = {"name": "site-a", "site_url": "sc-domain:a.test", "sheet_id": "m",
                 "reports_sheet_id": RID,
                 "deck": {"enabled": True, "site_label": "Example", "recipients": [], **fields}}
        return sites.parse_sites_json(json.dumps([entry]))[0]

    def recipients_of(self, **fields):
        """{'monthly': (To, Bcc), 'weekly': (To, Bcc)} of the decks sent for this deck block."""
        self.sent.clear()
        self.assertEqual(sorted(self.run_decks(site=self.parsed_site(**fields), rebuild=True).sent),
                         ["monthly", "weekly"])
        got = {}
        for m in self.sent:
            kind = "monthly" if "Week" not in m["Subject"] else "weekly"
            got[kind] = (m["To"], m["Bcc"])
        return got

    CLIENT = "client@example.test"

    def test_weekly_to_the_owner_only_and_monthly_also_to_the_client(self):
        got = self.recipients_of(weekly_recipients=[], monthly_recipients=[self.CLIENT])
        self.assertEqual(got["weekly"], (USER, None))
        self.assertEqual(got["monthly"], (self.CLIENT, USER))                 # the owner in blind copy

    def test_without_a_list_per_deck_type_recipients_applies_to_both(self):
        got = self.recipients_of(recipients=[self.CLIENT])
        self.assertEqual((got["weekly"], got["monthly"]), ((self.CLIENT, USER), (self.CLIENT, USER)))

    def test_an_empty_list_per_deck_type_overrides_recipients(self):
        got = self.recipients_of(recipients=[self.CLIENT], weekly_recipients=[])
        self.assertEqual(got["weekly"], (USER, None))
        self.assertEqual(got["monthly"], (self.CLIENT, USER))

    def test_no_lists_at_all_means_the_owner_only_as_before(self):
        got = self.recipients_of()
        self.assertEqual(got["weekly"], (USER, None))
        self.assertEqual(got["monthly"], (USER, None))

    def test_the_owner_is_never_listed_twice(self):
        got = self.recipients_of(monthly_recipients=[USER.upper(), self.CLIENT, self.CLIENT.upper(), "env:SMTP_USER"])
        self.assertEqual(got["monthly"], (self.CLIENT, USER))
        self.assertEqual(self.recipients_of(monthly_recipients=[USER])["monthly"], (USER, None))

    def test_a_client_address_kept_in_a_secret_works_too(self):
        self.env["CLIENT_MAIL"] = self.CLIENT
        got = self.recipients_of(monthly_recipients=["env:CLIENT_MAIL"])
        self.assertEqual(got["monthly"], (self.CLIENT, USER))

    def test_the_attachment_is_the_same_for_everyone_and_nothing_is_stored(self):
        self.recipients_of(monthly_recipients=[self.CLIENT])
        for m in self.sent:
            self.assertEqual(len(list(m.iter_attachments())), 1)
        self.assertEqual(glob.glob("*.pptx"), [])

    def test_a_failed_send_keeps_the_state_of_the_deck_that_did_send(self):
        def flaky(settings, message):
            if "Week" not in message["Subject"]:
                raise mailer.MailError(f"SMTPAuthenticationError: bad password for {USER}")
            self.sent.append(message)

        result = self.run_decks(send=flaky)
        self.assertEqual((result.sent, result.failed), (["weekly"], ["monthly"]))
        self.assertEqual(self.state(), {"weekly_deck_sent": "2026-W41"})
        self.assertTrue(any("decks: monthly failed" in l for l in self.logs))
        retry = self.run_decks()                                              # the next run sends the monthly one
        self.assertEqual(retry.sent, ["monthly"])
        self.assertEqual(len(self.sent), 2)

    def test_scrub_is_applied_to_the_failure_reason(self):
        def fail(settings, message):
            raise mailer.MailError(f"rejected {USER}")

        self.logs.clear()
        deck_runner.run_decks(self.sheets, RID, self.site, self.smtp, self.env, log=self.logs.append,
                              scrub=lambda s: s.replace(USER, "(hidden)"), send=fail,
                              fetch_logo=lambda u: None)
        text = "\n".join(self.logs)
        self.assertNotIn(USER, text)
        self.assertIn("(hidden)", text)

    def test_the_owner_gets_the_deck_even_when_the_secret_recipient_is_not_in_the_environment(self):
        result = self.run_decks(env={})
        self.assertEqual(result.failed, [])
        self.assertEqual([m["To"] for m in self.sent], [USER, USER])

    def test_no_report_data_yet(self):
        self.sheets = FakeSheets()
        self.sheets.add_book(RID)
        self.sa.records = []                       # nothing in Search Console: the report tabs stay empty
        self.run_reports(self.site)
        result = self.run_decks()
        self.assertEqual(result.sent, [])
        self.assertEqual(self.sent, [])
        self.assertTrue(any("no report data yet" in l for l in self.logs))

    def test_ai_summary_goes_into_the_weekly_deck_and_sends_only_aggregates(self):
        seen = {}

        def post(url, headers, body, timeout=60):
            seen["body"] = body.decode()
            return json.dumps({"content": [{"type": "text", "text":
                '{"bullets": ["Clicks fell sharply.", "Impressions fell too."], "next_steps": ["Check indexing."]}'}]}).encode()

        site = site_with(deck=deck(ai_summary=True))
        self.env["ANTHROPIC_API_KEY"] = "key-123"
        self.run_decks(site=site, post=post)
        prs = Presentation(io.BytesIO(list(self.sent[1].iter_attachments())[0].get_content()))
        text = "\n".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
        self.assertIn("Clicks fell sharply.", text)
        self.assertNotIn("Summary not generated automatically", text)
        for private in (RID, "key-123", USER, "main-id", "alpha.test"):
            self.assertNotIn(private, seen["body"])
        self.assertIn("Clicks:", seen["body"])

    def test_ai_summary_failure_falls_back_to_the_placeholder_without_failing(self):
        def post(*a, **k):
            raise OSError("api down")

        site = site_with(deck=deck(ai_summary=True))
        self.env["ANTHROPIC_API_KEY"] = "key-123"
        result = self.run_decks(site=site, post=post)
        self.assertEqual(result.failed, [])
        prs = Presentation(io.BytesIO(list(self.sent[1].iter_attachments())[0].get_content()))
        text = "\n".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
        self.assertIn("Summary not generated automatically", text)
        self.assertTrue(any("AI summary not available" in l for l in self.logs))

    def test_ai_summary_without_a_key_uses_the_placeholder(self):
        site = site_with(deck=deck(ai_summary=True))
        self.run_decks(site=site)
        prs = Presentation(io.BytesIO(list(self.sent[1].iter_attachments())[0].get_content()))
        text = "\n".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
        self.assertIn("Summary not generated automatically", text)

    def test_the_persian_deck_and_email(self):
        site = site_with(deck=deck(language="fa", site_label="نمونه"))
        self.run_decks(site=site)
        self.assertIn("گزارش سئو", self.sent[0]["Subject"])
        self.assertNotIn("PulseO", self.sent[0]["Subject"])
        self.assertEqual(list(self.sent[0].iter_attachments())[0].get_filename(), self.sent[0]["Subject"] + ".pptx")


class WeeklyMoversTests(ReportsTestCase):

    def test_latest_week_gets_its_top_movers_once(self):
        self.run_reports()
        rows = self.rows(reports.WEEKLY_MOVERS)
        self.assertTrue(rows)
        self.assertEqual({r[0] for r in rows}, {"2026-W41"})
        self.assertEqual({r[1] for r in rows}, {"Growing", "Declining"})
        for kind in ("Growing", "Declining"):
            self.assertLessEqual(len([r for r in rows if r[1] == kind]), 10)
        self.assertEqual(self.tab(reports.WEEKLY_MOVERS)[0], reports.KEYWORDS_HEADER)
        again = len(rows)
        self.run_reports()
        self.assertEqual(len(self.rows(reports.WEEKLY_MOVERS)), again)

    def test_a_new_week_adds_a_new_block_and_keeps_the_old_one(self):
        self.run_reports()
        before = [list(r) for r in self.rows(reports.WEEKLY_MOVERS)]
        self.run_reports(final=date(2026, 11, 4), today=date(2026, 11, 6))      # the week 26 Oct - 1 Nov spans two months
        rows = self.rows(reports.WEEKLY_MOVERS)
        self.assertEqual(rows[:len(before)], before)
        self.assertEqual({r[0] for r in rows[len(before):]}, {"2026-W44"})

    def test_the_list_length_is_configurable(self):
        self.run_reports(site_with(thresholds={"weekly_movers_top_n": 2}))
        for kind in ("Growing", "Declining"):
            self.assertLessEqual(len([r for r in self.rows(reports.WEEKLY_MOVERS) if r[1] == kind]), 2)


# --------------------------------------------------------------------------- #

ENV_KEYS = ("SITES_JSON", "SITE_URL", "SHEET_ID", "HISTORY_SHEET_ID", "DAYS", "GITHUB_ACTIONS",
            "DISCOVER_SITE", "REBUILD_ALARMS", "REBUILD_DECKS", "SMTP_HOST", "SMTP_PORT", "SMTP_USER",
            "SMTP_PASSWORD", "ANTHROPIC_API_KEY")
LABEL = "Private Client Label"
SITE = {"name": "site-a", "site_url": "sc-domain:alpha.test", "sheet_id": "A-MAIN-ID",
        "history_sheet_id": "A-HIST-ID", "reports_sheet_id": "A-REPORTS-ID",
        "deck": {"enabled": True, "language": "en", "site_label": LABEL, "recipients": ["env:SMTP_USER"],
                 "logo_url": "https://logos.example.test/private-logo.png"}}
SMTP_ENV = {"SMTP_HOST": "smtp.test", "SMTP_PORT": "587", "SMTP_USER": USER, "SMTP_PASSWORD": "very-secret-pw"}


class DecksInTheRunTests(unittest.TestCase):

    def setUp(self):
        patch_time(self)
        reset_smtp()
        today = date.today()
        self.sa = FakeSA(build_records(start=today - timedelta(days=520), end=today))
        self.sheets = FakeSheets()
        for book in ("A-MAIN-ID", "A-REPORTS-ID"):
            self.sheets.add_book(book)

    def run_main(self, env=None, site=SITE):
        clean = {k: v for k, v in os.environ.items() if k not in ENV_KEYS}
        clean.update({"SITES_JSON": json.dumps([site]), "GITHUB_ACTIONS": "true", **(env or {})})
        out, error = io.StringIO(), None
        with mock.patch.dict(os.environ, clean, clear=True), \
             mock.patch.object(sys, "argv", ["gsc_pulse.py", "--config", "none.yaml"]), \
             mock.patch.object(gsc_pulse, "build",
                               lambda n, v, credentials=None: self.sa if n == "webmasters" else self.sheets), \
             mock.patch.object(gsc_pulse, "get_credentials", lambda f: object()), \
             mock.patch.object(gsc_pulse, "PUBLIC_LOGS", True), \
             mock.patch.object(gsc_pulse, "_SENSITIVE", []), \
             mock.patch("smtplib.SMTP", FakeSMTP), \
             mock.patch("deck_runner.decks.fetch_logo", lambda url: None), \
             contextlib.redirect_stdout(out):
            try:
                gsc_pulse.main()
            except (Exception, SystemExit) as e:        # noqa: BLE001
                error = e
        return out.getvalue(), error

    def test_without_smtp_settings_nothing_is_built(self):
        out, err = self.run_main()
        self.assertIsNone(err)
        self.assertIn("decks: off (no SMTP settings)", out)
        self.assertEqual(FakeSMTP.sent, [])
        state_tab = self.sheets.books["A-REPORTS-ID"].get("State")
        self.assertFalse(state_tab and any(str(r[0]).endswith("_deck_sent") for r in state_tab.data[1:]))

    def test_decks_are_emailed_after_the_reports_and_before_history(self):
        site = dict(SITE, history_sheet_id="A-HIST-ID")
        out, err = self.run_main(SMTP_ENV, site)
        self.assertIsNone(err, out)
        self.assertLess(out.index("a · Reports"), out.index("a · Decks"))
        self.assertLess(out.index("a · Decks"), out.index("a · History"))
        self.assertRegex(out, r"decks: monthly deck sent \(1 recipient, \d+ slides\)")
        summary_line = [l for l in out.splitlines() if l.strip().startswith("site-a: ok")][0]
        self.assertIn("decks: sent monthly + weekly", summary_line)
        self.assertEqual(len(FakeSMTP.sent), 2)
        smtp = FakeSMTP.instances[0]
        self.assertIn(("login", USER, "very-secret-pw"), smtp.log)
        self.assertEqual(FakeSMTP.sent[0]["To"], USER)

    def test_an_unchanged_second_run_sends_nothing_new(self):
        self.run_main(SMTP_ENV)
        reset_smtp()
        out, err = self.run_main(SMTP_ENV)
        self.assertIsNone(err)
        self.assertEqual(FakeSMTP.sent, [])
        self.assertIn("already sent", out)

    def test_rebuild_decks_input_sends_again(self):
        self.run_main(SMTP_ENV)
        reset_smtp()
        out, err = self.run_main({**SMTP_ENV, "REBUILD_DECKS": "true"})
        self.assertIsNone(err)
        self.assertEqual(len(FakeSMTP.sent), 2)

    def test_nothing_private_reaches_the_log(self):
        FakeSMTP.fail_times = 99                                              # every send fails
        out, err = self.run_main(SMTP_ENV)
        self.assertIsInstance(err, gsc_pulse.GscPulseError)                   # the run ends red
        self.assertIn("decks: monthly failed", out)
        for private in (USER, "very-secret-pw", LABEL, "private-logo", "A-REPORTS-ID", "alpha.test"):
            self.assertNotIn(private, out)

    def test_a_deck_failure_does_not_block_history(self):
        FakeSMTP.fail_times = 99
        site = dict(SITE, history_sheet_id="A-HIST-ID")
        out, err = self.run_main(SMTP_ENV, site)
        self.assertIn("State", self.sheets.books["A-HIST-ID"])
        self.assertIn("site-a: failed", out)

    def test_decks_need_a_reports_sheet(self):
        site = {k: v for k, v in SITE.items() if k != "reports_sheet_id"}
        out, err = self.run_main(SMTP_ENV, site)
        self.assertIsNone(err)
        self.assertIn("decks: off (no reports sheet)", out)


class NamingAndPrivacyTests(unittest.TestCase):

    def read(self, name):
        with open(os.path.join(ROOT, name), encoding="utf-8") as f:
            return f.read()

    def test_the_product_is_called_pulseo(self):
        self.assertEqual(self.read("README.md").splitlines()[0], "# PulseO")
        self.assertNotIn("GSC Pulse", self.read("README.md"))
        self.assertTrue(self.read(".github/workflows/gsc_pulse.yml").startswith("name: PulseO weekly sync"))
        for lang in ("en", "fa"):                                  # ... but the emails and decks never say so
            self.assertNotIn("PulseO", deck_runner.i18n.table(lang)["mail_subject"])
            self.assertNotIn("PulseO", deck_runner.i18n.table(lang)["mail_sign"])

    def test_no_workflow_uploads_anything(self):
        for path in glob.glob(os.path.join(ROOT, ".github", "workflows", "*.yml")):
            with open(path, encoding="utf-8") as f:
                self.assertNotIn("upload-artifact", f.read(), path)

    def test_decks_are_never_written_to_disk_by_the_code(self):
        for module in ("decks.py", "deck_runner.py", "deck_data.py", "mailer.py", "summary.py"):
            source = self.read(module)
            self.assertNotRegex(source, r"open\([^)]*['\"]w", module)
            self.assertNotIn("prs.save(\"", source)
        self.assertIn("*.pptx", self.read(".gitignore"))

    def test_only_the_two_templates_are_in_the_repo(self):
        found = [os.path.relpath(p, ROOT).replace("\\", "/")
                 for p in glob.glob(os.path.join(ROOT, "**", "*.pptx"), recursive=True)]
        self.assertEqual(sorted(found), ["templates/gsc_pulse_template.pptx", "templates/gsc_pulse_template_fa.pptx"])

    def test_workflow_passes_the_smtp_secrets_and_the_rebuild_input(self):
        text = self.read(".github/workflows/gsc_pulse.yml")
        for name in ("SMTP_HOST", "SMTP_PORT", "SMTP_USER", "SMTP_PASSWORD", "ANTHROPIC_API_KEY", "REBUILD_DECKS"):
            self.assertIn(name, text)
        self.assertIn("rebuild_decks:", text)


if __name__ == "__main__":
    unittest.main()
