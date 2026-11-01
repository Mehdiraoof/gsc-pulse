"""
Tests for the client-ready monthly deck (v2.4): what goes on each slide (client_data) and how it
is drawn (client_deck). All data is invented (client_fixtures.py).
"""

import io
import re
import unittest

from pptx import Presentation

import client_data as cd
import client_deck
import client_fixtures as fx
import client_i18n
import deck_runner
import resolved

PLACEHOLDER = re.compile(r"\{(\w+)\}")


def report(**kw):
    return fx.monthly_report(**kw)


def client(lang="en", **kw):
    return cd.ClientReport(report(**kw), lang, "Example")


def texts_of(prs):
    out = []
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame:
                out.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    out += [c.text_frame.text for c in row.cells]
    return out


def build(lang="en", **kw):
    rep = report(**kw)
    data = client_deck.build_monthly(rep, fx.Cfg(lang), deck_runner.TEMPLATE_DIR)
    return Presentation(io.BytesIO(data))


class SlideSelectionTests(unittest.TestCase):

    def test_a_full_deck_has_fifteen_slides_in_the_agreed_order(self):
        ids = [sid for sid, _ in client().slides()]
        self.assertEqual(ids, ["headline", "scorecard", "climbs", "grew", "new", "work", "trend",
                               "sections", "wins", "pages", "ideas", "next", "watch", "focus"])
        self.assertEqual(len(build().slides), 15)

    def test_slides_without_data_are_left_out_but_the_work_slide_stays(self):
        ids = [sid for sid, _ in client(segments=False, opps=False, gone=0).slides()]
        self.assertIn("work", ids)
        for gone in ("sections", "wins", "pages", "ideas", "focus"):
            self.assertNotIn(gone, ids)
        self.assertEqual(len(build(segments=False, opps=False, gone=0).slides), 1 + len(ids))

    def test_no_data_at_all_still_gives_title_and_work_slide(self):
        rep = report(segments=False, opps=False)
        rep.keywords, rep.pages, rep.gone = [], [], []
        ids = [sid for sid, _ in cd.ClientReport(rep, "en", "Example").slides()]
        self.assertIn("work", ids)
        self.assertNotIn("climbs", ids)
        self.assertNotIn("grew", ids)
        self.assertNotIn("new", ids)

    def test_slides_point_at_each_other_by_their_real_numbers(self):
        c = client()
        nums = {sid: i + 2 for i, (sid, _) in enumerate(c.slides())}
        focus = dict(c.slides())["focus"]
        self.assertIn(f"slide {nums['wins']}", focus["cards"][0]["body"])
        self.assertIn(f"slide {nums['ideas']}", focus["cards"][2]["body"])
        nxt = dict(c.slides())["next"]
        self.assertTrue(any(f"slide {nums['pages']}" in b for b in nxt["bullets"]))

    def test_without_the_target_slide_the_text_does_not_name_a_slide(self):
        rep = report()
        rep.opps = [o for o in rep.opps if o["type"] not in ("update_page", "decliner")]
        slides = dict(cd.ClientReport(rep, "en", "Example").slides())
        self.assertNotIn("pages", slides)
        bullets = slides["next"]["bullets"]
        self.assertIn("Guides grow most next month: refresh the top pages first.", bullets)
        self.assertFalse(any("from slide" in b for b in bullets))


class HeadlineTests(unittest.TestCase):

    def head(self, mutate=None, **kw):
        rep = report(**kw)
        if mutate:
            mutate(rep)
        return cd.ClientReport(rep, "en", "Example").headline()

    def test_best_position_of_the_year_comes_first(self):
        d = self.head()
        self.assertEqual(d["headline"], "Example ranks higher than it has all year.")
        self.assertIn("Average position improved to 6.2", d["support"])
        self.assertIn("more searchers clicked through (CTR 3.61%)", d["support"])
        self.assertIn("Guides led the growth", d["support"])
        self.assertEqual(d["cards"][0]["title"], "Best ranking of the year")
        self.assertIn("up from 6.4 last month and 7.6 a year ago", d["cards"][0]["body"])

    def test_then_clicks_up_ten_percent_on_last_year(self):
        def worse(rep):
            rep.rows[(fx.MONTHS[-1], "All")]["Avg. Position"] = 8.0
        self.assertEqual(self.head(worse)["headline"], "Clicks are up 37% on last year.")

    def test_then_more_searchers_clicked_through(self):
        def mutate(rep):
            row = rep.rows[(fx.MONTHS[-1], "All")]
            row["Avg. Position"] = 8.0
            row["Clicks YoY %"] = 5.0
        self.assertEqual(self.head(mutate)["headline"], "More searchers clicked through.")

    def test_then_the_biggest_growing_section(self):
        def mutate(rep):
            row = rep.rows[(fx.MONTHS[-1], "All")]
            row["Avg. Position"] = 8.0
            row["Clicks YoY %"] = 5.0
            row["CTR (%)"] = 1.0
        self.assertEqual(self.head(mutate)["headline"], "Guides led the growth.")

    def test_otherwise_a_calm_sentence(self):
        def mutate(rep):
            row = rep.rows[(fx.MONTHS[-1], "All")]
            row["Avg. Position"] = 8.0
            row["Clicks YoY %"] = 5.0
            row["CTR (%)"] = 1.0
            for s in fx.SEGMENT_CLICKS:
                rep.rows[(fx.MONTHS[-1], s)]["Clicks MoM %"] = -3.0
        self.assertEqual(self.head(mutate)["headline"], "Example held its ground this month.")

    def test_clicks_vs_last_year_card_wording(self):
        d = self.head()
        self.assertEqual(d["cards"][1]["title"], "Clicks up 37% on last year")

        def steady(rep):
            rep.rows[(fx.MONTHS[-1], "All")]["Clicks YoY %"] = -1.1
            rep.rows[(fx.MONTHS[-1], "All")]["Impr YoY %"] = -7.0
        card = self.head(steady)["cards"][1]
        self.assertEqual(card["title"], "Clicks held steady vs last year")
        self.assertIn("within 1% of last September even with fewer searches", card["body"])

    def test_the_room_to_grow_card_sums_the_top_five_gains(self):
        d = self.head()
        self.assertEqual(d["cards"][2]["title"], "Big room to grow")
        self.assertIn("5 keywords", d["cards"][2]["body"])
        self.assertIn("475 extra clicks a month", d["cards"][2]["body"])

    def test_no_opportunities_means_no_third_card(self):
        self.assertEqual(len(self.head(opps=False)["cards"]), 2)


class ScorecardTests(unittest.TestCase):

    def card(self, mutate=None):
        rep = report()
        if mutate:
            mutate(rep)
        d = cd.ClientReport(rep, "en", "Example").scorecard()
        return {c["label"]: c for c in d["cards"]}, d

    def test_values_and_gain_colours(self):
        cards, d = self.card()
        self.assertEqual(cards["Clicks"]["value"], "1,300")
        self.assertEqual(cards["Impressions"]["value"], "36,000")
        self.assertEqual(cards["CTR"]["value"], "3.61%")
        self.assertEqual(cards["Avg. position"]["value"], "6.2")
        self.assertEqual(cards["Clicks"]["lines"][0]["tone"], cd.GOOD)
        self.assertEqual(cards["Avg. position"]["lines"][0]["text"], "↑ 0.2")
        self.assertEqual(cards["Avg. position"]["lines"][0]["tone"], cd.GOOD)
        self.assertEqual(d["kicker"], "September 2026 vs Aug and vs Sep 2025")

    def test_a_drop_from_fewer_searches_is_grey_not_red(self):
        cards, _ = self.card()
        line = cards["Impressions"]["lines"][0]
        self.assertTrue(line["text"].startswith("−"))
        self.assertEqual(line["tone"], cd.GREY)
        self.assertTrue(cards["Impressions"]["note"].startswith("Fewer searches this month"))

    def test_red_only_when_the_position_really_fell(self):
        def worse(rep):
            row = rep.rows[(fx.MONTHS[-1], "All")]
            row["Avg. Position"] = 7.0                      # 0.6 worse than 6.4
            row["Clicks MoM %"] = -5.0
        cards, _ = self.card(worse)
        self.assertEqual(cards["Impressions"]["lines"][0]["tone"], cd.RED)
        self.assertEqual(cards["Avg. position"]["lines"][0]["tone"], cd.RED)
        self.assertIn("rankings slipped", cards["Clicks"]["note"])

        def slightly(rep):
            rep.rows[(fx.MONTHS[-1], "All")]["Avg. Position"] = 6.7        # 0.3 worse: grey
        cards, _ = self.card(slightly)
        self.assertEqual(cards["Avg. position"]["lines"][0]["tone"], cd.GREY)

    def test_old_urls_leaving_are_mentioned_from_five_pages_on(self):
        def fewer_searches(rep):
            rep.gone = rep.gone[:4]
        cards, _ = self.card(fewer_searches)
        self.assertEqual(cards["Impressions"]["note"], "Fewer searches this month.")
        cards, _ = self.card()
        self.assertIn("6 old URLs no longer show in Google", cards["Impressions"]["note"])

    def test_aio_card_only_when_imported(self):
        cards, _ = self.card()
        self.assertNotIn("AIO impressions", cards)

        def imported(rep):
            rep.rows[(fx.MONTHS[-1], "All")]["AIO Impressions"] = 4200
        cards, _ = self.card(imported)
        self.assertEqual(cards["AIO impressions"]["value"], "4,200")

    def test_the_footnote_explains_grey(self):
        self.assertIn("grey, not red", self.card()[1]["footnote"])


class ProgressSlideTests(unittest.TestCase):

    def test_climbs_follow_the_rules_and_the_order(self):
        d = client().climbs()
        labels = [r["label"] for r in d["rows"]]
        self.assertEqual(labels, ["/guide/blue-shoes/", fx.PERSIAN[0], fx.PERSIAN[1], fx.ENGLISH[0],
                                  fx.ENGLISH[1]])
        self.assertNotIn(fx.ENGLISH[2], labels)               # too few impressions
        self.assertNotIn("/", labels)                          # only 0.7 steps
        first = d["rows"][0]
        self.assertEqual((first["prev"], first["now"], first["change"]), ("8.0", "5.9", "↑ 2.1"))
        self.assertEqual(d["card"]["from"], "7.6")
        self.assertEqual(d["card"]["to"], "6.2")
        self.assertEqual(d["header"][1:3], ["Aug", "Sep"])

    def test_the_footnote_counts_rankings_that_reached_page_one(self):
        d = client().climbs()
        self.assertIn(fx.PERSIAN[1], [r["label"] for r in d["rows"]])
        self.assertEqual(d["footnote"], "")                                        # the table already shows it

    def test_the_card_explains_level_clicks_when_impressions_fell_year_on_year(self):
        rep = report()
        row = rep.rows[(fx.MONTHS[-1], "All")]
        row["Impr YoY %"], row["Clicks YoY %"] = -7.0, -1.0
        d = cd.ClientReport(rep, "en", "Example").climbs()
        self.assertIn("clicks stayed level with last year while impressions fell 7%", d["card"]["body"])

    def test_no_climbs_means_no_slide(self):
        rep = report()
        rep.keywords = [k for k in rep.keywords if k["Type"] != "Growing"]
        rep.pages = []
        self.assertIsNone(cd.ClientReport(rep, "en", "Example").climbs())

    def test_what_grew(self):
        d = client().grew()
        self.assertEqual([r["change"] for r in d["keywords"]], ["+17", "+12", "+7", "+5", "+4"])
        self.assertEqual([r["label"] for r in d["pages"]], ["/", "/guide/blue-shoes/", "/shop/hat/"])
        self.assertEqual(d["pages"][0]["clicks"], "794")
        self.assertEqual(d["fact"], "Year on year, Guides are up 70% and /guide/blue-shoes/ gained 80 clicks.")

    def test_new_searches_are_the_three_with_most_impressions(self):
        d = client().new_searches()
        self.assertEqual([c["pill"] for c in d["cards"]], ["276 impressions", "120 impressions", "90 impressions"])
        self.assertEqual(d["cards"][0]["position"], "position 8.4")
        self.assertIn("for the first time this month", d["intro"])

    def test_the_work_slide_is_a_placeholder_with_four_prompts(self):
        d = client().work()
        self.assertEqual(d["fill"], "Mehdi fills this slide before sending.")
        self.assertEqual(len(d["prompts"]), 4)
        self.assertEqual(d["prompts"][3], "Anything waiting on the Example team")


class TrendAndSectionTests(unittest.TestCase):

    def test_trend_notes(self):
        d = client().trend()
        self.assertEqual(len(d["values"]), 13)
        self.assertEqual(d["labels"][0], "Sep 25")
        self.assertEqual(d["labels"][-1], "Sep 26")
        titles = [n["title"] for n in d["notes"]]
        self.assertEqual(titles, ["December", "Now", "Last year, the next 3 months"])
        self.assertEqual(d["notes"][0]["body"], "Record month: 1,400 clicks.")
        self.assertEqual(d["notes"][1]["tone"], "periwinkle")
        self.assertEqual(d["notes"][2]["body"], "Clicks changed by +23.2% on average compared with September.")
        self.assertEqual(d["title"], "13-month trend")

    def low_month(self, resolved_list=None):
        rep = report(resolved_list=resolved_list)
        rep.rows[("2026-03", "All")]["Clicks"] = 300
        return cd.ClientReport(rep, "en", "Example").trend()

    def test_a_very_low_month_is_red_unless_it_is_on_the_resolved_list(self):
        notes = {n["title"]: n for n in self.low_month()["notes"]}
        self.assertEqual(notes["March"]["tone"], "red")
        self.assertEqual(notes["March"]["body"], "Lowest month: 300 clicks.")
        notes = {n["title"]: n for n in self.low_month([["month", "2026-03", "outage", ""]])["notes"]}
        self.assertEqual(notes["March"]["tone"], "grey")
        self.assertEqual(notes["March"]["body"], "Known and fixed: 300 clicks.")

    def test_sections_chart_table_and_overlap_note(self):
        d = client().sections()
        self.assertEqual([r["label"] for r in d["rows"]], ["Guides", "Blog & other", "Shop"])
        self.assertEqual(d["chart_title"], "Clicks in September")
        guides = d["rows"][0]
        self.assertEqual(guides["mom"], "+11.1%")
        self.assertEqual(guides["mom_tone"], cd.GOOD)
        shop = d["rows"][2]
        self.assertEqual((shop["mom_tone"], shop["yoy_tone"]), (cd.GREY, cd.GREY))   # grey, never red
        self.assertIn("add up to about 8% more than the site total", d["footnote"])

    def test_sections_that_add_up_say_so_and_one_section_means_no_slide(self):
        rep = report()
        for name in fx.SEGMENT_CLICKS:
            rep.rows[(fx.MONTHS[-1], name)]["Clicks"] = 0
        rep.rows[(fx.MONTHS[-1], "Guides")]["Clicks"] = 700
        rep.rows[(fx.MONTHS[-1], "Shop")]["Clicks"] = 200
        rep.rows[(fx.MONTHS[-1], "Blog & other")]["Clicks"] = 400
        self.assertEqual(cd.ClientReport(rep, "en", "Example").sections()["footnote"],
                         "The sections add up to the site total.")
        self.assertIsNone(client(segments=False).sections())


class OpportunitySlideTests(unittest.TestCase):

    def test_quick_wins_one_row_per_keyword_biggest_gain_first(self):
        d = client().quick_wins()
        self.assertEqual([r["gain"] for r in d["rows"]], ["+145", "+100", "+95", "+90", "+45"])
        self.assertEqual(d["total"], 475)
        self.assertEqual(d["card"], "About +475 clicks a month together.")
        self.assertIn("Estimate based on typical top-3 click rates", d["card_note"])
        self.assertEqual(d["rows"][0]["searches"], "1,825")                  # 5,475 impressions in 90 days / 3
        self.assertEqual(d["rows"][1]["position"], "7.9")

    def test_pages_to_update_hold_the_diagnosis_and_the_action(self):
        d = client().pages_to_update()
        self.assertEqual([r["page"] for r in d["rows"]],
                         ["/guide/g/", "/guide/h/", "/guide/old-tips/"])
        self.assertEqual(d["rows"][0]["text"], "Ranks well, title doesn't win the click. Rewrite title + meta.")
        self.assertEqual(d["rows"][0]["ctr"], "1.05%")

    def test_ideas_questions_and_competing_pages(self):
        d = client().ideas()
        self.assertEqual([c["keyword"] for c in d["cards"]],
                         ["what is a trail shoe", "how to clean a bike", "why do boots leak"])
        self.assertEqual(d["cards"][0]["sub"], "1,756 impr. · position 8")
        self.assertEqual(d["cards"][2]["sub"], "625 impr. · CTR 0.90%")
        self.assertIn("Each of these keywords has 3 pages competing", d["merge"]["body"])
        self.assertEqual(d["merge"]["names"], [fx.ENGLISH[0], fx.ENGLISH[1]])       # on lines of their own, not in the sentence
        self.assertNotIn(fx.ENGLISH[0], d["merge"]["body"])

    def test_next_month_uses_last_years_change_and_real_slide_numbers(self):
        d = dict(client().slides())["next"]
        self.assertEqual(d["title"], "Get ready for October")
        self.assertEqual(d["intro"], "What happened from September to October last year:")
        self.assertEqual(d["cards"][0]["label"], "Whole site")
        self.assertEqual(d["cards"][0]["value"], "+6.3%")
        self.assertEqual(d["cards"][0]["sub"], "+4.8% impressions")
        self.assertEqual([c["label"] for c in d["cards"]], ["Whole site", "Guides", "Blog & other", "Shop"])
        self.assertEqual(d["cards"][3]["tone"], cd.GOOD)
        self.assertIn("Guides grow most next month: refresh the top pages from slide 11 first.", d["bullets"])
        self.assertTrue(any("Update /guide/k/ before October: it grew +50% into October last year." == b
                            for b in d["bullets"]))

    def test_watching_is_calm_and_short(self):
        d = client().watching()
        titles = [i["title"] for i in d["items"]]
        self.assertEqual(titles, ["Shop", "Old URLs", "/guide/old-tips/", "/guide/j/"])
        self.assertIn("Down 35% on last year", d["items"][0]["body"])
        self.assertEqual(d["items"][1]["body"], "6 old URLs no longer show in Google (expected after "
                                                 "redirects); we check they pass traffic on.")
        self.assertEqual(d["items"][2]["body"], "Clicks 137 → 55. Ranking fell.")
        self.assertLessEqual(len(d["items"]), 5)

    def test_the_old_urls_line_needs_five_pages(self):
        titles = [i["title"] for i in client(gone=4).watching()["items"]]
        self.assertNotIn("Old URLs", titles)

    def test_focus_cards(self):
        d = dict(client().slides())["focus"]
        self.assertEqual([c["title"] for c in d["cards"]],
                         ["Lift the 5 quick wins", "Merge competing pages", "Answer the questions",
                          "Prepare for October"])
        self.assertIn("About +475 clicks a month at stake", d["cards"][0]["body"])
        self.assertEqual(d["cards"][0]["tone"], "amber")
        self.assertIn("2 topics", d["cards"][1]["body"])


class ResolvedTests(unittest.TestCase):

    def ids_and_text(self, resolved_list):
        c = client(resolved_list=resolved_list)
        slides = dict(c.slides())
        return c, slides

    def test_the_parser_ignores_unknown_kinds_and_empty_values(self):
        r = resolved.parse([["page", "/a/", "n", "2026-10-01"], ["keyword", "Some Words", "", ""],
                            ["month", "2026-03", "", ""], ["segment", "Shop", "", ""],
                            ["colour", "x", "", ""], ["page", "", "", ""], ["keyword"]])
        self.assertTrue(r.page("https://example.test/a/"))
        self.assertTrue(r.page("/a/"))
        self.assertFalse(r.page("https://example.test/b/"))
        self.assertTrue(r.keyword("some words"))
        self.assertTrue(r.month("2026-03"))
        self.assertTrue(r.segment("SHOP"))
        self.assertEqual(r.note("page", "/a/"), "n")
        self.assertFalse(resolved.Resolved())

    def test_persian_pages_match_by_url_or_decoded_path(self):
        r = resolved.parse([["page", "/مایع-ضد/", "", ""]])
        self.assertTrue(r.page("https://example.test/%D9%85%D8%A7%DB%8C%D8%B9-%D8%B6%D8%AF/"))

    def test_a_resolved_keyword_leaves_the_opportunities_the_lists_and_the_climbs(self):
        c, slides = self.ids_and_text([["keyword", fx.ENGLISH[0], "", ""]])
        wins = [r["keyword"] for r in slides["wins"]["rows"]]
        self.assertNotIn(fx.ENGLISH[0], wins)
        self.assertNotIn(fx.ENGLISH[0], [r["label"] for r in slides["climbs"]["rows"]])
        self.assertNotIn(fx.ENGLISH[0], [r["label"] for r in slides["grew"]["keywords"]])
        self.assertNotIn("cannibal", [o["type"] for o in c.opps if o["keyword"] == fx.ENGLISH[0]])

    def test_a_resolved_page_leaves_every_list(self):
        c, slides = self.ids_and_text([["page", "/guide/old-tips/", "", ""]])
        watch = [i["title"] for i in slides["watch"]["items"]]
        self.assertNotIn("/guide/old-tips/", watch)
        pages = [r["page"] for r in slides["pages"]["rows"]]
        self.assertNotIn("/guide/old-tips/", pages)
        c, slides = self.ids_and_text([["page", "/guide/blue-shoes/", "", ""]])
        self.assertNotIn("/guide/blue-shoes/", [r["label"] for r in slides["climbs"]["rows"]])
        self.assertNotIn("/guide/blue-shoes/", [r["label"] for r in slides["grew"]["pages"]])

    def test_a_resolved_gone_page_is_not_counted(self):
        rows = [["page", f"https://example.test/old/p{i}", "", ""] for i in range(3)]
        c, slides = self.ids_and_text(rows)
        self.assertEqual(len(c.gone), 3)
        self.assertNotIn("Old URLs", [i["title"] for i in slides["watch"]["items"]])

    def test_a_resolved_segment_leaves_sections_watching_and_next_month(self):
        c, slides = self.ids_and_text([["segment", "Shop", "", ""]])
        self.assertNotIn("Shop", [r["label"] for r in slides["sections"]["rows"]])
        self.assertNotIn("Shop", [i["title"] for i in slides["watch"]["items"]])
        self.assertNotIn("Shop", [k["label"] for k in slides["next"]["cards"]])

    def test_the_weekly_deck_still_shows_everything(self):
        rep = report(resolved_list=[["page", "https://example.test/old/p0", "", ""]])
        self.assertEqual(len(rep.gone), 6)                   # the report itself is untouched


class RenderingTests(unittest.TestCase):

    def test_every_slide_has_the_footer_without_a_report_name_and_a_slide_number(self):
        prs = build()
        for slide in list(prs.slides)[1:]:
            footer = [s.text_frame.text for s in slide.shapes
                      if s.has_text_frame and "Example · September 2026" in s.text_frame.text]
            self.assertTrue(footer)
        numbers = [s for s in prs.slides[3].shapes if s.has_text_frame and s.text_frame.text.strip().isdigit()
                   or "slidenum" in s._element.xml]
        self.assertTrue(numbers)

    def test_title_slide_wording(self):
        prs = Presentation(io.BytesIO(client_deck.build_monthly(report(), fx.Cfg(), deck_runner.TEMPLATE_DIR)))
        text = "\n".join(texts_of(prs)[:6])
        self.assertIn("EXAMPLE", text)
        self.assertIn("Monthly SEO Report", text)
        self.assertIn("September 2026 \u00b7 Google Search performance, progress and next steps", text)
        layout = "\n".join(s.text_frame.text for s in prs.slides[0].slide_layout.shapes if s.has_text_frame)
        self.assertIn("Prepared by Mehdi Raoof", layout)
        self.assertNotIn("PulseO", layout)
        self.assertNotIn("Generated by", layout)

    def test_kickers_titles_and_the_headline_are_drawn(self):
        text = "\n".join(texts_of(build()))
        for needle in ("HEADLINE", "This month in one line", "SEPTEMBER 2026 VS AUG AND VS SEP 2025",
                       "Scorecard", "Rankings that climbed", "What grew this month",
                       "New searches finding Example", "Our work this month", "13-month trend",
                       "Performance by section", "Quick wins: almost on top", "Pages to update",
                       "New content ideas", "Get ready for October", "What we're watching",
                       "Next month's focus", "Mehdi fills this slide before sending.",
                       "Example ranks higher than it has all year."):
            self.assertIn(needle, text, needle)

    def test_the_trend_is_a_native_line_chart_and_sections_a_bar_chart(self):
        prs = build()
        kinds = [shp.chart.chart_type for slide in prs.slides for shp in slide.shapes if shp.has_chart]
        self.assertEqual(len(kinds), 2)
        xml = "".join(shp.chart._chartSpace.xml for slide in prs.slides for shp in slide.shapes
                      if shp.has_chart)
        self.assertIn("lineChart", xml)
        self.assertIn("barChart", xml)

    def test_no_slide_shows_a_raw_period_an_all_segment_or_a_regex(self):
        text = "\n".join(texts_of(build()))
        self.assertIsNone(re.search(r"\b\d{4}-\d{2}\b", text))
        self.assertNotRegex(text, r"\bAll\b")

    def test_declines_are_slate_not_red_in_the_tables_and_cards(self):
        prs = build()
        reds = 0
        for slide in prs.slides:
            for shp in slide.shapes:
                if shp.has_text_frame:
                    for p in shp.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.color and r.font.color.type is not None and \
                                    str(r.font.color.rgb) == "D64545":
                                reds += 1
        self.assertEqual(reds, 0)

    def test_a_persian_keyword_in_an_english_deck_is_right_aligned_in_vazirmatn(self):
        prs = build()
        found = False
        for slide in prs.slides:
            for shp in slide.shapes:
                if getattr(shp, "has_table", False) and shp.has_table:
                    for row in shp.table.rows:
                        cell = row.cells[0]
                        if fx.PERSIAN[0] in cell.text_frame.text:
                            p = cell.text_frame.paragraphs[0]
                            self.assertEqual(p._p.pPr.get("algn"), "r")
                            self.assertTrue(any(r.font.name == "Vazirmatn" for r in p.runs))
                            found = True
        self.assertTrue(found)

    def test_a_persian_deck_is_built_from_the_persian_template_and_texts(self):
        prs = build("fa")
        self.assertEqual(len(prs.slides), 15)
        text = "\n".join(texts_of(prs))
        self.assertIn("گزارش ماهانه سئو", text)
        self.assertIn("فعلاً اقدامی لازم نیست", text)
        self.assertNotIn("Quick wins", text)

    def test_the_document_properties_have_no_report_data(self):
        core = build().core_properties
        self.assertEqual(core.author, "Mehdi Raoof")
        self.assertIn("Example", core.title)


class TextsTests(unittest.TestCase):

    def test_both_languages_have_the_same_keys(self):
        self.assertEqual(set(client_i18n.EN), set(client_i18n.FA))

    def test_every_placeholder_matches_between_languages(self):
        for key, en in client_i18n.EN.items():
            fa = client_i18n.FA[key]
            if isinstance(en, list):
                self.assertEqual(len(en), len(fa), key)
                for a, b in zip(en, fa):
                    self.assertEqual(set(PLACEHOLDER.findall(a)), set(PLACEHOLDER.findall(b)), key)
            else:
                self.assertEqual(set(PLACEHOLDER.findall(en)), set(PLACEHOLDER.findall(fa)), key)

    def test_persian_content_is_built_too(self):
        c = client("fa")
        slides = c.slides()
        self.assertEqual(len(slides), 14)
        self.assertEqual(dict(slides)["headline"]["headline"], "Example بالاتر از تمام سال گذشته رتبه گرفته است.")


class RunnerTests(unittest.TestCase):

    def test_the_email_is_a_short_personal_note_without_the_tool_name_or_alarms(self):
        rep = report()
        rep.alarms = [{"severity": "high", "area": "Segment clicks", "impact": 5,
                       "message": "Clicks fell 30.0% vs August (100 to 70)."}]
        subject, body = deck_runner.email_text(rep, fx.Cfg(), "en", "September 2026")
        self.assertEqual(subject, "Example SEO Report \u00b7 September 2026")
        lines = [l for l in body.splitlines() if l.strip()]
        self.assertEqual(lines[0], "Hi,")
        self.assertEqual(lines[-1], "Mehdi")
        self.assertIn("Here is the SEO report for Example, September 2026.", body)
        self.assertRegex(body, r"Example got [\d,]+ clicks from Google(, up \d+% on August|, down \d+% on August|, about the same as August)\.")
        self.assertRegex(body, r"It showed up [\d,]+ times in search, and the average position was \d+\.\d( \(\d+\.\d in August\))?\.")
        for banned in ("PulseO", "alarm", "Clicks:", "Impressions:", "Avg. position", "- "):
            self.assertNotIn(banned, subject + body)
        self.assertLessEqual(len(lines), 5)                         # a greeting, a line, the numbers, the sign-off

    def test_the_persian_email_is_a_note_too(self):
        subject, body = deck_runner.email_text(report(), fx.Cfg(), "fa", "\u0633\u067e\u062a\u0627\u0645\u0628\u0631 2026")
        self.assertTrue(subject.startswith("\u06af\u0632\u0627\u0631\u0634 \u0633\u0626\u0648 Example"))
        self.assertEqual(body.splitlines()[0], "\u0633\u0644\u0627\u0645\u060c")
        self.assertEqual([l for l in body.splitlines() if l.strip()][-1], "\u0645\u0647\u062f\u06cc")
        self.assertNotIn("PulseO", subject + body)


class PolishTests(unittest.TestCase):
    """Small fixes to the monthly deck: numbers that fit, gone pages out of the watch list, grammar
    that follows the section name, keywords on lines of their own, and no tool name anywhere."""

    def test_big_numbers_shrink_to_fit_instead_of_wrapping(self):
        self.assertLessEqual(client_deck.fit_size("160,778", 1.73), 30)
        self.assertGreaterEqual(client_deck.fit_size("160,778", 2.34), 36)
        self.assertEqual(client_deck.fit_size("9", 1.0), 40)
        self.assertEqual(client_deck.fit_size("", 1.0), 40)
        self.assertGreaterEqual(client_deck.fit_size("123,456,789,012", 1.0), 22)

    def test_with_five_cards_the_big_number_stays_on_one_line(self):
        rep = report()
        row = rep.rows[(fx.MONTHS[-1], "All")]
        row["Impressions"], row["AIO Impressions"] = 160778, 12345
        prs = Presentation(io.BytesIO(client_deck.build_monthly(rep, fx.Cfg(), deck_runner.TEMPLATE_DIR)))
        sizes = []
        for slide in prs.slides:
            for shp in slide.shapes:
                if shp.has_text_frame and shp.text_frame.text == "160,778":
                    sizes += [r.font.size.pt for p in shp.text_frame.paragraphs for r in p.runs if r.font.size]
        self.assertTrue(sizes)
        self.assertLessEqual(max(sizes), 30)

    def test_gone_pages_are_not_checked_or_watched(self):
        rep = report()
        rep.pages.append(fx._pg("Declining", "/guide/zero/", 0, 30, 0, 500, 0.0, 5.0))      # lost every impression
        for page, kind in (("https://example.test/old/p1", "check_page"),                 # on the gone list
                           ("https://example.test/guide/zero/", "check_page"),
                           ("https://example.test/guide/zero/", "decliner")):
            rep.opps.append(fx._opp(kind, "", page, 300, 5, 12.0, 0.016, 12,
                                    "Impressions down 100% on last month.", "Check this page.", "Guides"))
        c = cd.ClientReport(rep, "en", "Example")
        pages = [o["page"] for o in c.opps]
        self.assertNotIn("https://example.test/old/p1", pages)
        self.assertNotIn("https://example.test/guide/zero/", pages)
        titles = [i["title"] for i in c.watching()["items"]]
        self.assertEqual(titles, ["Shop", "Old URLs", "/guide/old-tips/", "/guide/j/"])      # a real drop still shows
        self.assertTrue(c.vanished("https://example.test/old/p3"))
        self.assertFalse(c.vanished("https://example.test/guide/j/"))

    def test_plural_section_names_take_are_and_grow(self):
        labels = {"Guides": "Programs", "Shop": "Shop", "Blog & other": "Blog & other"}
        c = cd.ClientReport(report(deck_labels=labels), "en", "Example")
        slides = dict(c.slides())
        self.assertIn("Year on year, Programs are up 70%", slides["grew"]["fact"])
        self.assertTrue(any(b.startswith("Programs grow most next month") for b in slides["next"]["bullets"]))
        self.assertTrue(c.plural("Old node pages") and c.plural("Articles") and c.plural("Songs"))
        for singular in ("Shop", "English section", "Blog & other", "Business", "News"):
            self.assertFalse(c.plural(singular), singular)
        single = cd.ClientReport(report(deck_labels={"Guides": "Guide section"}), "en", "Example")
        self.assertIn("Guide section is up 70%", dict(single.slides())["grew"]["fact"])
        self.assertTrue(any(b.startswith("Guide section grows most") for b in dict(single.slides())["next"]["bullets"]))

    def test_persian_sentences_do_not_change_with_the_name(self):
        fa = client_i18n.table("fa")
        for key in ("c_yoy_both", "c_yoy_seg", "c_do_seg", "c_do_seg_slide"):
            self.assertEqual(fa[key], fa[key + "_pl"])

    def test_the_competing_keywords_are_chips_not_part_of_the_sentence(self):
        prs = build()
        slide = [s for s in prs.slides if any("one page per topic" in sh.text_frame.text
                                              for sh in s.shapes if sh.has_text_frame)][0]
        boxes = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        sentence = next(t for t in boxes if "one page per topic" in t)
        for keyword in (fx.ENGLISH[0], fx.ENGLISH[1]):
            self.assertNotIn(keyword, sentence)
            self.assertIn(keyword, boxes)                                        # a box of its own

    def test_a_persian_keyword_chip_is_right_aligned_and_in_vazirmatn(self):
        rep = report()
        for o in rep.opps:
            if o["type"] == "cannibal":
                o["keyword"] = fx.PERSIAN[0]
                break
        prs = Presentation(io.BytesIO(client_deck.build_monthly(rep, fx.Cfg(), deck_runner.TEMPLATE_DIR)))
        found = [sh for s in prs.slides for sh in s.shapes
                 if sh.has_text_frame and sh.text_frame.text == fx.PERSIAN[0]]
        chip = found[-1]                                              # the chip, after the other mentions on earlier slides
        p = chip.text_frame.paragraphs[0]
        self.assertEqual(p._p.pPr.get("algn"), "r")
        self.assertTrue(any(r.font.name == "Vazirmatn" for r in p.runs))

    def test_the_footnote_is_dropped_when_the_page_is_in_the_table_and_kept_when_it_is_not(self):
        rep = report()
        c = cd.ClientReport(rep, "en", "Example")
        self.assertEqual(c.climbs()["footnote"], "")
        for i in range(9):                                         # nine bigger climbs push the page-one move off the table
            rep.keywords.append(fx._kw("Growing", f"filler search {i}", 50, 20, 9000, 8000, 3.0, 20.0 - i))
        crowded = cd.ClientReport(rep, "en", "Example")
        rows = crowded.climbs()
        self.assertEqual(len(rows["rows"]), 7)
        self.assertIn("moved onto page 1 this month", rows["footnote"])

    def test_no_part_of_a_deck_names_the_tool(self):
        import zipfile
        for lang in ("en", "fa"):
            payloads = [client_deck.build_monthly(report(), fx.Cfg(lang), deck_runner.TEMPLATE_DIR)]
            for payload in payloads:
                with zipfile.ZipFile(io.BytesIO(payload)) as z:
                    for part in z.namelist():
                        if part.endswith((".xml", ".rels")):
                            data = z.read(part)
                            self.assertNotIn(b"PulseO", data, f"{lang}: {part}")
                            self.assertNotIn(b"pulse_amber", data, f"{lang}: {part}")
                        self.assertFalse(part.startswith("ppt/media/"), f"{lang}: {part}")   # the mark's picture is gone
        prs = build()
        self.assertEqual([sh.text_frame.text for sh in prs.slides[0].slide_layout.shapes
                          if sh.has_text_frame and "Prepared" in sh.text_frame.text], ["Prepared by Mehdi Raoof"])

class BrandAndAddressTests(unittest.TestCase):
    """Brand searches, addresses and front pages are never opportunities or new keywords."""

    def extra_opps(self):
        return [
            fx._opp("quick_win", "example web", "https://example.test/guide/a/", 9000, 100, 9.0, 0.011, 375,
                    "Ranks on page 1 but below the top 3.", "Improve the page.", "Guides"),
            fx._opp("update_page", "", "https://example.test/", 20000, 100, 6.0, 0.005, 400,
                    "Ranks well, title doesn't win the click.", "Rewrite title + meta.", ""),
            fx._opp("update_page", "", "https://example.test/en", 20000, 100, 6.0, 0.005, 390,
                    "Ranks well.", "Rewrite.", ""),
            fx._opp("cannibal", "example", "https://example.test/guide/b/", 9000, 90, 3.0, 0.01, 380,
                    "5 pages compete for this search.", "Merge into one page.", "Guides"),
            fx._opp("quick_win", "app.example.ir", "https://example.test/guide/c/", 3000, 30, 8.0, 0.01, 300,
                    "Ranks on page 1 but below the top 3.", "Improve the page.", "Guides"),
            fx._opp("quick_win", "plain topic", "https://example.test/", 3000, 30, 8.0, 0.01, 290,
                    "Ranks on page 1 but below the top 3.", "Improve the page.", ""),
        ]

    def test_the_deck_leaves_them_out_even_when_an_old_opportunities_tab_lists_them(self):
        rep = report()
        rep.opps = self.extra_opps() + rep.opps
        rep.brand_terms = ["example"]
        c = cd.ClientReport(rep, "en", "Example")
        slides = dict(c.slides())
        wins = [r["keyword"] for r in slides["wins"]["rows"]]
        for gone in ("example web", "app.example.ir", "plain topic"):
            self.assertNotIn(gone, wins)
        self.assertEqual(wins[0], fx.PERSIAN[0])                              # the real top item stays
        pages = [r["page"] for r in slides["pages"]["rows"]]
        self.assertNotIn("/", pages)
        self.assertNotIn("/en", pages)
        self.assertNotIn("5 pages", slides["ideas"]["merge"]["body"])
        self.assertIn("475 extra clicks", slides["headline"]["cards"][2]["body"])      # the total is unchanged

    def test_new_searches_skip_addresses_and_brand_words(self):
        rep = report()
        rep.brand_terms = ["example"]
        rep.keywords.append(fx._kw("New", "app.example.ir", 3, 0, 5000, 0, 1.2, 0))
        rep.keywords.append(fx._kw("New", "Example web store", 3, 0, 4000, 0, 1.2, 0))
        rep.keywords.append(fx._kw("New", "www site", 3, 0, 3000, 0, 1.2, 0))
        cards = cd.ClientReport(rep, "en", "Example").new_searches()["cards"]
        self.assertEqual([c["pill"] for c in cards], ["276 impressions", "120 impressions", "90 impressions"])

    def test_addresses_and_front_pages_by_rule(self):
        import opportunities as op
        for query in ("app.example.ir", "example.com", "www example", "http://x", "shop.site.ir/path"):
            self.assertTrue(op.looks_like_domain(query), query)
        for query in ("blue running shoes", "what is a trail shoe", "mr smith", "کلمه کلیدی", "a.b"):
            self.assertFalse(op.looks_like_domain(query), query)
        for url in ("https://x.test/", "https://x.test", "https://x.test/en", "https://x.test/en/",
                    "https://x.test/fa/", "https://x.test/pt-br"):
            self.assertTrue(op.is_homepage(url), url)
        for url in ("https://x.test/article/a", "https://x.test/about", "https://x.test/en/guide/"):
            self.assertFalse(op.is_homepage(url), url)


class PersianPathTests(unittest.TestCase):
    """A path or address always reads left to right, whatever its slug is written in."""

    SLUG = "کبوتر"

    def cells(self, prs, needle):
        found = []
        for slide in prs.slides:
            for shp in slide.shapes:
                if getattr(shp, "has_table", False) and shp.has_table:
                    for row in shp.table.rows:
                        for cell in row.cells:
                            if needle in cell.text_frame.text:
                                found.append(cell)
        return found

    def deck(self, lang="en"):
        rep = report(slugs=True)
        return Presentation(io.BytesIO(client_deck.build_monthly(rep, fx.Cfg(lang), deck_runner.TEMPLATE_DIR)))

    def test_a_persian_slug_is_left_to_right_left_aligned_marked_and_in_vazirmatn(self):
        for lang in ("en", "fa"):
            cells = self.cells(self.deck(lang), self.SLUG)
            self.assertTrue(cells, lang)
            for cell in cells:
                p = cell.text_frame.paragraphs[0]
                self.assertEqual(p._p.pPr.get("algn"), "l", lang)
                self.assertNotEqual(p._p.pPr.get("rtl"), "1", lang)
                self.assertTrue(cell.text_frame.text.startswith("‎/article/"), lang)
                self.assertTrue(cell.text_frame.text.endswith("‎"), lang)
                self.assertTrue(all(r.font.name == "Vazirmatn" for r in p.runs), lang)

    def test_plain_persian_keywords_stay_right_aligned(self):
        cells = self.cells(self.deck(), fx.PERSIAN[0])
        self.assertTrue(cells)
        self.assertTrue(all(c.text_frame.paragraphs[0]._p.pPr.get("algn") == "r" for c in cells))

    def test_long_slugs_are_cut_to_one_line(self):
        c = cd.ClientReport(report(slugs=True), "en", "Example")
        labels = [r["label"] for r in c.climbs()["rows"]]
        long = [l for l in labels if l.startswith("/guide/a-very-long")]
        self.assertEqual(len(long), 1)
        self.assertEqual(len(long[0]), cd.MAX_LABEL)
        self.assertTrue(long[0].endswith("…"))
        self.assertTrue(all(len(l) <= cd.MAX_LABEL for l in labels))
        self.assertEqual(c.short("short"), "short")

    def test_paths_inside_a_sentence_keep_their_order(self):
        import decks
        self.assertEqual(decks.lrm_urls("and /a/b gained 5"), "and ‎/a/b‎ gained 5")
        self.assertEqual(decks.lrm_urls("https://x.test/p is up"), "‎https://x.test/p‎ is up")
        self.assertEqual(decks.lrm_urls("and/or 50% / more"), "and/or 50% / more")
        self.assertEqual(decks.lrm_urls("‎/a/b‎"), "‎/a/b‎")             # not wrapped twice
        self.assertTrue(decks.is_url_like("/a"))
        self.assertTrue(decks.is_url_like("‎/a‎"))
        self.assertTrue(decks.is_url_like("https://x"))
        self.assertFalse(decks.is_url_like("a /b"))
        self.assertFalse(decks.is_url_like("کلمه"))

    def test_a_path_in_a_text_box_title_is_left_aligned(self):
        prs = Presentation(io.BytesIO(client_deck.build_monthly(
            report(), fx.Cfg("fa"), deck_runner.TEMPLATE_DIR)))
        titles = [s for slide in prs.slides for s in slide.shapes
                  if s.has_text_frame and s.text_frame.text.startswith("‎/guide/")]
        self.assertTrue(titles)
        for shp in titles:
            p = shp.text_frame.paragraphs[0]
            self.assertEqual(p._p.pPr.get("algn"), "l")
            self.assertNotEqual(p._p.pPr.get("rtl"), "1")


class HeadlineRepeatTests(unittest.TestCase):

    def c(self, mutate=None):
        rep = report()
        if mutate:
            mutate(rep)
        return cd.ClientReport(rep, "en", "Example")

    def test_the_click_rate_headline_is_not_repeated_in_the_supporting_sentence(self):
        def mutate(rep):
            row = rep.rows[(fx.MONTHS[-1], "All")]
            row["Avg. Position"], row["Clicks YoY %"] = 8.0, 5.0
        d = self.c(mutate).headline()
        self.assertEqual(d["rule"], "click rate")
        self.assertEqual(d["headline"], "More searchers clicked through.")
        self.assertNotIn("clicked through", d["support"])
        self.assertIn("CTR is 3.61%", d["support"])

    def test_the_section_headline_is_not_repeated_either(self):
        def mutate(rep):
            row = rep.rows[(fx.MONTHS[-1], "All")]
            row["Avg. Position"], row["Clicks YoY %"], row["CTR (%)"] = 8.0, 5.0, 1.0
        d = self.c(mutate).headline()
        self.assertEqual(d["rule"], "section")
        self.assertNotIn("led the growth", d["support"])

    def test_the_position_headline_keeps_both_numbers(self):
        d = self.c().headline()
        self.assertEqual(d["rule"], "position")
        self.assertIn("Average position improved to 6.2 and more searchers clicked through (CTR 3.61%)", d["support"])

    def test_a_tie_with_the_best_earlier_month_still_counts_as_the_best(self):
        def mutate(rep):
            rep.rows[(fx.MONTHS[-1], "All")]["Avg. Position"] = 6.4          # the best before was 6.4
        self.assertEqual(self.c(mutate).headline()["rule"], "position")

    def test_which_rule_won_is_logged_without_the_figures(self):
        c = self.c()
        c.slides()
        self.assertEqual(len(c.decisions), 1)
        line = c.decisions[0]
        self.assertTrue(line.startswith("headline rule: position."))
        self.assertIn("Best position of the last 12 months: yes", line)
        self.assertIn("gap -0.2 steps", line)
        for figure in ("6.2", "6.4", "1300", "36.8"):
            self.assertNotIn(figure, line)

        def worse(rep):
            rep.rows[(fx.MONTHS[-1], "All")]["Avg. Position"] = 6.9
        c2 = self.c(worse)
        c2.slides()
        self.assertIn("Best position of the last 12 months: no", c2.decisions[0])
        self.assertIn("gap +0.5 steps", c2.decisions[0])
        c2.slides()
        self.assertEqual(len(c2.decisions), 1)                               # not piled up by a second call

    def test_the_build_writes_the_decision_to_the_log(self):
        lines = []
        client_deck.build_monthly(report(), fx.Cfg(), deck_runner.TEMPLATE_DIR, log=lines.append)
        self.assertEqual(len(lines), 1)
        self.assertTrue(lines[0].startswith("decks: monthly headline rule: position."))


class DeckLabelTests(unittest.TestCase):
    LABELS = {"Guides": "Guides section", "Shop": "Shop", "Blog & other": "Everything else"}

    def c(self, **kw):
        return cd.ClientReport(report(deck_labels=self.LABELS, **kw), "en", "Example")

    def test_the_deck_uses_the_display_names_everywhere(self):
        slides = dict(self.c().slides())
        self.assertIn("Guides section led the growth", slides["headline"]["support"])
        self.assertEqual([r["label"] for r in slides["sections"]["rows"]],
                         ["Guides section", "Everything else", "Shop"])
        self.assertEqual([k["label"] for k in slides["next"]["cards"]],
                         ["Whole site", "Guides section", "Everything else", "Shop"])
        self.assertIn("Guides section grows most next month", slides["next"]["bullets"][0])
        self.assertIn("Year on year, Guides section is up 70%", slides["grew"]["fact"])
        self.assertEqual(slides["watch"]["items"][0]["title"], "Shop")

    def test_without_a_display_name_the_label_is_used(self):
        slides = dict(client().slides())
        self.assertEqual([r["label"] for r in slides["sections"]["rows"]], ["Guides", "Blog & other", "Shop"])

    def test_the_resolved_list_still_uses_the_segment_label(self):
        c = cd.ClientReport(report(deck_labels=self.LABELS, resolved_list=[["segment", "Guides", "", ""]]),
                            "en", "Example")
        self.assertNotIn("Guides section", [r["label"] for r in dict(c.slides())["sections"]["rows"]])

    def test_page_types_and_the_rest_segment_have_display_names_in_the_site_settings(self):
        import json
        import sites
        entry = {"name": "site-a", "site_url": "sc-domain:a.test", "sheet_id": "m", "first_match_wins": True,
                 "other_label": "Other", "other_deck_label": "Blog & other",
                 "page_types": [{"name": "en", "regex": "/en/", "deck_label": "English section"},
                                {"name": "Song", "regex": "/song/"}]}
        site = sites.parse_sites_json(json.dumps([entry]))[0]
        self.assertEqual(site.deck_label("en"), "English section")
        self.assertEqual(site.deck_label("Song"), "Song")                         # no deck_label: the label
        self.assertEqual(site.deck_label("Other"), "Blog & other")
        self.assertEqual(site.deck_label("All"), "All")
        self.assertEqual([t["name"] for t in site.page_types], ["en", "Song"])    # the segment names did not change
        for bad in ({"deck_label": "<x>"}, {"deck_label": 5}, {"deck_label": ""}):
            e = dict(entry, page_types=[dict(entry["page_types"][0], **bad)])
            with self.assertRaises(sites.SitesConfigError) as cm:
                sites.parse_sites_json(json.dumps([e]))
            self.assertNotIn("<x>", str(cm.exception))
        with self.assertRaises(sites.SitesConfigError):
            sites.parse_sites_json(json.dumps([dict(entry, other_deck_label="<x>")]))
        extra = dict(entry["page_types"][0], colour="red")
        with self.assertRaises(sites.SitesConfigError):
            sites.parse_sites_json(json.dumps([dict(entry, page_types=[extra])]))


class WatchingTests(unittest.TestCase):

    def test_a_segment_gets_a_calm_sentence_that_points_at_the_plan(self):
        d = dict(client().slides())["watch"]
        self.assertEqual(d["items"][0]["title"], "Shop")
        self.assertEqual(d["items"][0]["body"],
                         "Down 35% on last year; the update plan on slides 10 to 12 targets this.")

    def test_this_months_rise_is_mentioned_calmly(self):
        rep = report()
        rep.rows[(fx.MONTHS[-1], "Shop")]["Clicks MoM %"] = 4.2
        body = dict(cd.ClientReport(rep, "en", "Example").slides())["watch"]["items"][0]["body"]
        self.assertEqual(body, "Down 35% on last year, but up 4% this month; the update plan on slides 10 to "
                               "12 targets this.")

    def test_the_plan_sentence_follows_the_slides_that_exist(self):
        c = client()
        self.assertEqual(c.plan_sentence({"wins": 10, "pages": 11, "ideas": 12}),
                         "the update plan on slides 10 to 12 targets this")
        self.assertEqual(c.plan_sentence({"wins": 10}), "the update plan on slide 10 targets this")
        self.assertEqual(c.plan_sentence({"pages": 7, "ideas": 8}), "the update plan on slides 7 to 8 targets this")
        self.assertEqual(c.plan_sentence({}), "we are keeping an eye on it")
        self.assertEqual(c.plan_sentence(None), "we are keeping an eye on it")

    def test_a_decliner_needs_at_least_50_clicks_last_month(self):
        rep = report()
        for r in rep.pages:
            if r["Type"] == "Declining" and r["Path"] == "/guide/old-tips/":
                r["Clicks prev"] = 24
        slides = dict(cd.ClientReport(rep, "en", "Example").slides())
        self.assertNotIn("/guide/old-tips/", [i["title"] for i in slides["watch"]["items"]])
        self.assertNotIn("/guide/old-tips/", [r["page"] for r in slides["pages"]["rows"]])
        slides = dict(client().slides())                          # 137 clicks last month: it stays
        self.assertIn("/guide/old-tips/", [i["title"] for i in slides["watch"]["items"]])
        self.assertIn("/guide/old-tips/", [r["page"] for r in slides["pages"]["rows"]])

    def test_the_watching_slide_is_drawn_larger(self):
        prs = Presentation(io.BytesIO(client_deck.build_monthly(report(), fx.Cfg(), deck_runner.TEMPLATE_DIR)))
        slide = prs.slides[len(prs.slides) - 2]
        sizes = {r.font.size.pt for shp in slide.shapes if shp.has_text_frame
                 for p in shp.text_frame.paragraphs for r in p.runs if r.font.size}
        self.assertTrue({19.0, 16.0} <= sizes)


if __name__ == "__main__":
    unittest.main()
