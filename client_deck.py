"""
PulseO: the client-ready monthly deck, drawn (v2.4)
===================================================
Draws the slides of client_data.ClientReport with python-pptx on the layouts of the PulseO template
(footer, rule, PulseO mark and slide number come from the layouts). White slides, navy text, an
amber uppercase kicker above a 28pt title, rounded tinted cards, green for gains and slate grey for
declines. Slides without data are not drawn.
"""

import io
import os

from pptx import Presentation
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Pt

import charts
import client_data as cd
import deck_data as dd
import decks
import i18n
from decks import (AMBER, GREEN, INK, MUTED, PERIWINKLE, RED, SLATE, TINT, Cell, card, inch,
                   mirror_x, rgb, shape, text_box, write_paragraph)

GREEN_TINT, AMBER_TINT, PERI_TINT = "E3F4EC", "FDF1DE", "ECEEFB"
BROWN = "9A5B00"
LEFT, WIDTH = 0.6, 12.13
TONES = {"green": (GREEN_TINT, GREEN), "periwinkle": (PERI_TINT, PERIWINKLE),
         "amber": (AMBER_TINT, AMBER), "tint": (TINT, None), "grey": (TINT, SLATE), "red": (TINT, RED)}
TEXT_COLOR = {cd.GOOD: GREEN, cd.GREY: SLATE, cd.RED: RED}


def tone_color(tone: str) -> str:
    return TEXT_COLOR.get(tone, SLATE)


# --------------------------------------------------------------------------- #
# Pieces every slide uses
# --------------------------------------------------------------------------- #

def kicker(slide, ctx, text: str):
    """The small amber uppercase line above the title."""
    box = text_box(slide, mirror_x(ctx, LEFT, 9.0), 0.36, 9.0, 0.3,
                   [text if ctx.rtl else text.upper()], ctx, size=11, color=AMBER, bold=True,
                   align="r" if ctx.rtl else "l")
    if not ctx.rtl:                                        # letter spacing, like the sample
        for run in box.text_frame.paragraphs[0].runs:
            run._r.get_or_add_rPr().set("spc", "180")
    return box


def content_slide(ctx, kicker_text: str, title: str):
    slide, roles = decks.new_slide(ctx, "Alarms")
    title_ph = roles["title"]
    title_ph.left, title_ph.top = Emu(inch(LEFT)), Emu(inch(0.62))       # all four: a placeholder
    title_ph.width, title_ph.height = Emu(inch(WIDTH)), Emu(inch(0.85))  # has no position of its own
    decks.fill_placeholder(title_ph, [title if len(title) <= 70 else title[:69] + "…"], ctx,
                           anchor=MSO_ANCHOR.MIDDLE, size=28, bold=True)
    kicker(slide, ctx, kicker_text)
    return slide


def bar_card(slide, ctx, x, y, w, h, tone: str, radius: float = 0.05):
    """A tinted rounded card with a thin coloured bar on its leading edge."""
    fill, bar = TONES[tone]
    card(slide, x, y, w, h, fill, radius)
    if bar:
        bx = x if not ctx.rtl else x + w - 0.1
        shape(slide, MSO_SHAPE.RECTANGLE, bx, y + 0.02, 0.1, h - 0.04, bar)


def para(box, ctx, items: list, first: bool = True):
    """Adds paragraphs (text, {format}) to a text box after its first, empty one."""
    for i, (text, fmt) in enumerate(items):
        p = box.text_frame.paragraphs[0] if (i == 0 and first) else box.text_frame.add_paragraph()
        write_paragraph(p, text, ctx, **fmt)
    return box


def align_for(ctx, text: str) -> str:
    """Persian text sits on the right, other text follows the deck direction."""
    return "r" if (ctx.rtl or decks.dd_has_arabic(text)) else "l"


def foot_note(slide, ctx, text: str, y: float = 6.35):
    if text:
        text_box(slide, mirror_x(ctx, LEFT, WIDTH), y, WIDTH, 0.4, [text], ctx, size=12, color=MUTED,
                 align="r" if ctx.rtl else "l")


def intro(slide, ctx, text: str, y: float = 1.55):
    text_box(slide, mirror_x(ctx, LEFT, WIDTH), y, WIDTH, 0.5, [text], ctx, size=15, color=SLATE,
             align="r" if ctx.rtl else "l")


def fmt_text(tone: str) -> str:
    return tone_color(tone)


# --------------------------------------------------------------------------- #
# Slides
# --------------------------------------------------------------------------- #

def slide_title(ctx, c: cd.ClientReport):
    T = c.T
    slide, ph = decks.new_slide(ctx, "Title")
    decks.fill_placeholder(ph["site"], [c.label if ctx.rtl else c.label.upper()], ctx)
    decks.fill_placeholder(ph["title"], [T["monthly_report"]], ctx)
    decks.fill_placeholder(ph["subtitle"], [T["c_subtitle"].format(period=ctx.period_title)], ctx)
    if ctx.logo:
        try:
            slide.shapes.add_picture(io.BytesIO(ctx.logo), Emu(inch(10.8 if not ctx.rtl else 0.6)),
                                     Emu(inch(0.7)), height=Emu(inch(0.9)))
        except Exception:                                  # a bad logo never breaks the deck
            pass


def slide_headline(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_headline"], T["c_t_headline"])
    card(slide, mirror_x(ctx, LEFT, WIDTH), 1.65, WIDTH, 2.15, INK, 0.04)
    box = text_box(slide, mirror_x(ctx, 1.0, WIDTH - 0.8), 1.85, WIDTH - 0.8, 1.8,
                   [""], ctx, size=24, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE,
                   align="r" if ctx.rtl else "l")
    write_paragraph(box.text_frame.paragraphs[0], "", ctx, size=24,
                    runs=[(d["headline"] + " ", {"bold": True, "color": AMBER}),
                          (d["support"], {"color": "FFFFFF"})])
    cards = d["cards"]
    if cards:
        gap = 0.28
        w = (WIDTH - gap * (len(cards) - 1)) / len(cards)
        for i, k in enumerate(cards):
            x = mirror_x(ctx, LEFT + i * (w + gap), w)
            bar_card(slide, ctx, x, 4.1, w, 2.45, k["tone"])
            text_box(slide, x + 0.3, 4.3, w - 0.5, 0.6, [k["title"]], ctx, size=19, color=INK, bold=True,
                     align="r" if ctx.rtl else "l")
            text_box(slide, x + 0.3, 4.95, w - 0.5, 1.5, [k["body"]], ctx, size=15, color=SLATE,
                     align="r" if ctx.rtl else "l")


def slide_scorecard(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, d["kicker"], T["c_t_score"])
    cards, gap = d["cards"], 0.25
    w = (WIDTH - gap * (len(cards) - 1)) / len(cards)
    for i, k in enumerate(cards):
        x = mirror_x(ctx, LEFT + i * (w + gap), w)
        card(slide, x, 1.7, w, 4.3, TINT, 0.06)
        al = "r" if ctx.rtl else "l"
        text_box(slide, x + 0.25, 1.95, w - 0.5, 0.4, [k["label"]], ctx, size=15, color=SLATE, align=al)
        text_box(slide, x + 0.25, 2.5, w - 0.5, 0.95, [k["value"]], ctx, size=40, color=INK, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE, align=al)
        box = text_box(slide, x + 0.25, 3.6, w - 0.5, 1.0, [""], ctx, size=13, align=al)
        for j, line in enumerate(k["lines"]):
            p = box.text_frame.paragraphs[0] if j == 0 else box.text_frame.add_paragraph()
            p.space_after = Pt(6)
            write_paragraph(p, "", ctx, size=14, align=al,
                            runs=[(line["label"] + "  ", {"color": MUTED, "size": 13}),
                                  (line["text"], {"color": tone_color(line["tone"]), "bold": True})])
        text_box(slide, x + 0.25, 4.75, w - 0.5, 1.1, [k["note"]], ctx, size=13, color=SLATE, align=al)
    foot_note(slide, ctx, d["footnote"], y=6.2)


def slide_climbs(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_progress"], T["c_t_climb"])
    table_x, table_w = LEFT, WIDTH
    if d["card"]:
        k = d["card"]
        x = mirror_x(ctx, LEFT, 4.1)
        card(slide, x, 1.65, 4.1, 4.65, GREEN_TINT, 0.05)
        al = "r" if ctx.rtl else "l"
        text_box(slide, x + 0.3, 1.9, 3.5, 0.4, [k["title"]], ctx, size=17, color=SLATE, align=al)
        box = text_box(slide, x + 0.3, 2.45, 3.6, 1.0, [""], ctx, size=48, anchor=MSO_ANCHOR.MIDDLE,
                       align=al)
        write_paragraph(box.text_frame.paragraphs[0], "", ctx, size=48, align=al,
                        runs=[(k["from"], {"color": MUTED}), ("  →  ", {"color": MUTED, "size": 36}),
                              (k["to"], {"color": GREEN, "bold": True, "size": 52})])
        text_box(slide, x + 0.3, 3.55, 3.5, 0.4, [k["span"]], ctx, size=13, color=MUTED, align=al)
        text_box(slide, x + 0.3, 4.15, 3.5, 2.0, [k["body"]], ctx, size=15, color=INK, align=al)
        table_w = WIDTH - 4.1 - 0.25
        table_x = mirror_x(ctx, LEFT + 4.1 + 0.25, table_w)
    widths = [table_w - 3.3, 1.0, 1.0, 1.3]
    rows = [[Cell(r["label"], "query"), Cell(r["prev"], "num"), Cell(r["now"], "num", GREEN, True),
             Cell(r["change"], "change", GREEN, True)] for r in d["rows"]]
    decks.add_table(slide, ctx, table_x, 1.65, widths, d["header"], rows, row_height=0.58, size=14)
    foot_note(slide, ctx, d["footnote"], y=6.45)


def slide_grew(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_progress"], T["c_t_grew"])
    sides = [(LEFT, d["left_title"], d["keywords"], T["col_query"]),
             (LEFT + 6.3, d["right_title"], d["pages"], T["col_page"])]
    if ctx.rtl:
        sides.reverse()
    for x, heading, rows, col in sides:
        text_box(slide, x, 1.6, 5.83, 0.4, [heading], ctx, size=17, color=INK, bold=True,
                 align="r" if ctx.rtl else "l")
        if rows:
            cells = [[Cell(r["label"], "query"), Cell(r["clicks"], "num"),
                      Cell(r["change"], "change", GREEN, True)] for r in rows]
            decks.add_table(slide, ctx, x, 2.1, [3.6, 1.1, 1.13], [col, T["col_clicks"], T["col_change"]],
                            cells, row_height=0.52, size=13)
    if d["fact"]:
        bar_card(slide, ctx, mirror_x(ctx, LEFT, WIDTH), 5.45, WIDTH, 1.15, "green")
        box = text_box(slide, mirror_x(ctx, LEFT + 0.35, WIDTH - 0.7), 5.55, WIDTH - 0.7, 0.95, [""], ctx,
                       size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE, align="r" if ctx.rtl else "l")
        write_paragraph(box.text_frame.paragraphs[0], "", ctx, size=15,
                        runs=[(d["fact_title"] + " ", {"bold": True, "color": GREEN}),
                              (d["fact"], {"color": INK})])


def slide_new(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_progress"], T["c_t_new"].format(site=c.label))
    intro(slide, ctx, d["intro"])
    for i, k in enumerate(d["cards"]):
        y = 2.35 + i * 1.45
        card(slide, mirror_x(ctx, LEFT, WIDTH), y, WIDTH, 1.2, TINT, 0.06)
        text_box(slide, mirror_x(ctx, 0.95, 6.0), y + 0.1, 6.0, 1.0, [k["keyword"]], ctx, size=26,
                 color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=align_for(ctx, k["keyword"]))
        pill_x = mirror_x(ctx, 7.35, 2.6)
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, pill_x, y + 0.38, 2.6, 0.46, AMBER_TINT, 0.3)
        text_box(slide, pill_x, y + 0.38, 2.6, 0.46, [k["pill"]], ctx, size=14, color=BROWN, bold=True,
                 align="ctr", anchor=MSO_ANCHOR.MIDDLE)
        if k["position"]:
            text_box(slide, mirror_x(ctx, 10.3, 2.2), y + 0.3, 2.2, 0.6, [k["position"]], ctx, size=15,
                     color=SLATE, anchor=MSO_ANCHOR.MIDDLE, align="r" if ctx.rtl else "l")


def slide_work(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_did"], T["c_t_work"])
    box = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, mirror_x(ctx, LEFT, WIDTH), 1.65, WIDTH, 4.7,
                AMBER_TINT, 0.03)
    box.line.color.rgb = rgb(AMBER)
    box.line.width = Pt(1.5)
    box.line.dash_style = MSO_LINE.DASH
    text = text_box(slide, mirror_x(ctx, 1.0, WIDTH - 0.8), 1.95, WIDTH - 0.8, 4.0,
                    [(d["fill"], {"bold": True, "color": BROWN, "size": 20})], ctx, size=20,
                    color=SLATE, align="r" if ctx.rtl else "l")
    para(text, ctx, [("", {"size": 12})] + [("• " + p, {"size": 20, "color": SLATE})
                                              for p in d["prompts"]], first=False)


def slide_trend(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_big"], d["title"])
    chart_x = 0.6 if not ctx.rtl else 4.3
    charts.line_chart(slide, (inch(chart_x), inch(1.5), inch(8.4), inch(5.0)), d["labels"],
                      T["c_chart_clicks"], d["values"], AMBER, d["chart_title"], ctx.font)
    x = mirror_x(ctx, 9.2, 3.53)
    for i, n in enumerate(d["notes"]):
        y = 1.65 + i * 1.3
        bar_card(slide, ctx, x, y, 3.53, 1.18, n["tone"] if n["tone"] != "grey" else "grey", 0.05)
        text_box(slide, x + 0.25, y + 0.1, 3.1, 0.35, [n["title"]], ctx, size=14, color=INK, bold=True,
                 align="r" if ctx.rtl else "l")
        text_box(slide, x + 0.25, y + 0.45, 3.1, 0.7, [n["body"]], ctx, size=12, color=SLATE,
                 align="r" if ctx.rtl else "l")


def slide_sections(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_types"], T["c_t_sections"])
    chart_x = 0.6 if not ctx.rtl else 6.7
    charts.bar_chart(slide, (inch(chart_x), inch(1.5), inch(6.0), inch(4.8)),
                     [r["label"] for r in d["rows"]], T["series_clicks"],
                     [r["clicks"] for r in d["rows"]], PERIWINKLE, ctx.font, d["chart_title"])
    rows = [[Cell(r["label"], "query"), Cell(r["mom"], "change", tone_color(r["mom_tone"]), True),
             Cell(r["yoy"], "change", tone_color(r["yoy_tone"]), True)] for r in d["rows"]]
    decks.add_table(slide, ctx, 6.95 if not ctx.rtl else 0.6, 1.65, [2.5, 1.5, 1.6], d["header"], rows,
                    row_height=0.55, size=13)
    foot_note(slide, ctx, d["footnote"], y=6.25)


def slide_wins(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_opp"], T["c_t_wins"])
    intro(slide, ctx, d["intro"])
    rows = [[Cell(r["keyword"], "query"), Cell(r["searches"], "num"), Cell(r["position"], "num"),
             Cell(r["gain"], "change", GREEN, True)] for r in d["rows"]]
    decks.add_table(slide, ctx, mirror_x(ctx, LEFT, WIDTH), 2.2, [4.9, 2.4, 2.3, 2.53], d["header"], rows,
                    row_height=0.58, size=14)
    card(slide, mirror_x(ctx, LEFT, WIDTH), 5.75, WIDTH, 0.85, AMBER_TINT, 0.08)
    box = text_box(slide, mirror_x(ctx, LEFT + 0.35, WIDTH - 0.7), 5.8, WIDTH - 0.7, 0.75, [""], ctx,
                   size=15, anchor=MSO_ANCHOR.MIDDLE, align="r" if ctx.rtl else "l")
    write_paragraph(box.text_frame.paragraphs[0], "", ctx, size=15,
                    runs=[(d["card"], {"bold": True, "color": BROWN}),
                          (d["card_note"], {"color": INK})])


def slide_pages(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_opp"], T["c_t_pages"])
    intro(slide, ctx, d["intro"])
    rows = [[Cell(r["page"], "page"), Cell(r["impr"], "num"), Cell(r["position"], "num"),
             Cell(r["ctr"], "num"), Cell(r["text"], "text")] for r in d["rows"]]
    decks.add_table(slide, ctx, mirror_x(ctx, LEFT, WIDTH), 2.1, [3.3, 1.55, 1.15, 1.05, 5.08],
                    d["header"], rows, row_height=0.78, size=12)


def slide_ideas(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_opp"], T["c_t_ideas"])
    if d["cards"]:
        intro(slide, ctx, d["intro"])
    gap, w = 0.22, (WIDTH - 0.44) / 3
    for i, k in enumerate(d["cards"][:6]):
        col, row = i % 3, i // 3
        x = mirror_x(ctx, LEFT + col * (w + gap), w)
        y = 2.15 + row * 1.35
        card(slide, x, y, w, 1.15, PERI_TINT, 0.08)
        text_box(slide, x + 0.25, y + 0.12, w - 0.5, 0.6, [k["keyword"]], ctx, size=22, color=INK,
                 bold=True, anchor=MSO_ANCHOR.MIDDLE, align=align_for(ctx, k["keyword"]))
        text_box(slide, x + 0.25, y + 0.75, w - 0.5, 0.3, [k["sub"]], ctx, size=12, color=SLATE,
                 align=align_for(ctx, k["keyword"]))
    if d["merge"]:
        top = 5.0 if d["cards"] else 1.75
        card(slide, mirror_x(ctx, LEFT, WIDTH), top, WIDTH, 1.5, TINT, 0.05)
        box = text_box(slide, mirror_x(ctx, LEFT + 0.35, WIDTH - 0.7), top + 0.1, WIDTH - 0.7, 1.3, [""],
                       ctx, size=15, anchor=MSO_ANCHOR.MIDDLE, align="r" if ctx.rtl else "l")
        write_paragraph(box.text_frame.paragraphs[0], "", ctx, size=15,
                        runs=[(d["merge"]["title"] + " ", {"bold": True, "color": INK}),
                              (d["merge"]["body"], {"color": SLATE})])


def slide_next(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_season"], d["title"])
    intro(slide, ctx, d["intro"])
    cards = d["cards"]
    gap = 0.22
    w = (WIDTH - gap * (len(cards) - 1)) / len(cards)
    for i, k in enumerate(cards):
        x = mirror_x(ctx, LEFT + i * (w + gap), w)
        card(slide, x, 2.1, w, 2.35, GREEN_TINT if k["tone"] == cd.GOOD else TINT, 0.07)
        al = "r" if ctx.rtl else "l"
        text_box(slide, x + 0.22, 2.3, w - 0.4, 0.4, [k["label"]], ctx, size=14, color=SLATE, align=al)
        text_box(slide, x + 0.22, 2.8, w - 0.4, 0.8, [k["value"]], ctx, size=34, bold=True,
                 color=GREEN if k["tone"] == cd.GOOD else SLATE, anchor=MSO_ANCHOR.MIDDLE, align=al)
        lines = [k["word"]] + ([k["sub"]] if k["sub"] else [])
        text_box(slide, x + 0.22, 3.7, w - 0.4, 0.7, lines, ctx, size=12, color=INK, align=al)
    if d["bullets"]:
        card(slide, mirror_x(ctx, LEFT, WIDTH), 4.75, WIDTH, 1.85, TINT, 0.04)
        box = text_box(slide, mirror_x(ctx, LEFT + 0.35, WIDTH - 0.7), 4.85, WIDTH - 0.7, 1.7,
                       [(d["do_title"], {"bold": True, "color": INK, "size": 16})], ctx, size=15,
                       color=SLATE, align="r" if ctx.rtl else "l")
        para(box, ctx, [("• " + b, {"size": 15, "color": SLATE}) for b in d["bullets"][:4]],
             first=False)


def slide_watch(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_watch"], T["c_t_watch"])
    for i, item in enumerate(d["items"]):
        y = 1.65 + i * 1.02
        card(slide, mirror_x(ctx, LEFT, WIDTH), y, WIDTH, 0.88, TINT, 0.08)
        shape(slide, MSO_SHAPE.OVAL, mirror_x(ctx, 0.9, 0.4), y + 0.24, 0.4, 0.4, AMBER)
        text_box(slide, mirror_x(ctx, 1.55, 3.4), y + 0.08, 3.4, 0.72, [item["title"]], ctx, size=17,
                 color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=align_for(ctx, item["title"]))
        text_box(slide, mirror_x(ctx, 5.1, 7.4), y + 0.08, 7.4, 0.72, [item["body"]], ctx, size=14,
                 color=SLATE, anchor=MSO_ANCHOR.MIDDLE, align="r" if ctx.rtl else "l")


def slide_focus(ctx, c, d):
    T = c.T
    slide = content_slide(ctx, T["c_k_plan"], T["c_t_focus"])
    w, h, gap = 5.95, 2.3, 0.23
    for i, k in enumerate(d["cards"]):
        col, row = i % 2, i // 2
        x = mirror_x(ctx, LEFT + col * (w + gap), w)
        y = 1.65 + row * (h + gap)
        card(slide, x, y, w, h, AMBER_TINT if k["tone"] == "amber" else TINT, 0.05)
        cx = x + (0.3 if not ctx.rtl else w - 1.0)
        shape(slide, MSO_SHAPE.OVAL, cx, y + 0.3, 0.7, 0.7, INK)
        text_box(slide, cx, y + 0.3, 0.7, 0.7, [str(i + 1)], ctx, size=22, color=AMBER, bold=True,
                 align="ctr", anchor=MSO_ANCHOR.MIDDLE)
        tx = x + (1.2 if not ctx.rtl else 0.3)
        text_box(slide, tx, y + 0.3, w - 1.5, 0.6, [k["title"]], ctx, size=21, color=INK, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE, align="r" if ctx.rtl else "l")
        text_box(slide, tx, y + 1.0, w - 1.5, 1.2, [k["body"]], ctx, size=15, color=SLATE,
                 align="r" if ctx.rtl else "l")


RENDERERS = {"headline": slide_headline, "scorecard": slide_scorecard, "climbs": slide_climbs,
             "grew": slide_grew, "new": slide_new, "work": slide_work, "trend": slide_trend,
             "sections": slide_sections, "wins": slide_wins, "pages": slide_pages,
             "ideas": slide_ideas, "next": slide_next, "watch": slide_watch, "focus": slide_focus}


# --------------------------------------------------------------------------- #
# Building the deck
# --------------------------------------------------------------------------- #

def _set_prepared(layouts, text: str):
    """The title layout's static line at the bottom reads 'Prepared by ... built with PulseO'."""
    layout = layouts["Title"]
    for shp in layout.shapes:
        if shp.has_text_frame and "PulseO" in shp.text_frame.text:
            p = shp.text_frame.paragraphs[0]
            runs = p.runs
            if runs:
                runs[0].text = text
                for extra in runs[1:]:
                    extra._r.getparent().remove(extra._r)
            return


def build_monthly(report: dd.Report, cfg, template_dir: str, logo: bytes = None) -> bytes:
    """The client-ready monthly deck as .pptx bytes. cfg is the site's DeckConfig."""
    lang = cfg.language
    path = os.path.join(template_dir, "gsc_pulse_template_fa.pptx" if lang == "fa"
                        else "gsc_pulse_template.pptx")
    prs, layouts = decks.open_template(path)
    c = cd.ClientReport(report, lang, cfg.site_label)
    _set_prepared(layouts, c.T["c_prepared"])
    ctx = decks.Ctx(prs=prs, layouts=layouts, lang=lang, T=c.T, label=cfg.site_label, report=report,
                    report_title=c.T["monthly_report"], period_title=decks.period_title(report, lang),
                    accent=(cfg.accent_color or AMBER), brand=(cfg.brand_color or PERIWINKLE),
                    logo=logo, mark=decks.template_mark(layouts))
    slide_title(ctx, c)
    for sid, payload in c.slides():
        RENDERERS[sid](ctx, c, payload)
    core = prs.core_properties
    core.title = f"{c.T['monthly_report']} · {cfg.site_label} · {ctx.period_title}"
    core.author = "PulseO"
    core.subject = c.T["monthly_report"]
    core.comments = ""
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
